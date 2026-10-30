"""Binômio aquisição + transporte (modelo das abas BINOMIO_COTACOES / RESUMO_COTACOES / BINOMIO_ANP).

Materiais cotados (pétreos, madeira, capa asfáltica comercial, outros):
  preço (t) = preço cotado (t) ou preço (m³) ÷ densidade
  fator     = índice(data-base) ÷ índice(data da cotação) − 1          [índice: Conservação Rodoviária]
  preço reajustado = TRUNC(preço × (1 + fator), 2)
  transporte (t) = DMT pavimentada × custo t·km pav + DMT não pavimentada × custo t·km não pav
  TOTAL = preço reajustado (t) + transporte  →  vence o menor (vantajosidade)

Materiais betuminosos (ANP – Portarias DNIT 1977/2017 e 434/2017):
  aquisição s/ ICMS = preço ANP (R$/t)
  transporte s/ ICMS = TRUNC((26,939 + 0,253·D1 + 0,299·D2 + 0,412·D3)·(1+Kr) + (20,7256 + 0,1603·D)·(1+Kf), 2)
      D1 pavimentada, D2 revestimento primário, D3 leito natural, D fluvial;
      Kr = TRUNC((I − 270,237)/270,237, 3)  [Pavimentação, jul/2014];  Kf = TRUNC((I − 303,752)/303,752, 3) [nov/2016]
  aquisição c/ ICMS = TRUNC(aq / (1 − (ICMS − (PIS+COFINS)·ICMS)), 2)
  transporte c/ ICMS = TRUNC(tr / (1 − (ICMS + COFINS + PIS − (PIS+COFINS)·ICMS)), 2)
  aquisição c/ BDI = TRUNC(aq_icms × (1 + BDI), 2)
  transporte c/ BDI = TRUNC((tr_icms + pedágio/60 + balsa/60) × (1 + 15%), 2)
  TOTAL = aquisição + transporte  →  vence o menor
"""
from __future__ import annotations

import datetime as _dt
import math

import openpyxl

from . import normalizar as nz

INDICE_PETREOS = "Conservação Rodoviária"
INDICE_ASFALTO = "Pavimentação"
BASE_RODOVIARIO = 270.237   # Pavimentação jul/2014
BASE_FLUVIAL = 303.752      # Pavimentação nov/2016

GRUPOS = {"petreo": "Materiais pétreos", "madeira": "Madeira", "betuminoso_comercial": "Betuminoso comercial (capa asfáltica)",
          "outro": "Outros materiais"}


def trunc(v: float, casas: int = 2) -> float:
    f = 10 ** casas
    return math.trunc(round(v * f, 9)) / f


def mes(v) -> str | None:
    if isinstance(v, (_dt.date, _dt.datetime)):
        return v.strftime("%Y-%m")
    s = nz.texto(v)
    if not s:
        return None
    d = nz.data(s)
    return d[:7] if d and len(d) >= 7 and d[4] == "-" else s[:7]


# ------------------------------------------------------------------ índices
def ler_indices(caminho: str) -> dict[str, dict[str, float]]:
    """Lê a aba de índices de reajustamento (blocos com a linha 'Descrição dos índices' + meses)."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        nome_aba = next((n for n in wb.sheetnames if "INDICE" in nz.rotulo(n)), None)
        if not nome_aba:
            raise ValueError("A planilha não tem aba de índices de reajustamento.")
        out: dict[str, dict[str, float]] = {}
        meses = None
        for r in wb[nome_aba].iter_rows(values_only=True):
            if not r:
                continue
            datas = [mes(x) if isinstance(x, (_dt.date, _dt.datetime)) else None for x in r]
            if sum(1 for x in datas if x) >= 3:
                meses = datas
                continue
            nome = next((x for x in r if isinstance(x, str) and x.strip()), None)
            if meses and nome:
                for j, m in enumerate(meses):
                    if m and j < len(r) and isinstance(r[j], (int, float)):
                        out.setdefault(nome.strip().replace("\n", " "), {})[m] = float(r[j])
        return out
    finally:
        wb.close()


def indice(indices: dict, nome: str, m: str | None) -> float | None:
    serie = indices.get(nome) or next((v for k, v in indices.items() if nz.sem_acento(k).upper().startswith(
        nz.sem_acento(nome).upper()[:18])), None)
    if not serie or not m:
        return None
    if m in serie:
        return serie[m]
    anteriores = sorted(k for k in serie if k <= m)  # usa o último índice publicado
    return serie[anteriores[-1]] if anteriores else None


# --------------------------------------------------- materiais cotados
def calcular_material(mat: dict, indices: dict, data_base: str, preco_sicro: dict | None = None) -> dict:
    """mat: nome, codigo, grupo, indice, custo_pav, custo_npav, cotacoes[...]. preco_sicro: {preco, unidade}."""
    nome_ind = mat.get("indice") or INDICE_PETREOS
    i_base = indice(indices, nome_ind, data_base)
    cpav, cnpav = nz.numero(mat.get("custo_pav")) or 0.0, nz.numero(mat.get("custo_npav")) or 0.0
    linhas = []
    cots = list(mat.get("cotacoes") or [])
    if preco_sicro and preco_sicro.get("preco") is not None and not any(c.get("sicro") for c in cots):
        cots.append(dict(fornecedor="SICRO", sicro=True))
    for c in cots:
        dens = nz.numero(c.get("densidade")) or nz.numero(mat.get("densidade")) or 1.0
        un = (c.get("unidade") or "t").lower()
        if c.get("sicro") and preco_sicro:
            un = "m3" if "m" in (preco_sicro.get("unidade") or "").lower() and "3" in (preco_sicro.get("unidade") or "").replace("³", "3") else "t"
            preco = preco_sicro["preco"]
            data_c = data_base
        else:
            preco = nz.numero(c.get("preco"))
            data_c = mes(c.get("data")) or data_base
        if preco is None:
            continue
        p_t = preco if un == "t" else preco / dens
        p_m3 = preco if un == "m3" else preco * dens
        i_c = indice(indices, nome_ind, data_c)
        fator = (i_base / i_c - 1) if i_base and i_c else 0.0
        pt_r, pm3_r = trunc(p_t * (1 + fator)), trunc(p_m3 * (1 + fator))
        dp, dn = nz.numero(c.get("dmt_pav")) or 0.0, nz.numero(c.get("dmt_npav")) or 0.0
        transp = dp * cpav + dn * cnpav
        linhas.append(dict(c, fornecedor=c.get("fornecedor") or ("SICRO" if c.get("sicro") else "—"), unidade=un,
                           densidade=dens, preco_t=round(p_t, 4), preco_m3=round(p_m3, 4), data=data_c,
                           indice_cotacao=i_c, indice_base=i_base, fator=round(fator, 6),
                           preco_t_reaj=pt_r, preco_m3_reaj=pm3_r, transporte=round(transp, 4),
                           total=round(pt_r + transp, 4)))
    venc = min(linhas, key=lambda x: x["total"]) if linhas else None
    for x in linhas:
        x["vencedor"] = x is venc
    return dict(mat, linhas=linhas, vencedor=venc, indice_nome=nome_ind, data_base=data_base,
                sem_indice=i_base is None)


# ----------------------------------------------------------- betuminosos
BDI_TRANSPORTE = {"SD": 0.15, "CD": 0.1867}  # BDI do transporte de betuminosos (planilha de referência: 15% / 18,67%)


def calcular_betuminoso(prod: dict, indices: dict, data_base: str, bdi: float, regime: str = "SD") -> dict:
    p = prod
    pis, cofins = nz.numero(p.get("pis", 0.0065)), nz.numero(p.get("cofins", 0.03))
    chave_tr = "bdi_transporte" if regime == "SD" else "bdi_transporte_cd"
    bdi_tr = nz.numero(p.get(chave_tr)) if p.get(chave_tr) not in (None, "") else BDI_TRANSPORTE[regime]
    ton_viagem = nz.numero(p.get("ton_viagem", 60)) or 60
    i_at = nz.numero(p.get("indice_atual")) or indice(indices, p.get("indice") or INDICE_ASFALTO, data_base)
    base_r, base_f = nz.numero(p.get("base_rodoviario", BASE_RODOVIARIO)), nz.numero(p.get("base_fluvial", BASE_FLUVIAL))
    kr = trunc((i_at - base_r) / base_r, 3) if i_at else 0.0
    kf = trunc((i_at - base_f) / base_f, 3) if i_at else 0.0
    fluvial_sempre = p.get("fluvial_sempre", True)
    linhas = []
    for o in p.get("origens") or []:
        aq = nz.numero(o.get("preco_t"))
        if aq is None:
            continue
        d, d1, d2, d3 = (nz.numero(o.get(k)) or 0.0 for k in ("d_fluvial", "d_pav", "d_rp", "d_ln"))
        rod = (26.939 + 0.253 * d1 + 0.299 * d2 + 0.412 * d3) * (1 + kr)
        flu = (20.7256 + 0.1603 * d) * (1 + kf) if (d > 0 or fluvial_sempre) else 0.0
        tr = trunc(rod + flu)
        icms = nz.numero(o.get("icms")) or 0.0
        icms = icms / 100 if icms > 1 else icms
        aq_i = trunc(aq / (1 - (icms - (pis + cofins) * icms)))
        tr_i = trunc(tr / (1 - (icms + cofins + pis - (pis + cofins) * icms)))
        ped = (nz.numero(o.get("pedagio")) or 0.0) / ton_viagem
        bal = (nz.numero(o.get("balsa")) or 0.0) / ton_viagem
        aq_b = trunc(aq_i * (1 + bdi))
        tr_b = trunc((tr_i + ped + bal) * (1 + bdi_tr))
        linhas.append(dict(o, aq_sem_icms=aq, tr_sem_icms=tr, aq_icms=aq_i, tr_icms=tr_i, pedagio_t=round(ped, 4),
                           balsa_t=round(bal, 4), aq_bdi=aq_b, tr_bdi=tr_b, total=round(aq_b + tr_b, 2)))
    venc = min(linhas, key=lambda x: x["total"]) if linhas else None
    for x in linhas:
        x["vencedor"] = x is venc
    return dict(prod, linhas=linhas, vencedor=venc, kr=kr, kf=kf, indice_atual=i_at, bdi=bdi, data_base=data_base)


# --------------------------------------------- importar de uma planilha PATO
def importar_planilha(caminho: str) -> dict:
    """Lê BINOMIO_COTACOES e BINOMIO_ANP de uma planilha PATO (valores)."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    materiais, produtos = [], []
    try:
        if "BINOMIO_COTACOES" in wb.sheetnames:
            rows = list(wb["BINOMIO_COTACOES"].iter_rows(values_only=True))
            atual = None
            for r in rows:
                r = tuple(r) + (None,) * 25
                nome, cod, forn = r[1], r[2], r[3]
                if isinstance(forn, str) and forn.strip().upper() == "FORNECEDOR":
                    continue
                if isinstance(nome, str) and nome.strip() and not nome.upper().startswith(("VERIFICA", "COTA", "OBSERVA")) and forn:
                    atual = dict(nome=nome.strip(), codigo=nz.texto(cod), grupo="petreo", indice=INDICE_PETREOS,
                                 custo_pav=nz.numero(r[16]), custo_npav=nz.numero(r[17]), cotacoes=[])
                    materiais.append(atual)
                if atual is None or not forn or not isinstance(r[4] if r[4] is not None else 0, (int, float, str)):
                    continue
                if isinstance(forn, str) and forn.strip().upper() == "SICRO":
                    atual["cotacoes"].append(dict(fornecedor="SICRO", sicro=True, densidade=nz.numero(r[5]),
                                                  dmt_pav=nz.numero(r[14]), dmt_npav=nz.numero(r[15])))
                    continue
                preco_t, preco_m3 = nz.numero(r[4]), nz.numero(r[6])
                if preco_t is None and preco_m3 is None:
                    continue
                atual["cotacoes"].append(dict(fornecedor=nz.texto(forn), preco=preco_t if preco_t is not None else preco_m3,
                                              unidade="t" if preco_t is not None else "m3", densidade=nz.numero(r[5]),
                                              data=mes(r[7]), dmt_pav=nz.numero(r[14]), dmt_npav=nz.numero(r[15])))
        if "BINOMIO_ANP" in wb.sheetnames:
            rows = list(wb["BINOMIO_ANP"].iter_rows(values_only=True))
            for i, r in enumerate(rows):
                r = tuple(r) + (None,) * 18
                if isinstance(r[1], str) and i + 3 < len(rows) and isinstance(rows[i + 1][1], str) and \
                        nz.rotulo(rows[i + 1][1]).startswith("ORIGEM"):
                    prod = dict(produto=r[1].strip(), indice=INDICE_ASFALTO, pis=0.0065, cofins=0.03, bdi_transporte=0.15,
                                ton_viagem=60, fluvial_sempre=True, origens=[])
                    for rr in rows[i + 3:i + 10]:
                        rr = tuple(rr) + (None,) * 18
                        if not (isinstance(rr[1], str) and len(rr[1]) <= 4 and "-" in rr[1]):
                            break
                        prod["origens"].append(dict(nome=rr[1], uf=nz.texto(rr[2]), local=nz.texto(rr[3]),
                                                    d_fluvial=nz.numero(rr[4]), d_pav=nz.numero(rr[5]), d_rp=nz.numero(rr[6]),
                                                    d_ln=nz.numero(rr[7]), icms=nz.numero(rr[8]),
                                                    preco_t=(nz.numero(rr[12]) or 0) * 1000 if rr[12] is not None else None,
                                                    pedagio=nz.numero(rr[13]), balsa=nz.numero(rr[14])))
                        prod["pis"] = nz.numero(rr[10]) or prod["pis"]
                        prod["cofins"] = nz.numero(rr[9]) or prod["cofins"]
                    produtos.append(prod)
        return dict(materiais=materiais, produtos=produtos)
    finally:
        wb.close()
