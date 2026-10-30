"""Transportes locais (modelo das abas CROQUI, VMDf, FIT e TRANSPORTES_CONSERVA do PATO).

Croqui: ocorrências (canteiro/centro sede, pedreira, areal, usina, bota-fora…) pelo km no segmento [a, b],
com distâncias de acesso (LN, RP, P) e travessias urbanas.

DMT média de uma ocorrência no km x até a pista (destinos distribuídos ao longo do segmento):
    DMT = ((x − a)² + (b − x)²) / (2·(b − a))
Parcela urbana (Pu), método da planilha de referência: travessia que contém x → distância média dentro dela
((x − u1)² + (u2 − x)²) / (2·(u2 − u1)); demais travessias → metade da extensão. Pr = P − Pu.
Entre duas ocorrências: distância pela rodovia; Pu = extensão das travessias no caminho.

VMDf da viagem: média dos VMD por faixa dos segmentos atravessados, ponderada pela extensão, no sentido da viagem
(da ocorrência para a pista: sentido do lado mais longo do segmento, intervalo [x, x ± DMT]).

FIT (nova metodologia — Manual de Custos de Infraestrutura de Transportes, vol. 05):
    A1 = 70 + 0,0035·VMDf   (s)      A2 = 60·DMTu   (s)      FIT = 1 + (A1 + A2) / (60·DMTp)
aplicado ao transporte em rodovia pavimentada saindo de uma ocorrência, com DMTp ≤ 100 km (como na planilha).

Quadro: momento (t·km) = TRUNC(quantidade anual × fator (t/un) × DMT × FIT, 3), somado por composição de transporte.
"""
from __future__ import annotations

import math
import re

from . import composicoes as cp
from . import normalizar as nz

TIPOS = {"canteiro": "Canteiro / centro sede", "pedreira": "Pedreira", "areal": "Areal", "usina": "Usina",
         "deposito": "Depósito", "bota_fora": "Bota-fora", "emprestimo": "Empréstimo (jazida)",
         "aterro": "Aterro sanitário", "cidade": "Cidade / referência", "outro": "Outro"}
PISTA = "Pista"
FIT_DMT_MAX = 100.0


def trunc(v: float, casas: int = 3) -> float:
    """TRUNC como no Excel (o valor é arredondado a 15 algarismos significativos antes de truncar)."""
    f = 10 ** casas
    return math.trunc(float(f"{float(f'{v:.15g}') * f:.15g}")) / f


def _chave(nome: str | None) -> str:
    return nz.sem_acento(nome or "").upper().replace("-", " ").strip()


def _media_ate_pista(x: float, a: float, b: float) -> float:
    if b <= a:
        return abs(x - a)
    if a <= x <= b:
        return ((x - a) ** 2 + (b - x) ** 2) / (2 * (b - a))
    return abs((a + b) / 2 - x)


def _sobreposicao(p: float, q: float, u1: float, u2: float) -> float:
    lo, hi = min(p, q), max(p, q)
    return max(0.0, min(hi, u2) - max(lo, u1))


def _urbano_exato(x: float, a: float, b: float, u1: float, u2: float, passos: int = 2000) -> float:
    """E[extensão urbana percorrida] com o destino uniforme em [a, b] (integração numérica)."""
    if b <= a:
        return 0.0
    h = (b - a) / passos
    return sum(_sobreposicao(x, a + (i + 0.5) * h, u1, u2) for i in range(passos)) * h / (b - a)


class Croqui:
    def __init__(self, dados: dict, km_ini: float, km_fim: float):
        d = dados or {}
        self.a = nz.numero(d.get("km_ini")) if d.get("km_ini") not in (None, "") else km_ini
        self.b = nz.numero(d.get("km_fim")) if d.get("km_fim") not in (None, "") else km_fim
        self.metodo = d.get("metodo") or "referencia"
        self.fit_max = nz.numero(d.get("fit_dmt_max")) or FIT_DMT_MAX
        self.ocorrencias = [o for o in (d.get("ocorrencias") or []) if o.get("nome")]
        self.travessias = [t for t in (d.get("travessias") or []) if t.get("km_ini") not in (None, "")
                           and t.get("km_fim") not in (None, "")]
        self.vmd = sorted((s for s in (d.get("vmd") or []) if s.get("km_ini") not in (None, "")),
                          key=lambda s: nz.numero(s["km_ini"]))
        self.manuais = {(_chave(r.get("origem")), _chave(r.get("destino"))): r for r in (d.get("rotas_manuais") or [])}
        self._idx = {_chave(o["nome"]): o for o in self.ocorrencias}

    # -------------------------------------------------------------- ocorrências
    def ocorrencia(self, nome: str) -> dict | None:
        k = _chave(nome)
        if k in self._idx:
            return self._idx[k]
        # aceita o tipo ("Areal", "Usina", "Centro Sede" → canteiro)
        alias = {"CENTRO SEDE": "canteiro", "CANTEIRO": "canteiro", "ATERRO SANITARIO": "aterro", "BOTA FORA": "bota_fora",
                 "DEPOSITO": "deposito", "EMPRESTIMO": "emprestimo"}
        tipo = alias.get(k) or next((t for t, r in TIPOS.items() if _chave(r) == k or _chave(t) == k), None)
        return next((o for o in self.ocorrencias if o.get("tipo") == tipo), None) if tipo else None

    @staticmethod
    def _km(o: dict) -> float | None:
        return nz.numero(o.get("km")) if o and o.get("km") not in (None, "") else None

    @staticmethod
    def _acesso(o: dict | None) -> tuple[float, float, float]:
        if not o:
            return 0.0, 0.0, 0.0
        return tuple(nz.numero(o.get(k)) or 0.0 for k in ("acesso_ln", "acesso_rp", "acesso_p"))  # type: ignore

    def _trav(self):
        for t in self.travessias:
            u1, u2 = sorted((nz.numero(t["km_ini"]), nz.numero(t["km_fim"])))
            yield u1, u2, t

    # ---------------------------------------------------------------- DMT
    def rota(self, origem: str, destino: str) -> dict:
        """Distâncias da rota: ln, rp, p (= pr + pu), pr, pu, vmdf, fit, sentido e como foi calculada."""
        man = self.manuais.get((_chave(origem), _chave(destino)))
        if man:
            p = nz.numero(man.get("p")) or 0.0
            pu = nz.numero(man.get("pu")) or 0.0
            r = dict(origem=origem, destino=destino, ln=nz.numero(man.get("ln")) or 0.0, rp=nz.numero(man.get("rp")) or 0.0,
                     p=p, pu=pu, pr=p - pu, metodo="informada", aplicar_fit=bool(man.get("aplicar_fit")),
                     vmdf=nz.numero(man.get("vmdf")))
            return self._com_fit(r, origem, destino)
        o_pista, d_pista = _chave(origem) == "PISTA", _chave(destino) == "PISTA"
        oo = None if o_pista else self.ocorrencia(origem)
        od = None if d_pista else self.ocorrencia(destino)
        if (not o_pista and oo is None) or (not d_pista and od is None):
            falta = origem if (not o_pista and oo is None) else destino
            return dict(origem=origem, destino=destino, ln=0.0, rp=0.0, p=0.0, pu=0.0, pr=0.0, fit=1.0, vmdf=None,
                        erro=f"“{falta}” não está no croqui.", metodo="—", aplicar_fit=False)
        if o_pista and d_pista:
            return dict(origem=origem, destino=destino, ln=0.0, rp=0.0, p=0.0, pu=0.0, pr=0.0, fit=1.0, vmdf=None,
                        metodo="—", aplicar_fit=False)
        ln, rp, p = (sum(v) for v in zip(self._acesso(oo), self._acesso(od)))
        pu = 0.0
        sentido = None
        intervalo = None
        if o_pista or d_pista:
            o = od if o_pista else oo
            x = self._km(o)
            if x is None:  # ocorrência com distância fixa até a pista
                fx = [nz.numero(o.get(k)) or 0.0 for k in ("dist_ln", "dist_rp", "dist_p")]
                if o.get("meia_extensao"):
                    fx[2] = (self.b - self.a) / 2
                ln, rp, p = ln + fx[0], rp + fx[1], p + fx[2]
                r = dict(origem=origem, destino=destino, ln=ln, rp=rp, p=p, pu=0.0, pr=p, metodo="distância fixa",
                         aplicar_fit=False)
                return self._com_fit(r, origem, destino)
            rod = _media_ate_pista(x, self.a, self.b)
            for u1, u2, _ in self._trav():
                if self.metodo == "exato":
                    pu += _urbano_exato(x, self.a, self.b, max(u1, self.a), min(u2, self.b))
                elif u1 <= x <= u2:
                    pu += ((x - u1) ** 2 + (u2 - x) ** 2) / (2 * (u2 - u1)) if u2 > u1 else 0.0
                elif self.a <= u1 and u2 <= self.b:
                    pu += (u2 - u1) / 2
            p += rod
            sentido = 1 if (self.b - x) >= (x - self.a) else -1
            intervalo = (x, x + sentido * rod)
            if o_pista:  # volta: sentido contrário
                sentido = -sentido
                intervalo = (intervalo[1], intervalo[0])
            metodo = "média até a pista"
            aplicar = not o_pista  # sem FIT para recolhimento na pista (como na planilha)
        else:
            x1, x2 = self._km(oo), self._km(od)
            if x1 is None or x2 is None:
                fx = [sum(nz.numero(o.get(k)) or 0.0 for o in (oo, od) if self._km(o) is None)
                      for k in ("dist_ln", "dist_rp", "dist_p")]
                ln, rp, p = ln + fx[0], rp + fx[1], p + fx[2]
                r = dict(origem=origem, destino=destino, ln=ln, rp=rp, p=p, pu=0.0, pr=p, metodo="distância fixa",
                         aplicar_fit=False)
                return self._com_fit(r, origem, destino)
            p += abs(x2 - x1)
            pu = sum(_sobreposicao(x1, x2, u1, u2) for u1, u2, _ in self._trav())
            sentido = 1 if x2 >= x1 else -1
            intervalo = (x1, x2)
            metodo = "entre ocorrências"
            aplicar = True
        r = dict(origem=origem, destino=destino, ln=ln, rp=rp, p=p, pu=pu, pr=p - pu, metodo=metodo,
                 aplicar_fit=aplicar, sentido="CRESCENTE" if sentido == 1 else "DECRESCENTE",
                 vmdf=self.vmdf(intervalo[0], intervalo[1]) if intervalo else None)
        return self._com_fit(r, origem, destino)

    # ---------------------------------------------------------------- VMDf
    def vmdf(self, k0: float, k1: float) -> float | None:
        """VMD por faixa no sentido k0 → k1, ponderado pela extensão em cada segmento."""
        if not self.vmd or k0 == k1:
            return None
        cres = k1 > k0
        lo, hi = sorted((k0, k1))
        tot = soma = 0.0
        segs = self.vmd
        for i, s in enumerate(segs):
            s0 = nz.numero(s["km_ini"])
            s1 = nz.numero(s["km_fim"])
            # o primeiro e o último segmento valem também além do segmento cadastrado
            ini = -math.inf if i == 0 else s0
            fim = math.inf if i == len(segs) - 1 else s1
            ext = max(0.0, min(hi, fim) - max(lo, ini))
            if ext <= 0:
                continue
            v = vmdf_segmento(s, "C" if cres else "D")
            if v is None:
                continue
            tot += ext
            soma += ext * v
        return round(soma / tot) if tot else None

    def _com_fit(self, r: dict, origem: str, destino: str) -> dict:
        r["dmt"] = r["ln"] + r["rp"] + r["p"]
        r["fit"] = fit_transporte(r.get("vmdf"), r["pu"], r["p"]) if (
            r.get("aplicar_fit") and r["p"] > 0 and r["p"] <= self.fit_max and r.get("vmdf") is not None) else 1.0
        if r.get("aplicar_fit") and r.get("vmdf") is None and r["p"] > 0:
            r["aviso"] = "Sem VMD para o FIT: informe o VMD dos segmentos (FIT = 1)."
        if r.get("aplicar_fit") and r["p"] > self.fit_max:
            r["aviso"] = f"DMT pavimentada acima de {self.fit_max:g} km: FIT não aplicado."
        if r["fit"] != 1.0:
            r["a1"] = 70 + 0.0035 * r["vmdf"]
            r["a2"] = 60 * r["pu"]
        return r

    def desenho(self) -> dict:
        """Dados para desenhar o croqui (o SVG é montado na tela)."""
        return dict(km_ini=self.a, km_fim=self.b, ocorrencias=self.ocorrencias, travessias=self.travessias, vmd=self.vmd)


def vmdf_segmento(s: dict, sentido: str) -> float | None:
    """VMD por faixa do segmento: informado (vmdf_c / vmdf_d) ou VMDa do sentido ÷ faixas por sentido;
    sem o sentido, metade do VMDa total."""
    k = "c" if sentido == "C" else "d"
    if s.get(f"vmdf_{k}") not in (None, ""):
        return nz.numero(s[f"vmdf_{k}"])
    faixas = nz.numero(s.get("faixas_sentido")) or 1
    v = nz.numero(s.get(f"vmda_{k}"))
    if v is None and nz.numero(s.get("vmda")) is not None:
        v = nz.numero(s["vmda"]) / 2
    return v / faixas if v is not None else None


def fit_transporte(vmdf: float, dmt_u: float, dmt_p: float) -> float:
    a1 = 70 + 0.0035 * vmdf
    a2 = 60 * dmt_u
    return 1 + (a1 + a2) / (dmt_p * 60)


# ------------------------------------------------------------- quadro de transportes
_RESIDUOS = [  # (palavras na descrição do material, origem, destino)
    (("animais",), PISTA, "Aterro Sanitário"),
    (("pneu", "emborrachad"), PISTA, "Canteiro"),
    (("fresad",), PISTA, "Depósito"),
    (("vegetal", "galho", "arvore", "arbore", "sucata", "entulho", "lixo", "vidro", "engradad", "demolid", "graos",
      "material removido", "residuo", "solo removido"), PISTA, "Bota-fora"),
]
_PETREOS = ("brita", "pedra", "pedrisco", "po de pedra", "rachao", "areia", "cascalho", "seixo")


def material_do_transporte(desc: str | None) -> str:
    """'Areia média - Caminhão basculante 28 m³' → 'Areia média'."""
    return re.split(r"\s+-\s+(?=Caminh|Cavalo|Carreta|Veículo|Veiculo|Guindauto|Carroceria|Betoneira|Basculante)", desc or "")[0].strip()


def origem_padrao(insumo: str, material: str, contexto_desc: str, ocorrencias: list[dict]) -> tuple[str | None, str | None]:
    """Origem/destino sugeridos (o usuário ajusta em 'Origem dos materiais'). None = CIF (sem transporte local)."""
    u = nz.sem_acento(material).lower()
    ctx = nz.sem_acento(contexto_desc or "").lower()
    tipos = {o.get("tipo"): o["nome"] for o in ocorrencias}
    destino = tipos.get("usina", "Usina") if ctx.startswith("usinagem") else PISTA
    for chaves, o, d in _RESIDUOS:
        if any(k in u for k in chaves):
            return o, d
    if any(k in u for k in ("revestimento asfaltico", "mistura", "concreto asfaltico", "usinagem", "pre-misturado",
                            "pre misturado", "cbuq")):
        return tipos.get("usina", "Usina"), PISTA
    if any(k in u for k in ("solo", "material de jazida", "cascalho")) and "solo removido" not in u:
        return tipos.get("emprestimo", "Empréstimo"), destino
    if any(k in u for k in _PETREOS):
        if "areia" in u and "areal" in tipos and "lavada" in u:
            return tipos["areal"], destino
        return tipos.get("pedreira") or tipos.get("canteiro") or "Centro Sede", destino
    return None, None


def gerar_linhas(motor: cp.Motor, servicos: list[dict]) -> list[dict]:
    """Linhas do quadro a partir dos transportes das composições (auxiliares expandidas, fatores multiplicados)."""
    out: list[dict] = []

    def expandir(cod: str, fator: float, servico: str, pilha: tuple):
        if cod in pilha or cod not in motor.comps:
            return
        r = motor.calcular(cod)
        for t in r["transportes"]:
            q = nz.numero(t.get("quant")) or 0.0
            if not q:
                continue
            out.append(dict(servico=servico, insumo=t.get("insumo"), material=material_do_transporte(t.get("descricao")),
                            fator=round(fator * q, 6), transporte=t.get("composicao"), contexto=cod,
                            contexto_desc=r["descricao"]))
        for a in r["auxiliares"]:
            q = nz.numero(a.get("quant")) or 0.0
            if q and a["codigo"] in motor.comps:
                expandir(a["codigo"], fator * q, servico, pilha + (cod,))

    for s in servicos:
        expandir(s["codigo"], 1.0, s["codigo"], ())
    # agrupa linhas iguais (mesmo serviço, insumo, transporte e contexto de destino)
    agr: dict[tuple, dict] = {}
    for x in out:
        k = (x["servico"], x["insumo"], x["transporte"], x["contexto_desc"].lower().startswith("usinagem"))
        if k in agr:
            agr[k]["fator"] = round(agr[k]["fator"] + x["fator"], 6)
        else:
            agr[k] = x
    return list(agr.values())


def superficies_veiculo(motor: cp.Motor, transp: str, extras: dict | None = None) -> dict[str, str]:
    """{'ln': cód, 'rp': cód, 'p': cód} do mesmo veículo da composição de transporte `transp`."""
    sint = dict(motor.base.get("sintetico") or {})
    for c, d in motor.comps.items():
        sint.setdefault(c, dict(descricao=d.get("descricao")))
    por = cp.transportes_por_veiculo(sint, extras)
    desc = (sint.get(transp) or {}).get("descricao") or ""
    m = cp._VEIC_TR.match(desc)
    if m and m.group(1).strip() in por:
        return por[m.group(1).strip()]
    return {"p": transp}


def superficie(desc: str) -> str:
    u = nz.sem_acento(desc or "").lower()
    return "ln" if "leito natural" in u else "rp" if "revestimento primario" in u else "p"


def calcular_quadro(croqui: Croqui, linhas: list[dict], quantidades: dict[str, float], motor: cp.Motor | None = None,
                    descricoes: dict[str, str] | None = None) -> dict:
    """linhas: [{servico, material, fator, origem, destino, transporte, ativo}] → momentos e totais por composição."""
    descricoes = descricoes or {}
    rotas: dict[tuple, dict] = {}
    res, totais = [], {}
    cache_sup: dict[str, dict] = {}
    for ln in linhas:
        if ln.get("ativo") is False or not ln.get("origem") or not ln.get("destino"):
            res.append(dict(ln, momentos={}, cif=not ln.get("origem")))
            continue
        k = (ln["origem"], ln["destino"])
        if k not in rotas:
            rotas[k] = croqui.rota(*k)
        r = rotas[k]
        qtd = nz.numero(ln.get("quantidade")) if ln.get("quantidade") not in (None, "") else quantidades.get(ln["servico"], 0.0)
        peso = (qtd or 0.0) * (nz.numero(ln.get("fator")) or 0.0)
        tr = ln.get("transporte")
        partes = (("p", r["p"], r["fit"]), ("rp", r["rp"], 1.0), ("ln", r["ln"], 1.0))
        if ln.get("transporte_fixo"):  # linha de um bloco da planilha: só a distância da superfície do bloco
            s_bloco = superficie(descricoes.get(tr) or "")
            sup = {s_bloco: tr}
            partes = tuple(x for x in partes if x[0] == s_bloco)
        else:
            if motor is not None and tr not in cache_sup:
                cache_sup[tr] = superficies_veiculo(motor, tr)
            sup = cache_sup.get(tr) or {"p": tr}
        momentos = {}
        for s_, dist, fit in partes:
            if dist > 0 and peso > 0:
                cod = sup.get(s_) or tr
                mom = trunc(peso * dist * fit, 3)
                momentos[cod] = round(momentos.get(cod, 0.0) + mom, 3)
                t = totais.setdefault(cod, dict(codigo=cod, descricao=descricoes.get(cod), superficie=s_, momento=0.0,
                                                linhas=0))
                t["momento"] = round(t["momento"] + mom, 3)
                t["linhas"] += 1
        res.append(dict(ln, quantidade_anual=qtd, peso=round(peso, 6), dmt=r["dmt"], dmt_p=r["p"], dmt_rp=r["rp"],
                        dmt_ln=r["ln"], dmt_u=r["pu"], fit=r["fit"], vmdf=r.get("vmdf"), momentos=momentos,
                        erro=r.get("erro")))
    return dict(linhas=res, rotas=list(rotas.values()), totais=sorted(totais.values(), key=lambda t: t["codigo"]))


# ------------------------------------------------------- importar da planilha PATO
def importar_planilha(caminho: str) -> dict:
    """Lê VMDf (segmentos), TRANSPORTES_CONSERVA (linhas do quadro) e as quantidades do PATO (código → quantidade)."""
    import openpyxl
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    out: dict = dict(vmd=[], linhas=[], quantidades={}, descricoes={})
    try:
        if "PATO" in wb.sheetnames:
            for r in wb["PATO"].iter_rows(min_row=1, values_only=True):
                r = tuple(r) + (None,) * 12
                cod, qtd = r[3], r[9]
                if cod not in (None, "") and isinstance(qtd, (int, float)) and r[2]:
                    c = cp._cod(str(cod))
                    out["quantidades"][c] = round(out["quantidades"].get(c, 0.0) + float(qtd), 6)
                    out["descricoes"][c] = r[4]
        if "VMDf" in wb.sheetnames:
            for r in wb["VMDf"].iter_rows(values_only=True):
                r = tuple(r) + (None,) * 12
                seg, k0, k1, vc, vd = r[4], r[5], r[6], r[9], r[10]
                if isinstance(k0, (int, float, str)) and nz.numero(k0) is not None and nz.numero(k1) is not None \
                        and nz.numero(vc) is not None and str(seg or "").strip().isdigit():
                    out["vmd"].append(dict(segmento=str(seg).strip(), km_ini=nz.numero(k0), km_fim=nz.numero(k1),
                                           vmdf_c=nz.numero(vc), vmdf_d=nz.numero(vd), fonte="planilha PATO (VMDf)"))
        if "TRANSPORTES_CONSERVA" in wb.sheetnames:
            bloco = None
            for r in wb["TRANSPORTES_CONSERVA"].iter_rows(values_only=True):
                r = tuple(r) + (None,) * 16
                b = r[1]
                if b is None:
                    continue
                txt = str(b).strip()
                if r[3] is None and r[6] is None and txt.upper() not in ("CÓDIGO", "TOTAL:") and re.fullmatch(r"[\dA-Z]+", txt):
                    bloco = cp._cod(txt)
                    continue
                if bloco and isinstance(r[6], (int, float)) and r[3]:
                    out["linhas"].append(dict(servico=cp._cod(txt), material=str(r[3]).strip(), fator=float(r[6]),
                                              origem=nz.texto(r[9]), destino=nz.texto(r[10]), transporte=bloco,
                                              transporte_fixo=True, origem_linha="planilha"))
        return out
    finally:
        wb.close()
