"""Operações do sistema (usadas pela API web e pelo servidor MCP)."""
from __future__ import annotations

import copy
import os
import re
import shutil
import tempfile
import time
import uuid
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import range_boundaries

from . import correcao, geo, padrao_dnit
from . import normalizar as nz
from .calculo import medidas, resumo
from .catalogo import CAMPOS_NUM, CATEGORIAS
from .db import PASTA_DADOS, Banco
from .exportador import exportar_inventario, preencher_modelo_pato
from .importador import importar_xlsx
from .xlsx_patch import XlsxPatcher

MODELO_PADRAO = Path(__file__).parent / "modelos" / "PATO_modelo.xlsx"
CAMPOS_TOPO = {"snv", "tipo_pista", "km_ini", "km_fim", "lat_ini", "lon_ini", "lat_fim", "lon_fim",
               "lado", "condicao", "tipo", "obs", "ordem", "origem"}


def normalizar_item(entrada: dict, categoria: str | None = None) -> dict:
    """Aceita um dicionário 'plano' (campos do item + atributos) e separa em colunas e dados."""
    cat = (entrada.get("categoria") or categoria or "").upper()
    if cat not in CATEGORIAS:
        raise ValueError(f"Categoria inválida: {cat!r}. Use uma de: {', '.join(CATEGORIAS)}")
    it: dict[str, Any] = {"categoria": cat}
    dados = dict(entrada.get("dados") or {})
    for k, v in entrada.items():
        if k in ("categoria", "dados", "id", "projeto_id"):
            continue
        if k in CAMPOS_TOPO:
            it[k] = v
        else:
            dados[k] = v
    for k in ("km_ini", "km_fim"):
        if k in it:
            it[k] = nz.numero(it[k])
    for k in ("lat_ini", "lat_fim"):
        if k in it:
            it[k] = nz.coordenada(it[k], "lat")
    for k in ("lon_ini", "lon_fim"):
        if k in it:
            it[k] = nz.coordenada(it[k], "lon")
    if "lado" in it:
        it["lado"] = nz.lado(it["lado"])
    if "condicao" in it:
        it["condicao"] = nz.condicao(it["condicao"])
    if "tipo_pista" in it:
        it["tipo_pista"] = nz.tipo_pista(it["tipo_pista"])
    for k in list(dados):
        if k in CAMPOS_NUM and dados[k] not in (None, ""):
            n = nz.numero(dados[k])
            dados[k] = n if n is not None else dados[k]
        elif dados[k] == "":
            dados[k] = None
    it["dados"] = dados
    return it


class Servico:
    def __init__(self, banco: Banco | None = None, pasta: Path | str | None = None):
        self.pasta = Path(pasta or PASTA_DADOS)
        self.pasta.mkdir(parents=True, exist_ok=True)
        (self.pasta / "exportacoes").mkdir(exist_ok=True)
        (self.pasta / "modelos").mkdir(exist_ok=True)
        self.db = banco or Banco(self.pasta / "pato.db")
        self._previas: dict[str, dict] = {}
        from .custos import Custos
        self.custos = Custos(self)
        from .binomio_servico import Binomios
        self.binomio = Binomios(self)
        from .transporte_servico import Transportes
        self.transporte = Transportes(self)
        from .orcamento_servico import Orcamentos
        self.orcamento = Orcamentos(self)

    # ------------------------------------------------------------- projetos
    def listar_projetos(self) -> list[dict]:
        return self.db.projetos()

    def projeto(self, pid: int) -> dict:
        p = self.db.projeto(pid)
        if not p:
            raise KeyError(f"Projeto {pid} não encontrado")
        return p

    def criar_projeto(self, dados: dict) -> dict:
        dados = {k: (None if v == "" else v) for k, v in dados.items()}
        for k in ("km_ini", "km_fim", "prazo"):
            if dados.get(k) is not None:
                dados[k] = nz.numero(dados[k])
        pid = self.db.criar_projeto(dados)
        self.db.registrar(pid, "criar", dados.get("nome") or "")
        return self.projeto(pid)

    def atualizar_projeto(self, pid: int, dados: dict) -> dict:
        self.projeto(pid)
        for k in ("km_ini", "km_fim", "prazo"):
            if k in dados and dados[k] not in (None, ""):
                dados[k] = nz.numero(dados[k])
        self.db.atualizar_projeto(pid, dados)
        self.db.registrar(pid, "editar projeto", ", ".join(dados))
        return self.projeto(pid)

    def excluir_projeto(self, pid: int) -> None:
        self.db.excluir_projeto(pid)

    def duplicar_projeto(self, pid: int, nome: str | None = None) -> dict:
        p = self.projeto(pid)
        novo = {k: v for k, v in p.items() if k not in ("id", "criado_em", "alterado_em", "n_itens")}
        novo["nome"] = nome or f"{p.get('nome') or 'Projeto'} (cópia)"
        nid = self.db.criar_projeto(novo)
        self.db.inserir_itens(nid, self.db.itens(pid))
        return self.projeto(nid)

    # ------------------------------------------------------------- importar
    def importar_previa(self, caminho: str, nome_arquivo: str | None = None) -> dict:
        res = importar_xlsx(caminho)
        token = uuid.uuid4().hex[:12]
        res["arquivo"] = nome_arquivo or os.path.basename(caminho)
        seguro = re.sub(r"[^\w.\-]+", "_", res["arquivo"])
        guarda = self.pasta / "originais" / f"{token}_{seguro}"
        guarda.parent.mkdir(exist_ok=True)
        shutil.copyfile(caminho, guarda)
        res["copia"] = str(guarda)
        self._previas[token] = res
        # limpa prévias antigas
        while len(self._previas) > 5:
            self._previas.pop(next(iter(self._previas)))
        por_cat: dict[str, int] = {}
        for it in res["itens"]:
            por_cat[it["categoria"]] = por_cat.get(it["categoria"], 0) + 1
        return dict(token=token, arquivo=res["arquivo"], projeto=res["projeto"], por_categoria=por_cat,
                    relatorio=[{k: v for k, v in r.items() if k != "mapa"} for r in res["relatorio"]],
                    ignoradas=res["ignoradas"], total=len(res["itens"]),
                    padronizacao=res["padronizacao_resumo"], vocabulario=res["vocabulario"])

    def importar_confirmar(self, token: str, modo: str = "novo", projeto_id: int | None = None,
                           categorias: list[str] | None = None, nome: str | None = None) -> dict:
        """modo: 'novo' (cria projeto), 'substituir' (troca as categorias importadas) ou 'acrescentar'."""
        res = self._previas.get(token)
        if not res:
            raise KeyError("Prévia expirada — envie o arquivo novamente.")
        itens = [i for i in res["itens"] if not categorias or i["categoria"] in categorias]
        if modo == "novo" or projeto_id is None:
            dados = dict(res["projeto"])
            dados["nome"] = nome or " ".join(x for x in (dados.get("rodovia"), dados.get("segmento")) if x) or res["arquivo"]
            dados["origem_arquivo"] = res["arquivo"]
            dados["arquivo_original"] = res.get("copia")
            pid = self.db.criar_projeto(dados)
        else:
            pid = projeto_id
            self.projeto(pid)
            if modo == "substituir":
                cats = sorted({i["categoria"] for i in itens})
                if cats:
                    self.db.excluir_categoria(pid, cats)
        self.db.inserir_itens(pid, itens)
        self.db.inserir_log(pid, "vocabulário", [dict(campo=v["campo"], de=v["de"], para=v["para"], situacao="padronizado",
                                                     regra="vocabulário normalizado na leitura", quantidade=v["quantidade"])
                                                for v in res.get("vocabulario", [])])
        self.db.registrar(pid, "importar", f"{res['arquivo']} — {len(itens)} itens ({modo})")
        sug = self.gerar_sugestoes(pid)
        return dict(projeto=self.projeto(pid), importados=len(itens), sugestoes=sug["pendentes"],
                    ajustes_automaticos=sug["aplicadas_automaticamente"])

    def importar_arquivo(self, caminho: str, modo: str = "novo", projeto_id: int | None = None,
                         categorias: list[str] | None = None, nome: str | None = None) -> dict:
        prev = self.importar_previa(caminho)
        out = self.importar_confirmar(prev["token"], modo, projeto_id, categorias, nome)
        out["relatorio"] = prev["relatorio"]
        return out

    # ---------------------------------------------------------------- itens
    def itens(self, pid: int, categoria: str | None = None, km_de: float | None = None,
              km_ate: float | None = None, texto: str | None = None) -> list[dict]:
        lst = self.db.itens(pid, categoria)
        if km_de is not None:
            lst = [i for i in lst if (i.get("km_fim") or i.get("km_ini") or -1e9) >= km_de]
        if km_ate is not None:
            lst = [i for i in lst if (i.get("km_ini") if i.get("km_ini") is not None else 1e9) <= km_ate]
        if texto:
            t = nz.sem_acento(texto).upper()
            lst = [i for i in lst if t in nz.sem_acento(str(i)).upper()]
        for i in lst:
            i["medidas"] = medidas(i)
        return lst

    def adicionar_itens(self, pid: int, itens: list[dict], categoria: str | None = None) -> list[dict]:
        self.projeto(pid)
        normal = [normalizar_item(i, categoria) for i in itens]
        for it in normal:
            it.setdefault("origem", "manual")
        ids = self.db.inserir_itens(pid, normal)
        novos = [self.db.item(i) for i in ids]
        self.db.substituir_pendentes(pid, self._sugestoes_padronizacao(novos), apagar=False)
        for s_ in self.db.sugestoes(pid, "pendente", "padronização"):
            if s_["item_id"] in ids and s_.get("certa"):
                self.aplicar_sugestao(s_["id"], automatica=True)
        novos = [self.db.item(i) for i in ids]
        self.db.registrar(pid, "adicionar itens", f"{len(ids)}")
        return novos

    def atualizar_item(self, iid: int, mudancas: dict) -> dict:
        atual = self.db.item(iid)
        if not atual:
            raise KeyError(f"Item {iid} não encontrado")
        norm = normalizar_item({**mudancas, "categoria": atual["categoria"]})
        campos = {k: v for k, v in norm.items() if k in mudancas or k == "dados"}
        if "dados" in mudancas or any(k not in CAMPOS_TOPO for k in mudancas):
            campos["dados"] = norm["dados"]
        else:
            campos.pop("dados", None)
        it = self.db.atualizar_item(iid, campos)
        it["medidas"] = medidas(it)
        return it

    def excluir_itens(self, ids: list[int]) -> int:
        return self.db.excluir_itens(ids)

    # ----------------------------------------------------------- trecho
    def alterar_trecho(self, pid: int, km_ini: float, km_fim: float, recortar: bool = False,
                       segmento: str | None = None) -> dict:
        """Altera o km inicial/final do segmento. Com recortar=True remove o que fica fora e
        corta as feições lineares na divisa (km e coordenadas reinterpoladas)."""
        km_ini, km_fim = sorted((nz.numero(km_ini), nz.numero(km_fim)))
        p = self.projeto(pid)
        seg = segmento or f"km {str(round(km_ini, 3)).replace('.', ',')} ao km {str(round(km_fim, 3)).replace('.', ',')}"
        self.db.atualizar_projeto(pid, {"km_ini": km_ini, "km_fim": km_fim, "segmento": seg})
        removidos = cortados = 0
        if recortar:
            itens = self.db.itens(pid)
            eixo = geo.Eixo(itens)
            excluir = []
            for it in itens:
                a, b = it.get("km_ini"), it.get("km_fim")
                if a is None:
                    continue
                lo, hi = (min(a, b), max(a, b)) if b is not None else (a, a)
                if hi < km_ini or lo > km_fim:
                    if not (hi == lo == km_ini or hi == lo == km_fim):
                        excluir.append(it["id"])
                    continue
                if b is not None and (lo < km_ini or hi > km_fim):
                    mud: dict[str, Any] = {}
                    crescente = b >= a
                    novo_lo, novo_hi = max(lo, km_ini), min(hi, km_fim)
                    na, nb = (novo_lo, novo_hi) if crescente else (novo_hi, novo_lo)
                    if na != a:
                        mud["km_ini"] = na
                        pos = eixo.posicao(na)
                        if pos:
                            mud["lat_ini"], mud["lon_ini"] = pos
                    if nb != b:
                        mud["km_fim"] = nb
                        pos = eixo.posicao(nb)
                        if pos:
                            mud["lat_fim"], mud["lon_fim"] = pos
                    d = dict(it["dados"])
                    for k in ("extensao", "extensao_km", "km_faixa", "km_acost", "area", "area_ha",
                              "area_caiacao", "area_capina", "quantidade"):
                        d.pop(k, None)  # recalculados a partir do km
                    mud["dados"] = d
                    self.db.atualizar_item(it["id"], {**mud})
                    cortados += 1
            removidos = self.db.excluir_itens(excluir)
        self.db.registrar(pid, "alterar trecho",
                          f"{p.get('km_ini')}–{p.get('km_fim')} → {km_ini}–{km_fim}; removidos {removidos}, cortados {cortados}")
        return dict(projeto=self.projeto(pid), removidos=removidos, cortados=cortados)

    def gerar_aceiro(self, pid: int, largura_ate_cerca: float = 30.0, pct_arboreo: float = 0.2,
                     substituir: bool = True) -> dict:
        """Cria o memorial de aceiro (CAD_ACEIRO) a partir das cercas inventariadas."""
        cercas = self.db.itens(pid, "CERCA")
        novos = []
        for c in cercas:
            if c.get("km_ini") is None or c.get("km_fim") is None:
                continue
            novos.append(dict(categoria="ACEIRO", km_ini=c["km_ini"], km_fim=c["km_fim"],
                              lat_ini=c.get("lat_ini"), lon_ini=c.get("lon_ini"), lat_fim=c.get("lat_fim"),
                              lon_fim=c.get("lon_fim"), lado=c.get("lado"), snv=c.get("snv"),
                              dados={"cerca": "SIM", "largura_ate_cerca": largura_ate_cerca, "pct_arboreo": pct_arboreo},
                              origem=f"gerado da cerca {c['id']}"))
        if substituir:
            self.db.excluir_categoria(pid, ["ACEIRO"])
        self.db.inserir_itens(pid, novos)
        self.db.registrar(pid, "gerar aceiro", f"{len(novos)} trechos")
        return dict(gerados=len(novos))

    # ------------------------------------------------------- consultas
    def resumo(self, pid: int) -> dict:
        p = self.projeto(pid)
        itens = self.db.itens(pid)
        out = resumo(p, itens)
        alertas = geo.validar(p, itens)
        out["alertas"] = alertas
        out["n_alertas"] = {n: sum(1 for a in alertas if a["nivel"] == n) for n in ("erro", "aviso", "info")}
        out["projeto"] = p
        out["sugestoes_pendentes"] = len(self.db.sugestoes(pid, "pendente"))
        out["ajustes_aplicados"] = len(self.db.sugestoes(pid, "aplicada"))
        return out

    def validar(self, pid: int) -> list[dict]:
        return geo.validar(self.projeto(pid), self.db.itens(pid))

    def geojson(self, pid: int, categorias: list[str] | None = None) -> dict:
        return geo.geojson(self.projeto(pid), self.db.itens(pid), set(categorias) if categorias else None)

    def kml(self, pid: int) -> str:
        return geo.kml(self.projeto(pid), self.db.itens(pid))

    def historico(self, pid: int) -> list[dict]:
        return self.db.historico(pid)

    # ------------------------------------------------------- exportar
    def modelo_pato(self) -> Path:
        usuario = self.pasta / "modelos" / "PATO_modelo.xlsx"
        return usuario if usuario.exists() else MODELO_PADRAO

    def definir_modelo(self, caminho: str) -> Path:
        destino = self.pasta / "modelos" / "PATO_modelo.xlsx"
        XlsxPatcher(caminho)  # valida
        shutil.copyfile(caminho, destino)
        return destino

    def _nome(self, p: dict, sufixo: str, ext: str) -> Path:
        base = re.sub(r"[^\w\-]+", "_", f"{p.get('rodovia') or 'projeto'}_{p.get('km_ini')}_{p.get('km_fim')}").strip("_")
        return self.pasta / "exportacoes" / f"{base}_{sufixo}_{time.strftime('%Y%m%d_%H%M%S')}.{ext}"

    def exportar_pato(self, pid: int, modelo: str | None = None, limpar_ausentes: bool = True) -> dict:
        p = self.projeto(pid)
        destino = self._nome(p, "PATO", "xlsx")
        log = preencher_modelo_pato(p, self.db.itens(pid), str(modelo or self.modelo_pato()), str(destino),
                                    limpar_ausentes)
        self.db.registrar(pid, "exportar PATO", destino.name)
        return dict(arquivo=str(destino), nome=destino.name, log=log)

    def exportar_inventario(self, pid: int) -> dict:
        p = self.projeto(pid)
        destino = self._nome(p, "inventario", "xlsx")
        exportar_inventario(p, self.db.itens(pid), str(destino))
        return dict(arquivo=str(destino), nome=destino.name)

    def exportar_geo(self, pid: int, formato: str = "geojson") -> dict:
        import json
        p = self.projeto(pid)
        destino = self._nome(p, "inventario", "kml" if formato == "kml" else "geojson")
        if formato == "kml":
            destino.write_text(self.kml(pid), encoding="utf-8")
        else:
            destino.write_text(json.dumps(self.geojson(pid), ensure_ascii=False), encoding="utf-8")
        return dict(arquivo=str(destino), nome=destino.name)

    def conferir_totais(self, arquivo: str) -> dict:
        """Recalcula a planilha exportada (LibreOffice) e devolve os totais do PATO."""
        from .recalculo import recalcular
        with tempfile.TemporaryDirectory() as tmp:
            saida = os.path.join(tmp, "r.xlsx")
            recalcular(arquivo, saida)
            wb = openpyxl.load_workbook(saida, data_only=True, read_only=True)
            try:
                ws = wb["RESUMO_PATO"]
                linhas = []
                for row in ws.iter_rows(min_row=17, max_row=26, min_col=2, max_col=7, values_only=True):
                    if row[0] or row[1]:
                        linhas.append(dict(item=row[0], familia=row[1], sem_deson=row[2], com_deson=row[4]))
                return dict(linhas=linhas)
            finally:
                wb.close()

    # ------------------------------------------- verificação do SNV / correções
    # ---------------------------------------- sugestões de correção (revisão)
    _CAMPOS_KM = ("km_ini", "km_fim")
    _DERIVADOS = ("extensao", "extensao_km", "km_faixa", "km_acost", "area", "area_ha", "area_caiacao",
                  "area_capina", "quantidade")

    def _contagem_base(self) -> list:
        """Tipos de todos os projetos + referência da planilha modelo (para sugerir o 'mais usual')."""
        import json
        ref = Path(__file__).parent / "modelos" / "referencia_tipos.json"
        extra = [tuple(x) for x in json.loads(ref.read_text(encoding="utf-8"))["contagem"]] if ref.exists() else []
        return self.db.contagem_tipos() + extra

    @staticmethod
    def _valor_atual(it: dict, campo: str):
        return it.get(campo) if campo in CAMPOS_TOPO else (it.get("dados") or {}).get(campo)

    def _sugestoes_padronizacao(self, itens: list[dict], ctx: dict | None = None) -> list[dict]:
        """Nomenclatura × álbum DNIT. Grafia ajustável = correção certa (automática); casos sem conversão direta
        recebem a sugestão mais assertiva (mais usual na rodovia/base, ordem dos códigos antigos…)."""
        if ctx is None:
            pid = itens[0]["projeto_id"] if itens else None
            ctx = padrao_dnit.contexto(self.db.itens(pid) if pid else itens, self._contagem_base())
        ordem_conf = {"alta": 0, "média": 1, "baixa": 2}
        out = []
        for it in itens:
            log = padrao_dnit.padronizar_item(copy.deepcopy(it))
            if not log:
                continue
            for sit in (padrao_dnit.PADRONIZADO, padrao_dnit.CONFERIR, padrao_dnit.VERIFICAR):
                ents = [e for e in log if e["situacao"] == sit]
                if not ents:
                    continue
                mud = {e["campo"]: e["para"] for e in ents if e.get("para") not in (None, "")}
                motivos = [e["regra"] for e in ents]
                conf = "alta" if sit == padrao_dnit.PADRONIZADO else "média"
                opcoes = []
                for e in ents:
                    escolha = padrao_dnit.escolher(e, it, ctx) if sit == padrao_dnit.VERIFICAR else None
                    if escolha:
                        valor, just, c = escolha
                        mud[e["campo"]] = valor
                        motivos.append(f"sugestão: {e['campo']} = {valor} — {just}")
                        conf = max(conf, c, key=lambda x: ordem_conf.get(x, 3))
                    ops = list(e.get("opcoes") or [])
                    if escolha and escolha[0] in ops:
                        ops.remove(escolha[0])
                        ops.insert(0, escolha[0])
                    for v in ops:
                        opcoes.append(dict(rotulo=f"{e['campo']} = {v}" + (" (sugerido)" if escolha and v == escolha[0] else ""),
                                           mudancas={e["campo"]: v}))
                if sit == padrao_dnit.VERIFICAR and not mud:
                    conf = "—"
                campos = sorted({e["campo"] for e in ents})
                out.append(dict(item_id=it["id"], tipo="padronização", origem=it.get("origem"), categoria=it["categoria"],
                                km_ini=it.get("km_ini"), km_fim=it.get("km_fim"), situacao=sit, confianca=conf,
                                motivo="; ".join(dict.fromkeys(motivos)), metodo=padrao_dnit.ALBUM,
                                atual={c: self._valor_atual(it, c) for c in campos}, mudancas=mud, opcoes=opcoes,
                                certa=int(sit == padrao_dnit.PADRONIZADO and bool(mud))))
        return out

    def _sugestoes_correcao(self, p: dict, itens: list[dict], oficial, versao,
                            tol_km: float = correcao.TOL_OK) -> tuple[list[dict], dict]:
        an = correcao.analisar(p, itens, oficial, versao, tol_km)
        por_id = {i["id"]: i for i in itens}
        grupos: dict[tuple, list] = {}
        for a in an["acoes"]:
            grupos.setdefault((a["id"], a["situacao"], a["motivo"] if a["situacao"] == correcao.VERIFICAR else ""), []).append(a)
        out = []
        for (iid, sit, _), acs in grupos.items():
            it = por_id[iid]
            a0 = acs[0]
            mud, opcoes = {}, []
            if sit == correcao.REMOVIDO:
                mud = {"_excluir": True}
            elif sit != correcao.VERIFICAR:
                mud = {a["campo"]: a["para"] for a in acs if a["campo"] in correcao.CAMPOS_EDITAVEIS}
            else:
                pt = a0.get("ponta", "ini")
                if a0.get("km_pela_coordenada") is not None:
                    opcoes.append(dict(rotulo=f"Usar o km da coordenada (km {a0['km_pela_coordenada']})",
                                       mudancas={f"km_{pt}": a0["km_pela_coordenada"]}))
                if a0.get("pos_pelo_km"):
                    la, lo = a0["pos_pelo_km"]
                    opcoes.append(dict(rotulo=f"Manter o km e reposicionar a coordenada pelo eixo ({la}, {lo})",
                                       mudancas={f"lat_{pt}": la, f"lon_{pt}": lo}))
                if a0.get("coord_anterior"):
                    la, lo = a0["coord_anterior"]
                    opcoes.append(dict(rotulo=f"Usar a coordenada final do segmento anterior ({la}, {lo})",
                                       mudancas={"lat_ini": la, "lon_ini": lo}))
                if a0["campo"] == "km_ini" and a0.get("km_pela_coordenada") is not None:
                    opcoes.append(dict(rotulo="Retirar do inventário (fora do trecho)", mudancas={"_excluir": True}))
            campos = sorted({c for a in acs for c in a["campo"].split("/") if c in correcao.CAMPOS_EDITAVEIS}
                            | {c for o in opcoes for c in o["mudancas"] if not c.startswith("_")})
            motivos = list(dict.fromkeys(a["motivo"] for a in acs))
            out.append(dict(item_id=iid, tipo="correção", origem=it.get("origem"), categoria=it["categoria"],
                            km_ini=it.get("km_ini"), km_fim=it.get("km_fim"), situacao=sit, confianca=a0["confianca"],
                            motivo="; ".join(motivos), metodo="; ".join(dict.fromkeys(a["metodo"] for a in acs)),
                            atual={c: self._valor_atual(it, c) for c in campos}, mudancas=mud, opcoes=opcoes,
                            certa=int(a0["confianca"] == "alta" and bool(mud) and sit != correcao.VERIFICAR)))
        return out, an

    def gerar_sugestoes(self, pid: int, arquivo_snv: str | None = None, versao_snv: str | None = None,
                        tol_km: float = correcao.TOL_OK, automatico: bool = True, snv_incluido: bool = False) -> dict:
        """Analisa o inventário (SNV, coordenadas × km, trecho e nomenclatura do álbum DNIT).
        Correções certas são aplicadas automaticamente (vão para a aba Ajustes); as demais ficam pendentes
        na aba Revisão já com a sugestão mais assertiva preenchida."""
        import json
        p = self.projeto(pid)
        itens = self.db.itens(pid)
        oficial = correcao.ler_snv_oficial(arquivo_snv, p.get("rodovia")) if arquivo_snv else None
        if not arquivo_snv and snv_incluido:  # SNV oficial incluído no sistema (versão 202610A, MT)
            from . import snv as snv_mod
            oficial = snv_mod.tabela_oficial(p.get("rodovia")) or None
            versao_snv = versao_snv or (snv_mod.versao() if oficial else None)
        if arquivo_snv and not oficial:
            raise ValueError("A planilha do SNV não tem segmentos da rodovia " + str(p.get("rodovia")))
        corr, an = self._sugestoes_correcao(p, itens, oficial, versao_snv, tol_km)
        sai = {s_["item_id"] for s_ in corr if (s_["mudancas"] or {}).get("_excluir") and s_["certa"]}
        ctx = padrao_dnit.contexto(itens, self._contagem_base())
        novas = corr + self._sugestoes_padronizacao([i for i in itens if i["id"] not in sai], ctx)
        if oficial:
            novas.append(dict(item_id=None, tipo="SNV", origem="cadastro de segmentos", categoria="SEG",
                              situacao=correcao.CONFERIR, confianca="alta",
                              motivo=f"Substituir o cadastro de segmentos pela tabela oficial do SNV {versao_snv or ''}".strip(),
                              metodo="SNV oficial", atual={"segmentos": an["snv"]["cadastro_segmentos"]},
                              mudancas={"_segmentos": oficial, "_versao": versao_snv}, opcoes=[], certa=1))
        self.db.substituir_pendentes(pid, novas)
        self.db.atualizar_projeto(pid, {"analise_snv": json.dumps(dict(an["snv"], tolerancias=an["tolerancias"]),
                                                                  ensure_ascii=False)})
        auto = 0
        if automatico:
            pend = self.db.sugestoes(pid, "pendente")
            # ordem: SNV oficial, trecho/coordenadas/SNV dos itens, nomenclatura
            for s_ in sorted((x for x in pend if x.get("certa")), key=lambda x: {"SNV": 0, "correção": 1}.get(x["tipo"], 2)):
                try:
                    self.aplicar_sugestao(s_["id"], automatica=True)
                    auto += 1
                except (KeyError, ValueError):
                    pass
        out = self.resumo_sugestoes(pid)
        out["aplicadas_automaticamente"] = auto
        self.db.registrar(pid, "análise", f"{auto} ajustes automáticos, {out['pendentes']} para revisão")
        return out

    def resumo_sugestoes(self, pid: int) -> dict:
        import json
        p = self.projeto(pid)
        todas = self.db.sugestoes(pid)
        cont: dict = {}
        for s_ in todas:
            k = (s_["status"], s_["tipo"], s_["situacao"])
            cont[k] = cont.get(k, 0) + 1
        return dict(pendentes=sum(1 for s_ in todas if s_["status"] == "pendente"),
                    automaticas=sum(1 for s_ in todas if s_["status"] == "aplicada" and s_.get("automatica")),
                    manuais=sum(1 for s_ in todas if s_["status"] == "aplicada" and not s_.get("automatica")),
                    contagem=[dict(status=k[0], tipo=k[1], situacao=k[2], n=v) for k, v in sorted(cont.items())],
                    snv=json.loads(p["analise_snv"]) if p.get("analise_snv") else None)

    def sugestoes(self, pid: int, status: str | None = None, tipo: str | None = None) -> list[dict]:
        return self.db.sugestoes(pid, status, tipo)

    def aplicar_sugestao(self, sid: int, mudancas: dict | None = None, opcao: int | None = None,
                         automatica: bool = False) -> dict:
        s_ = self.db.sugestao(sid)
        if not s_:
            raise KeyError(f"Sugestão {sid} não encontrada")
        if s_["status"] != "pendente":
            raise ValueError(f"Sugestão já está '{s_['status']}'.")
        if mudancas:
            mud = mudancas
        elif opcao is not None:
            mud = s_["opcoes"][int(opcao)]["mudancas"]
        else:
            mud = s_["mudancas"] or {}
        if not mud:
            raise ValueError("Esta sugestão não tem correção automática: escolha uma opção ou informe o valor.")
        pid = s_["projeto_id"]
        if "_segmentos" in mud:
            antes = self.db.itens(pid, "SEG")
            self._atualizar_segmentos(pid, mud["_segmentos"], mud.get("_versao"))
            self.db.resolver_sugestao(sid, "aplicada", {"_segmentos": True}, {"_snapshot_seg": antes}, automatica=automatica)
            return self.db.sugestao(sid)
        it = self.db.item(s_["item_id"]) if s_["item_id"] else None
        if not it:
            raise KeyError("O item desta sugestão não existe mais.")
        if mud.get("_excluir"):
            self.db.excluir_itens([it["id"]])
            self.db.resolver_sugestao(sid, "aplicada", {"_excluir": True}, {"_snapshot": it}, automatica=automatica)
        else:
            atual = {c: self._valor_atual(it, c) for c in mud}
            limpar = {}
            if any(c in self._CAMPOS_KM for c in mud):
                for c in self._DERIVADOS:
                    if c in (it.get("dados") or {}):
                        atual[c] = it["dados"][c]
                        limpar[c] = None  # recalculado a partir do km
            self.atualizar_item(it["id"], {**mud, **limpar})
            self.db.resolver_sugestao(sid, "aplicada", mud, atual, automatica=automatica)
        if not automatica:
            self.db.registrar(pid, "sugestão aplicada", f"#{sid} {s_['origem'] or ''} {mud}")
        return self.db.sugestao(sid)

    def aplicar_lote(self, pid: int, ids: list[int] | None = None, confianca: str | None = None,
                     tipo: str | None = None) -> dict:
        ok = falhas = 0
        for s_ in self.db.sugestoes(pid, "pendente", tipo):
            if ids is not None and s_["id"] not in ids:
                continue
            if confianca and s_["confianca"] != confianca:
                continue
            if not s_["mudancas"]:
                continue
            try:
                self.aplicar_sugestao(s_["id"])
                ok += 1
            except (KeyError, ValueError):
                falhas += 1
        return dict(aplicadas=ok, falhas=falhas)

    def dispensar_sugestao(self, sid: int) -> dict:
        s_ = self.db.sugestao(sid)
        if not s_ or s_["status"] != "pendente":
            raise ValueError("Só é possível dispensar sugestões pendentes.")
        self.db.resolver_sugestao(sid, "dispensada")
        return self.db.sugestao(sid)

    def desfazer_sugestao(self, sid: int) -> dict:
        s_ = self.db.sugestao(sid)
        if not s_ or s_["status"] not in ("aplicada", "dispensada"):
            raise ValueError("Só é possível desfazer sugestões aplicadas ou dispensadas.")
        if s_["status"] == "dispensada":
            self.db.resolver_sugestao(sid, "pendente")
            return self.db.sugestao(sid)
        atual, apl = s_["atual"] or {}, s_["aplicado"] or {}
        if apl.get("_excluir"):
            snap = dict(atual["_snapshot"])
            novo_id = self.db.inserir_itens(s_["projeto_id"], [snap])[0]
            self.db.resolver_sugestao(sid, "pendente", None, {}, item_id=novo_id)
        elif apl.get("_segmentos"):
            self.db.excluir_categoria(s_["projeto_id"], ["SEG"])
            self.db.inserir_itens(s_["projeto_id"], atual.get("_snapshot_seg") or [])
            self.db.resolver_sugestao(sid, "pendente")
        else:
            self.atualizar_item(s_["item_id"], atual)
            self.db.resolver_sugestao(sid, "pendente")
        self.db.registrar(s_["projeto_id"], "sugestão desfeita", f"#{sid}")
        return self.db.sugestao(sid)

    def log(self, pid: int, tipo: str | None = None) -> list[dict]:
        return self.db.log(pid, tipo)

    def relatorio_pontos(self, pid: int, pendentes: bool = True, automaticos: bool = True, manuais: bool = True,
                         dispensadas: bool = False, destinatario: str | None = None, responsavel: str | None = None,
                         observacoes: str | None = None) -> dict:
        """Relatório (.xlsx) com os pontos a corrigir, para enviar à empresa que levantou o inventário."""
        import json
        from . import relatorio_pontos
        p = self.projeto(pid)
        sel = []
        for s_ in self.db.sugestoes(pid):
            if (s_["status"] == "pendente" and pendentes) or (s_["status"] == "dispensada" and dispensadas) or \
               (s_["status"] == "aplicada" and ((s_.get("automatica") and automaticos) or (not s_.get("automatica") and manuais))):
                sel.append(s_)
        itens = {i["id"]: i for i in self.db.itens(pid)}
        dest = self._nome(p, "pontos_a_corrigir", "xlsx")
        out = relatorio_pontos.gerar(p, sel, itens, str(dest), destinatario, responsavel, observacoes,
                                     json.loads(p["analise_snv"]) if p.get("analise_snv") else None)
        self.db.registrar(pid, "relatório de pontos a corrigir", f"{out['pontos']} pontos · {destinatario or ''}")
        return dict(out, nome=dest.name)

    def exportar_composicoes(self, pid: int) -> dict:
        from .custos import exportar_xlsx
        p = self.projeto(pid)
        dest = self._nome(p, "composicoes", "xlsx")
        exportar_xlsx(self.custos, pid, str(dest))
        return dict(arquivo=str(dest), nome=dest.name)

    def exportar_original(self, pid: int) -> dict:
        """Cópia da planilha de inventário importada com os valores das sugestões aplicadas no sistema."""
        p = self.projeto(pid)
        orig = p.get("arquivo_original")
        if not orig or not os.path.exists(orig):
            raise ValueError("A planilha original deste projeto não está guardada (importe o arquivo novamente).")
        entradas = []
        for s_ in self.db.sugestoes(pid, "aplicada"):
            for campo, valor in (s_["aplicado"] or {}).items():
                if not campo.startswith("_"):
                    entradas.append(dict(origem=s_["origem"], campo=campo, para=valor, km_ini=s_["km_ini"], km_fim=s_["km_fim"]))
        base = re.sub(r"[^\w\-]+", "_", os.path.splitext(p.get("origem_arquivo") or "inventario.xlsx")[0])
        dest = self.pasta / "exportacoes" / f"{base}_CORRIGIDO_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        g = correcao.gravar_no_original(orig, str(dest), entradas)
        return dict(arquivo=str(dest), nome=dest.name, celulas=g["celulas"])

    def _atualizar_segmentos(self, pid: int, oficial: list[dict], versao: str | None) -> None:
        """Substitui o cadastro de segmentos pela tabela oficial do SNV (coordenadas pelo eixo)."""
        p = self.projeto(pid)
        itens = self.db.itens(pid)
        eixo = geo.Eixo(itens)
        k0, k1 = p.get("km_ini"), p.get("km_fim")
        novos = []
        for s in oficial:
            a, b = max(s["km_ini"], k0 if k0 is not None else s["km_ini"]), min(s["km_fim"], k1 if k1 is not None else s["km_fim"])
            if b <= a:
                continue
            pa, pb = eixo.posicao(a), eixo.posicao(b)
            novos.append(dict(categoria="SEG", snv=s["codigo"], tipo_pista="PISTA SIMPLES", km_ini=a, km_fim=b,
                              lat_ini=pa[0] if pa else None, lon_ini=pa[1] if pa else None,
                              lat_fim=pb[0] if pb else None, lon_fim=pb[1] if pb else None,
                              dados={"local_ini": s.get("local_ini"), "local_fim": s.get("local_fim")},
                              origem=f"SNV oficial {versao or ''}".strip()))
        self.db.excluir_categoria(pid, ["SEG"])
        self.db.inserir_itens(pid, novos)
        upd = {"snv": f"{novos[0]['snv']} - {novos[-1]['snv']}" if novos else None}
        if versao:
            upd["snv_versao"] = versao
        self.db.atualizar_projeto(pid, upd)

    # ---------------------------------------------- planilhas (genérico p/ MCP)
    @staticmethod
    def ler_celulas(arquivo: str, aba: str, intervalo: str, formulas: bool = False) -> list[list]:
        wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=not formulas)
        try:
            ws = wb[aba]
            c1, r1, c2, r2 = range_boundaries(intervalo)
            out = []
            for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2, values_only=True):
                out.append([v.isoformat() if hasattr(v, "isoformat") else v for v in row])
            return out
        finally:
            wb.close()

    @staticmethod
    def listar_abas(arquivo: str) -> list[dict]:
        wb = openpyxl.load_workbook(arquivo, read_only=True)
        try:
            return [dict(nome=ws.title, dimensao=ws.calculate_dimension() if hasattr(ws, "calculate_dimension") else None)
                    for ws in wb.worksheets]
        finally:
            wb.close()

    @staticmethod
    def escrever_celulas(arquivo: str, aba: str, valores: dict[str, Any], saida: str | None = None) -> str:
        """Grava valores/fórmulas ('=...') preservando o restante do arquivo."""
        p = XlsxPatcher(arquivo)
        for ref, v in valores.items():
            p.set(aba, ref.upper().replace("$", ""), v)
        destino = saida or arquivo
        tmp = destino + ".tmp"
        p.save(tmp)
        os.replace(tmp, destino)
        return destino
