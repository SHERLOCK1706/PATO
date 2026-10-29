"""Relatório de pontos a corrigir — para enviar à empresa que levantou o inventário.

Lista, por aba e linha da planilha entregue, o problema encontrado e a correção (aplicada no sistema ou sugerida),
com links do mapa e uma coluna para a resposta da empresa.
"""
from __future__ import annotations

import datetime as _dt
from collections import Counter

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .catalogo import CATEGORIAS, ROTULOS

AZUL = PatternFill("solid", fgColor="1F3A5F")
FINA = Side(style="thin", color="BFC5CD")
BORDA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
CORES = {"Corrigir na planilha": "FFF59D", "Confirmar a sugestão": "FFE0B2", "Verificar em campo": "F8BBD0",
         "Retirar da planilha (fora do trecho)": "EF9A9A"}


def _maps(lat, lon):
    try:
        return f"https://www.google.com/maps?q={float(lat)},{float(lon)}" if lat is not None and lon is not None else None
    except (TypeError, ValueError):
        return None


def _fmt(v):
    if v is None or v == "":
        return "—"
    if isinstance(v, float):
        return f"{v:.6f}".rstrip("0").rstrip(".")
    return str(v)


def _acao_empresa(s: dict) -> str:
    m = (s.get("aplicado") if s["status"] == "aplicada" else s.get("mudancas")) or {}
    if m.get("_excluir"):
        return "Retirar da planilha (fora do trecho)"
    if s["status"] == "aplicada":
        return "Corrigir na planilha"
    if m:
        return "Confirmar a sugestão"
    return "Verificar em campo"


def linhas_relatorio(sugestoes: list[dict], itens: dict[int, dict]) -> list[dict]:
    out = []
    for s in sugestoes:
        if s["tipo"] == "SNV":
            continue
        it = itens.get(s.get("item_id")) or ((s.get("atual") or {}).get("_snapshot")) or {}
        m = (s.get("aplicado") if s["status"] == "aplicada" else s.get("mudancas")) or {}
        atual = {k: v for k, v in (s.get("atual") or {}).items() if not k.startswith("_")}
        campos = [c for c in dict.fromkeys(list(m) + list(atual)) if not c.startswith("_")]
        aba, linha = (s.get("origem") or "—!").rsplit("!", 1) if "!" in (s.get("origem") or "") else (s.get("origem") or "—", "")
        # coordenadas: informada (original) e corrigida/sugerida
        lat_i = atual.get("lat_ini", atual.get("lat_fim", it.get("lat_ini")))
        lon_i = atual.get("lon_ini", atual.get("lon_fim", it.get("lon_ini")))
        lat_c = m.get("lat_ini", m.get("lat_fim"))
        lon_c = m.get("lon_ini", m.get("lon_fim"))
        acao = _acao_empresa(s)
        if m.get("_excluir"):
            correcao = "retirar o item da planilha deste segmento"
        elif m:
            correcao = "; ".join(f"{ROTULOS.get(c, c)}: {_fmt(m.get(c))}" for c in campos if c in m)
        else:
            correcao = "—" if not s.get("opcoes") else "opções: " + "; ".join(o["rotulo"] for o in s["opcoes"][:4])
        out.append(dict(
            aba=aba, linha=int(linha) if str(linha).isdigit() else linha,
            categoria=CATEGORIAS.get(s["categoria"], {}).get("nome", s["categoria"]),
            km_ini=s.get("km_ini"), km_fim=s.get("km_fim"), lado=it.get("lado"), snv=it.get("snv"),
            campos=", ".join(ROTULOS.get(c, c) for c in campos) or "item",
            informado="; ".join(f"{ROTULOS.get(c, c)}: {_fmt(atual.get(c))}" for c in campos if c in atual) or "—",
            problema=s.get("motivo"), correcao=correcao, acao=acao,
            origem_correcao=("aplicada automaticamente pelo sistema" if s.get("automatica") else "aplicada pela fiscalização")
            if s["status"] == "aplicada" else "sugerida pelo sistema" if m else "—",
            mapa_informado=_maps(lat_i, lon_i), mapa_corrigido=_maps(lat_c, lon_c) if lat_c is not None else None,
        ))
    out.sort(key=lambda r: (r["aba"], r["linha"] if isinstance(r["linha"], int) else 10 ** 9))
    return out


def gerar(projeto: dict, sugestoes: list[dict], itens: dict[int, dict], saida: str,
          destinatario: str | None = None, responsavel: str | None = None, observacoes: str | None = None,
          snv: dict | None = None) -> dict:
    linhas = linhas_relatorio(sugestoes, itens)
    wb = openpyxl.Workbook()

    # ------------------------------------------------------------ resumo
    ws = wb.active
    ws.title = "Resumo"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 70
    ws["A1"] = "RELATÓRIO DE PONTOS A CORRIGIR NO INVENTÁRIO"
    ws["A1"].font = Font(bold=True, size=14, color="1F3A5F")
    cab = [("Rodovia", projeto.get("rodovia")), ("Trecho", projeto.get("trecho")), ("Subtrecho", projeto.get("subtrecho")),
           ("Segmento", projeto.get("segmento")), ("SNV", projeto.get("snv")),
           ("Planilha de inventário analisada", projeto.get("origem_arquivo")),
           ("Data do inventário", projeto.get("data_inventario")),
           ("Destinatário", destinatario), ("Responsável pela análise", responsavel),
           ("Data do relatório", _dt.date.today().strftime("%d/%m/%Y"))]
    r = 3
    for a, b in cab:
        if b in (None, ""):
            continue
        ws.cell(r, 1, a).font = Font(bold=True)
        ws.cell(r, 2, b)
        r += 1
    r += 1
    ws.cell(r, 1, "Como usar").font = Font(bold=True)
    r += 1
    for t in ("A aba “Pontos a corrigir” lista cada ocorrência pela aba e linha da planilha entregue.",
              "Ação “Corrigir na planilha”: a correção já foi feita na base da fiscalização; repita-a no arquivo do inventário.",
              "Ação “Confirmar a sugestão”: há uma correção provável indicada; confirme ou informe o valor correto.",
              "Ação “Verificar em campo”: não foi possível determinar o valor; levantar novamente no local.",
              "Ação “Retirar da planilha”: o item está fora do segmento (km informado fora do trecho).",
              "Preencha a coluna “Resposta da empresa” e devolva o arquivo junto com o inventário revisado.",
              "Regra adotada: o km informado prevalece; coordenadas divergentes são reposicionadas no eixo da rodovia no km."):
        ws.cell(r, 2, "• " + t).alignment = Alignment(wrap_text=True)
        r += 1
    if observacoes:
        r += 1
        ws.cell(r, 1, "Observações").font = Font(bold=True)
        ws.cell(r, 2, observacoes).alignment = Alignment(wrap_text=True)
        r += 1
    if snv and snv.get("observacoes"):
        r += 1
        ws.cell(r, 1, "SNV").font = Font(bold=True)
        for o in snv["observacoes"]:
            ws.cell(r, 2, o).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 60
            r += 1
    r += 1
    ws.cell(r, 1, "Totais por ação").font = Font(bold=True)
    for k, n in Counter(l["acao"] for l in linhas).most_common():
        r += 1
        ws.cell(r, 1, k).fill = PatternFill("solid", fgColor=CORES.get(k, "FFFFFF"))
        ws.cell(r, 2, n)
    r += 2
    ws.cell(r, 1, "Totais por aba da planilha").font = Font(bold=True)
    for k, n in sorted(Counter(l["aba"] for l in linhas).items()):
        r += 1
        ws.cell(r, 1, k)
        ws.cell(r, 2, n)

    # ------------------------------------------------- pontos a corrigir
    ws = wb.create_sheet("Pontos a corrigir")
    cols = [("Nº", 6), ("Aba da planilha", 30), ("Linha", 7), ("Categoria", 22), ("km inicial", 10), ("km final", 10),
            ("Lado", 6), ("SNV", 13), ("Campo(s)", 18), ("Valor informado", 34), ("Problema identificado", 60),
            ("Correção", 40), ("Ação da empresa", 24), ("Origem da correção", 24), ("Mapa (informado)", 14),
            ("Mapa (corrigido)", 14), ("Resposta da empresa", 34)]
    for j, (t, w) in enumerate(cols, 1):
        c = ws.cell(1, j, t)
        c.fill, c.font, c.border = AZUL, Font(bold=True, color="FFFFFF"), BORDA
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 30
    for i, l in enumerate(linhas, 2):
        vals = [i - 1, l["aba"], l["linha"], l["categoria"], l["km_ini"], l["km_fim"], l["lado"], l["snv"], l["campos"],
                l["informado"], l["problema"], l["correcao"], l["acao"], l["origem_correcao"],
                l["mapa_informado"], l["mapa_corrigido"], None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            c.border = BORDA
            c.alignment = Alignment(wrap_text=j in (10, 11, 12), vertical="top")
            if j in (15, 16) and v:
                c.hyperlink, c.value = v, "abrir mapa"
                c.font = Font(color="1F4E79", underline="single")
        ws.cell(i, 13).fill = PatternFill("solid", fgColor=CORES.get(l["acao"], "FFFFFF"))
    if linhas:
        dv = DataValidation(type="list", formula1='"Corrigido,Mantido (justificar),Verificado em campo,Outro"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"Q2:Q{len(linhas) + 1}")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(2, len(linhas) + 1)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    wb.save(saida)
    return dict(arquivo=saida, pontos=len(linhas), por_acao=dict(Counter(l["acao"] for l in linhas)))
