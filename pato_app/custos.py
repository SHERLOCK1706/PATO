"""Custos do projeto: bases SICRO (estado + data-base + regime), BDI, serviços escolhidos, composições
recalculadas e alterações (preço de insumo, tempo fixo, veículo, transporte) com registro e desfazer."""
from __future__ import annotations

import json
import threading
from pathlib import Path
import zlib
from collections import Counter
from typing import Any

from . import composicoes as cp
from . import normalizar as nz

BDI_PARCELAS_PADRAO = {  # método DNIT (planilha BDI): AC e Lucro sobre o CD; demais parcelas sobre o PV
    "SD": dict(ac=9.0, lucro=12.0, df=1.15, seguros=0.25, riscos=0.5, pis=0.65, cofins=3.0, iss=5.0, cprb=0.0),
    "CD": dict(ac=9.0, lucro=12.0, df=1.15, seguros=0.25, riscos=0.5, pis=0.65, cofins=3.0, iss=5.0, cprb=2.7),
}
ROTULO_PARCELA = dict(ac="Administração central (% CD)", lucro="Lucro (% CD)",
                      df="Despesas financeiras (% sobre PV − lucro)", seguros="Seguros e garantias (% PV)",
                      riscos="Riscos (% PV)", pis="PIS (% PV)", cofins="COFINS (% PV)", iss="ISSQN (% PV)",
                      cprb="Contribuição previdenciária – CPRB (% PV)")
# Lei 14.973/2024: CPRB de 4,5% com base reduzida na transição
CPRB_POR_ANO = {2024: 4.5, 2025: 3.6, 2026: 2.7, 2027: 1.8, 2028: 0.0}
ISS_LIMITE = 5.0


def calcular_bdi(p: dict) -> float:
    """Fórmula da planilha BDI do DNIT:
    BDI = (1 + AC + L − DF·L) / (1 − (DF + R + S + PIS + COFINS + ISS + CPRB)) − 1   (parcelas em %)."""
    ac, l_, df = p.get("ac", 0) / 100, p.get("lucro", 0) / 100, p.get("df", 0) / 100
    num = 1 + ac + l_ - df * l_
    den = 1 - (df + sum(p.get(k, 0) for k in ("seguros", "riscos", "pis", "cofins", "iss", "cprb")) / 100)
    return num / den - 1


def df_pela_selic(selic_aa: float) -> float:
    """DF mensal (%) = [(1 + SELIC)^(1/12) − 1] × 100."""
    return ((1 + selic_aa / 100) ** (1 / 12) - 1) * 100


def calcular_iss(municipios: list[dict]) -> dict:
    """ISSQN efetivo (% do PV) dos municípios da área de influência da obra:
    Σ peso × alíquota × base de cálculo / Σ peso, limitado a 5% (observação da planilha BDI).
    Cada município: nome, uf, aliquota (%), base (% do PV tributável pelo ISS; 100 se não houver dedução de
    materiais na lei municipal), peso (ex.: km do trecho no município ou % do contrato)."""
    linhas, soma_p, soma = [], 0.0, 0.0
    for m in municipios:
        aliq = nz.numero(m.get("aliquota")) or 0.0
        base = nz.numero(m.get("base")) if m.get("base") not in (None, "") else 100.0
        peso = nz.numero(m.get("peso")) if m.get("peso") not in (None, "") else 1.0
        efetivo = aliq * base / 100
        linhas.append(dict(m, aliquota=aliq, base=base, peso=peso, efetivo=round(efetivo, 4)))
        soma_p += peso
        soma += peso * efetivo
    media = soma / soma_p if soma_p else 0.0
    return dict(municipios=linhas, iss_calculado=round(media, 4), iss=round(min(media, ISS_LIMITE), 4),
                limitado=media > ISS_LIMITE)


def _empacotar(b: dict) -> bytes:
    return zlib.compress(json.dumps(b, ensure_ascii=False).encode("utf-8"), 6)


def _desempacotar(raw: bytes) -> dict:
    return json.loads(zlib.decompress(raw).decode("utf-8"))


_TRAVA_EXEMPLO = threading.Lock()
INCLUIDAS = Path(__file__).parent / "modelos" / "bases_sicro"  # datas-base oficiais do SICRO que acompanham o sistema
FONTES = {"sicro": "SICRO oficial (DNIT)", "planilha": "planilha PATO de origem"}


def fonte_base(b: dict) -> str:
    """'sicro' para relatórios oficiais do DNIT; 'planilha' para bases lidas de uma planilha PATO (abas CPU)."""
    return "sicro" if str(b.get("origem") or "").startswith("SICRO oficial") else "planilha"


class Custos:
    def __init__(self, sv):
        self.sv = sv
        self.db = sv.db
        self._cache: dict[int, dict] = {}

    # ------------------------------------------------------------- bases
    def importar_base(self, caminhos: list[str] | str, uf: str | None = None, data_base: str | None = None,
                      descricao: str | None = None, origem: str | None = None) -> list[dict]:
        if isinstance(caminhos, str):
            caminhos = [caminhos]
        pacotes = [c for c in caminhos if c.lower().endswith((".zip", ".7z", ".rar"))]
        soltos = [c for c in caminhos if c not in pacotes]
        grupos = ([soltos] if soltos else []) + [[p] for p in pacotes]  # cada pacote do SICRO é uma data-base
        lidas = []
        for g in grupos:
            bs = cp.ler_bases(g, uf, data_base)
            lidas += list(bs.items())
        if not lidas:
            raise ValueError("Não encontrei composições analíticas nos arquivos enviados (relatório analítico do SICRO "
                             "ou abas CPU de uma planilha PATO).")
        out = []
        for reg, b in lidas:
            if not b.get("uf") or not b.get("data_base"):
                raise ValueError("Não foi possível identificar o estado e a data-base; informe-os na importação.")
            b["descricao"] = descricao or f"SICRO {b['uf']} {b['data_base']} — {cp.REGIMES[reg]}"
            b["origem"] = origem or "; ".join(b.get("fontes", []))
            if any(c.get("oficial") for c in b["composicoes"].values()) and not b["origem"].startswith("SICRO oficial"):
                b["origem"] = "SICRO oficial — " + b["origem"]
            b["n_composicoes"] = len(b["composicoes"])
            b["n_sintetico"] = len(b.get("sintetico") or {})
            bid = self.db.inserir_base(b, _empacotar(b))
            out.append(dict(id=bid, uf=b["uf"], data_base=b["data_base"], regime=reg, descricao=b["descricao"],
                            n_composicoes=b["n_composicoes"], n_sintetico=b["n_sintetico"],
                            bdi_planilha=b.get("bdi_planilha")))
        return out

    def bases(self) -> list[dict]:
        return self.db.bases()

    def excluir_base(self, bid: int) -> None:
        self.db.excluir_base(bid)
        self._cache.pop(bid, None)

    def _base(self, bid: int) -> dict:
        if bid not in self._cache:
            raw = self.db.base_dados(bid)
            if raw is None:
                raise KeyError(f"Base {bid} não encontrada")
            self._cache[bid] = _desempacotar(raw)
            while len(self._cache) > 6:
                self._cache.pop(next(iter(self._cache)))
        return self._cache[bid]

    def garantir_base_exemplo(self) -> None:
        """Na primeira execução, importa a base MT 04/2026 contida na planilha PATO de exemplo e carrega as datas-base
        oficiais do SICRO incluídas no sistema (modelos/bases_sicro) que ainda não estiverem cadastradas."""
        if getattr(self, "_bases_ok", False):
            return
        with _TRAVA_EXEMPLO:  # requisições simultâneas não podem importar a base duas vezes
            if not self.db.bases():
                from .servico import MODELO_PADRAO
                self.importar_base(str(MODELO_PADRAO), descricao=None,
                                   origem="Planilha PATO de exemplo (BR-158/MT) — 132 composições analíticas + lista sintética")
            self._carregar_incluidas()
            self._bases_ok = True

    def _carregar_incluidas(self) -> int:
        idx = INCLUIDAS / "indice.json"
        if not idx.exists():
            return 0
        tem = {(b["uf"], b["data_base"], b["regime"], fonte_base(b)) for b in self.db.bases()}
        n = 0
        for e in json.loads(idx.read_text(encoding="utf-8")):
            if (e["uf"], e["data_base"], e["regime"], "sicro") in tem:
                continue
            raw = (INCLUIDAS / e["arquivo"]).read_bytes()
            self.db.inserir_base(dict(e, n_sintetico=None), raw)
            n += 1
        return n

    def opcoes_data_base(self) -> list[dict]:
        """Datas-base disponíveis: estado, mês e fonte (SICRO oficial ou planilha), com as bases sem/com desoneração."""
        self.garantir_base_exemplo()
        out: dict[tuple, dict] = {}
        for b in self.db.bases():
            k = (b["uf"], b["data_base"], fonte_base(b))
            o = out.setdefault(k, dict(uf=k[0], data_base=k[1], fonte=k[2], chave="|".join(k), bases={},
                                       rotulo=f"{k[0]} {k[1][5:]}/{k[1][:4]} — {FONTES[k[2]]}"))
            o["bases"].setdefault(b["regime"], b["id"])  # a mais recente de cada regime
        return sorted(out.values(), key=lambda o: (o["uf"], o["data_base"], o["fonte"]), reverse=True)

    def _id_base(self, uf: str, data_base: str, regime: str, fonte: str | None) -> int | None:
        cands = [b for b in self.db.bases() if b["uf"] == uf and b["data_base"] == data_base and b["regime"] == regime]
        if fonte:
            pref = [b for b in cands if fonte_base(b) == fonte]
            cands = pref or cands
        return max((b["id"] for b in cands), default=None)

    # ---------------------------------------------------------- projeto
    def config(self, pid: int) -> dict:
        p = self.sv.projeto(pid)
        self.garantir_base_exemplo()
        bases = self.db.bases()
        uf, dbase = p.get("custo_uf"), p.get("custo_data_base")
        fonte = (self.db.dados_projeto(pid, "custo_base") or {}).get("fonte")
        if not uf and bases:
            # sem escolha: a base da planilha de origem (modelo do sistema); senão a data-base mais recente
            plan = [b for b in bases if fonte_base(b) == "planilha"]
            preferida = plan[0] if plan else next((b for b in bases if b["uf"] == "MT"), bases[0])
            uf = preferida["uf"]
            dbase = preferida["data_base"] if plan else max(b["data_base"] for b in bases if b["uf"] == uf)
            fonte = fonte or fonte_base(preferida)
        regime = p.get("custo_regime") or "SD"
        parcelas = json.loads(p["bdi_parcelas"]) if p.get("bdi_parcelas") else None
        bdi = {"SD": p.get("bdi_sd"), "CD": p.get("bdi_cd")}
        for reg in ("SD", "CD"):
            if bdi[reg] is None:
                bid = self._id_base(uf, dbase, reg, fonte) if uf else None
                bp = self._base(bid).get("bdi_planilha") if bid else None
                bdi[reg] = cp.arred(bp, 4) if bp else cp.arred(calcular_bdi(BDI_PARCELAS_PADRAO[reg]), 4)
        disp = sorted({(b["uf"], b["data_base"]) for b in bases}, key=lambda x: (x[0], x[1]), reverse=True)
        par = {r: dict(BDI_PARCELAS_PADRAO[r], **((parcelas or {}).get(r) or {})) for r in ("SD", "CD")}
        extras = {k: (parcelas or {}).get(k) for k in ("municipios", "selic", "ano")}
        if not fonte:
            tem_plan = any(b["uf"] == uf and b["data_base"] == dbase and fonte_base(b) == "planilha" for b in bases)
            fonte = "planilha" if tem_plan else "sicro"
        return dict(uf=uf, data_base=dbase, fonte=fonte, fonte_rotulo=FONTES.get(fonte), opcoes=self.opcoes_data_base(),
                    regime=regime, bdi=bdi, parcelas=par, bdi_extras=extras,
                    rotulos_parcelas=ROTULO_PARCELA, disponiveis=[dict(uf=u, data_base=d,
                    regimes=[b["regime"] for b in bases if b["uf"] == u and b["data_base"] == d]) for u, d in disp])

    def calcular_bdi_completo(self, parcelas: dict | None = None, municipios: list | None = None,
                              selic: float | None = None, ano: int | None = None) -> dict:
        """Calculadora de BDI: aplica SELIC → DF, CPRB do ano (regime com desoneração) e ISSQN dos municípios."""
        par = {r: dict(BDI_PARCELAS_PADRAO[r], **((parcelas or {}).get(r) or {})) for r in ("SD", "CD")}
        notas = []
        if selic not in (None, ""):
            df = round(df_pela_selic(nz.numero(selic)), 4)
            for r in par:
                par[r]["df"] = df
            notas.append(f"DF = (1 + {nz.numero(selic)}%)^(1/12) − 1 = {df:.4f}%")
        if ano not in (None, ""):
            a = int(nz.numero(ano))
            cprb = CPRB_POR_ANO.get(a, 0.0 if a > 2027 else 4.5)
            par["CD"]["cprb"] = cprb
            notas.append(f"CPRB {a}: {cprb}% (Lei 14.973/2024)")
        iss = None
        if municipios:
            iss = calcular_iss(municipios)
            for r in par:
                par[r]["iss"] = iss["iss"]
            notas.append(f"ISSQN efetivo {iss['iss']:.4f}%" + (" (limitado a 5%)" if iss["limitado"] else ""))
        return dict(parcelas=par, bdi={r: cp.arred(calcular_bdi(par[r]), 4) for r in par}, iss=iss, notas=notas,
                    municipios=municipios or [], selic=selic, ano=ano)

    def configurar(self, pid: int, uf: str | None = None, data_base: str | None = None, regime: str | None = None,
                   bdi_sd: float | None = None, bdi_cd: float | None = None, parcelas: dict | None = None,
                   fonte: str | None = None) -> dict:
        mud: dict[str, Any] = {}
        if fonte in FONTES:
            self.db.gravar_dados_projeto(pid, "custo_base", dict(fonte=fonte))
        if uf:
            mud["custo_uf"] = uf
        if data_base:
            mud["custo_data_base"] = data_base
        if regime:
            mud["custo_regime"] = regime
        if parcelas:
            mud["bdi_parcelas"] = json.dumps(parcelas, ensure_ascii=False)
            for reg in ("SD", "CD"):
                if isinstance(parcelas.get(reg), dict):
                    mud[f"bdi_{reg.lower()}"] = cp.arred(calcular_bdi(parcelas[reg]), 4)
        if bdi_sd is not None:
            mud["bdi_sd"] = cp.arred(nz.numero(bdi_sd) / (100 if nz.numero(bdi_sd) > 1 else 1), 4)
        if bdi_cd is not None:
            mud["bdi_cd"] = cp.arred(nz.numero(bdi_cd) / (100 if nz.numero(bdi_cd) > 1 else 1), 4)
        self.db.atualizar_projeto(pid, mud)
        self.db.registrar(pid, "custos: configuração", json.dumps({k: v for k, v in mud.items() if k != "bdi_parcelas"}))
        return self.config(pid)

    def _base_projeto(self, pid: int, regime: str, uf: str | None = None, data_base: str | None = None) -> dict:
        cfg = self.config(pid)
        fonte = cfg.get("fonte") if not (uf or data_base) else None
        bid = self._id_base(uf or cfg["uf"], data_base or cfg["data_base"], regime, fonte)
        if not bid:
            raise ValueError(f"Não há base SICRO {uf or cfg['uf']} {data_base or cfg['data_base']} {cp.REGIMES[regime]}. "
                             "Importe os relatórios do SICRO na aba Custos.")
        return self._base(bid)

    def motor(self, pid: int, regime: str, uf: str | None = None, data_base: str | None = None,
              sem_alteracoes: bool = False) -> cp.Motor:
        base = self._base_projeto(pid, regime, uf, data_base)
        cfg = self.config(pid)
        alts = [] if sem_alteracoes else [a["dados"] for a in self.db.alteracoes_custo(pid, so_ativas=True)]
        return cp.Motor(base, cfg["bdi"][regime], alts, self.db.proprias())

    # -------------------------------------------------------- serviços
    def buscar(self, pid: int, texto: str = "", limite: int = 60) -> list[dict]:
        cfg = self.config(pid)
        reg = cfg["regime"]
        base = self._base_projeto(pid, reg)
        t = nz.sem_acento(texto or "").upper().split()
        escolhidos = {s["codigo"] for s in self.db.servicos(pid)}
        out = []
        fontes = [(c, d.get("descricao"), d.get("unidade"), True) for c, d in base["composicoes"].items()]
        fontes += [(c, d.get("descricao"), d.get("unidade"), False) for c, d in (base.get("sintetico") or {}).items()
                   if c not in base["composicoes"]]
        m = None
        for cod, desc, un, analitica in fontes:
            alvo = nz.sem_acento(f"{cod} {desc}").upper()
            if t and not all(x in alvo for x in t):
                continue
            if analitica:
                m = m or self.motor(pid, reg)
                preco_cd = m.calcular(cod)["custo_direto"]
            else:
                preco_cd = (base["sintetico"].get(cod) or {}).get("preco")
            out.append(dict(codigo=cod, descricao=desc, unidade=un, custo_direto=preco_cd, analitica=analitica,
                            escolhido=cod in escolhidos))
            if len(out) >= limite:
                break
        return out

    def servicos(self, pid: int) -> dict:
        cfg = self.config(pid)
        linhas = []
        motores = {}
        for reg in ("SD", "CD"):
            try:
                motores[reg] = (self.motor(pid, reg), self.motor(pid, reg, sem_alteracoes=True))
            except ValueError:
                pass
        for s in self.db.servicos(pid):
            d = dict(s)
            for reg, (m, m0) in motores.items():
                if s["codigo"] in m.comps:
                    r, r0 = m.calcular(s["codigo"]), m0.calcular(s["codigo"])
                    d.update({f"descricao": r["descricao"], "unidade": r["unidade"],
                              f"cd_{reg}": r["custo_direto"], f"preco_{reg}": r["preco"],
                              f"preco_original_{reg}": r0["preco"], "analitica": True})
                else:
                    sint = (m.base.get("sintetico") or {}).get(s["codigo"]) or {}
                    d.setdefault("descricao", sint.get("descricao"))
                    d.setdefault("unidade", sint.get("unidade"))
                    if sint.get("preco") is not None and m.base.get("sintetico_regime") == reg:
                        d[f"cd_{reg}"] = sint["preco"]
                        d[f"preco_{reg}"] = cp.arred(sint["preco"] * (1 + m.bdi), 2)
                        d[f"preco_original_{reg}"] = d[f"preco_{reg}"]
                    d.setdefault("analitica", False)
            linhas.append(d)
        return dict(config=cfg, servicos=linhas)

    def adicionar_servico(self, pid: int, codigo: str, item: str | None = None) -> dict:
        self.db.adicionar_servico(pid, cp._cod(codigo), item)
        self.db.registrar(pid, "custos: serviço adicionado", codigo)
        return self.servicos(pid)

    def remover_servico(self, pid: int, codigo: str) -> dict:
        self.db.remover_servico(pid, cp._cod(codigo))
        return self.servicos(pid)

    def composicao(self, pid: int, codigo: str, regime: str | None = None) -> dict:
        reg = regime or self.config(pid)["regime"]
        m = self.motor(pid, reg)
        m0 = self.motor(pid, reg, sem_alteracoes=True)
        cod = cp._cod(codigo)
        if cod not in m.comps:
            raise KeyError(f"A composição {codigo} não tem relatório analítico nesta base.")
        r = m.calcular(cod)
        r["preco_original"] = m0.calcular(cod)["preco"]
        r["custo_direto_original"] = m0.calcular(cod)["custo_direto"]
        r["regime"] = reg
        r["bdi_taxa"] = m.bdi
        r["data_base"] = m.base.get("data_base")
        r["uf"] = m.base.get("uf")
        return r

    # ------------------------------------------------------- insumos
    def insumos(self, pid: int, regime: str | None = None, texto: str = "", tipo: str | None = None,
                todos: bool = False) -> list[dict]:
        """Insumos usados nos serviços escolhidos (com as auxiliares), com preço original e atual."""
        reg = regime or self.config(pid)["regime"]
        m = self.motor(pid, reg)
        m0 = self.motor(pid, reg, sem_alteracoes=True)
        cods = list(m.comps) if todos else [s["codigo"] for s in self.db.servicos(pid) if s["codigo"] in m.comps]
        vistos: set = set()
        uso: Counter = Counter()
        info: dict[str, dict] = {}

        def visitar(c, raiz):
            if (c, raiz) in vistos:
                return
            vistos.add((c, raiz))
            r = m.calcular(c)
            r0 = m0.calcular(c)
            for e, e0 in zip(r["equipamentos"], r0["equipamentos"]):
                k = f"EQ:{e['codigo']}"
                uso[k] += 1
                info.setdefault(k, dict(tipo="EQ", codigo=e["codigo"], descricao=e["descricao"], unidade="h",
                                        custo_op=e["custo_op"], custo_imp=e["custo_imp"],
                                        original_op=e0["custo_op"], original_imp=e0["custo_imp"]))
            for x, x0 in zip(r["mao_obra"], r0["mao_obra"]):
                k = f"MO:{x['codigo']}"
                uso[k] += 1
                info.setdefault(k, dict(tipo="MO", codigo=x["codigo"], descricao=x["descricao"], unidade=x["unidade"],
                                        preco=x["custo"], original=x0["custo"]))
            for x, x0 in zip(r["materiais"], r0["materiais"]):
                k = f"MAT:{x['codigo']}"
                uso[k] += 1
                info.setdefault(k, dict(tipo="MAT", codigo=x["codigo"], descricao=x["descricao"], unidade=x["unidade"],
                                        preco=x["preco"], original=x0["preco"]))
            for x in r["tempo_fixo"]:
                if x["retirado"]:
                    continue
                k = f"TF:{x['composicao']}"
                uso[k] += 1
                orig = next((y["preco"] for y in m0.comps[c]["tempo_fixo"] if y["composicao"] == x["composicao"]), x["preco"])
                info.setdefault(k, dict(tipo="TF", codigo=x["composicao"], descricao=self._desc(m, x["composicao"]),
                                        unidade="t", preco=x["preco"], original=orig))
            for a in r["auxiliares"]:
                if a["codigo"] in m.comps:
                    visitar(a["codigo"], raiz)
        for c in cods:
            visitar(c, c)
        t = nz.sem_acento(texto or "").upper()
        out = []
        for k, d in info.items():
            if tipo and d["tipo"] != tipo:
                continue
            if t and t not in nz.sem_acento(f"{d['codigo']} {d['descricao']}").upper():
                continue
            d["usos"] = uso[k]
            d["alterado"] = (d.get("preco") != d.get("original")) if d["tipo"] != "EQ" else (
                d["custo_op"] != d["original_op"] or d["custo_imp"] != d["original_imp"])
            out.append(d)
        return sorted(out, key=lambda d: ("EQ MO MAT TF".split().index(d["tipo"]), d["codigo"]))

    @staticmethod
    def _desc(m, cod):
        return ((m.base.get("sintetico") or {}).get(cod) or {}).get("descricao") or (m.comps.get(cod) or {}).get("descricao") or cod

    # ----------------------------------------------------- alterações
    def _registrar(self, pid, tipo, dados, descricao) -> dict:
        aid = self.db.inserir_alteracao_custo(pid, tipo, dados, descricao)
        self.db.registrar(pid, "custos: alteração", descricao)
        return dict(id=aid, descricao=descricao)

    def alterar_insumo(self, pid: int, tipo: str, codigo: str, valores: dict, regime: str = "AMBOS",
                       motivo: str | None = None) -> dict:
        """tipo: EQ (valores custo_op / custo_imp), MO / MAT / TF (valores preco)."""
        tipo = tipo.upper()
        vals = {k: nz.numero(v) for k, v in valores.items() if nz.numero(v) is not None}
        if not vals:
            raise ValueError("Informe o novo preço.")
        dados = dict(tipo="insumo", insumo_tipo=tipo, codigo=cp._cod(codigo), valores=vals, regime=regime)
        desc = f"Preço do insumo {tipo} {codigo} → " + ", ".join(f"{k} = {v}" for k, v in vals.items()) + \
               (f" ({cp.REGIMES.get(regime, 'ambos os regimes')})") + (f" — {motivo}" if motivo else "")
        return self._registrar(pid, "insumo", dados, desc)

    def opcoes_tempo_fixo(self, pid: int, regime: str | None = None) -> list[dict]:
        reg = regime or self.config(pid)["regime"]
        base = self._base_projeto(pid, reg)
        props = {p["codigo"]: p for p in self.db.proprias()}
        extras = {c: dict(descricao=p["descricao"]) for c, p in props.items()}
        sint = dict(base.get("sintetico") or {})
        m = None
        for c, p in props.items():
            if p.get("categoria") == "tempo_fixo":
                m = m or self.motor(pid, reg)
                sint[c] = dict(descricao=p["descricao"], unidade=p.get("unidade") or "t",
                               preco=m.calcular(c)["custo_direto"] if c in m.comps else None)
        ops = cp.opcoes_tempo_fixo(sint)
        for o in ops:
            o["propria"] = o["codigo"] in props
            o["veiculo"], o["transporte"] = cp.transporte_do_tf(o["descricao"], base.get("sintetico") or {}, extras)
        return sorted(ops, key=lambda o: (not o["propria"], o["descricao"]))

    def _preco_comp(self, pid: int, reg: str, cod: str) -> tuple[float | None, str]:
        m = self.motor(pid, reg)
        if cod in m.comps:
            return m.calcular(cod)["custo_direto"], "analítico"
        s = (m.base.get("sintetico") or {}).get(cod)
        if s and s.get("preco") is not None:
            aviso = "" if m.base.get("sintetico_regime") == reg else f" (preço {cp.REGIMES.get(m.base.get('sintetico_regime'), '')})"
            return s["preco"], "sintético" + aviso
        return None, "sem preço"

    def alterar_tempo_fixo(self, pid: int, acao: str, composicao: str | None = None, insumo: str | None = None,
                           de_composicao: str | None = None, valor: float | None = None,
                           para_composicao: str | None = None, trocar_transporte: bool = True) -> dict:
        """acao: quantidade | retirar | incluir | trocar. composicao=None aplica a todas as composições
        (filtrando por insumo e/ou pela composição de tempo fixo atual `de_composicao`)."""
        alvo = f"composição {composicao}" if composicao else "todas as composições"
        filtro = (f", insumo {insumo}" if insumo else "") + (f", tempo fixo {de_composicao}" if de_composicao else "")
        regs = [r for r in ("SD", "CD") if self.db.base_id(self.config(pid)["uf"], self.config(pid)["data_base"], r)]
        res = []
        for reg in regs:
            base = {"composicao": cp._cod(composicao) if composicao else None, "insumo": insumo,
                    "de_composicao": cp._cod(de_composicao) if de_composicao else None, "regime": reg}
            if acao == "quantidade":
                q = nz.numero(valor)
                if q is None:
                    raise ValueError("Informe a quantidade.")
                dados = dict(base, tipo="tf_quantidade", valor=q)
                desc = f"Tempo fixo: quantidade = {q} ({alvo}{filtro})"
            elif acao == "retirar":
                dados = dict(base, tipo="tf_retirar")
                desc = f"Tempo fixo retirado ({alvo}{filtro})"
            elif acao == "incluir":
                m = self.motor(pid, reg)
                cods_tf = {x["composicao"] for c in m.comps.values() for x in c["tempo_fixo"]}
                precos = {c: self._preco_comp(pid, reg, c)[0] for c in cods_tf}
                dados = dict(base, tipo="tf_incluir", precos={k: v for k, v in precos.items() if v is not None})
                desc = f"Tempo fixo incluído ({alvo}{filtro})"
            elif acao == "trocar":
                para = cp._cod(para_composicao)
                preco, fonte = self._preco_comp(pid, reg, para)
                if preco is None:
                    raise ValueError(f"A composição {para} não tem preço na base.")
                m = self.motor(pid, reg)
                desc_para = self._desc(m, para)
                dados = dict(base, tipo="tf_trocar", para_composicao=para, preco=preco, descricao=desc_para)
                desc = f"Tempo fixo trocado para {para} – {desc_para} ({alvo}{filtro}; preço {preco} {fonte})"
                if trocar_transporte:
                    extras = {p["codigo"]: dict(descricao=p["descricao"]) for p in self.db.proprias()}
                    veic, tr = cp.transporte_do_tf(desc_para, m.base.get("sintetico") or {}, extras)
                    if tr:
                        custos = {}
                        for sup, chave in (("ln", "custo_ln"), ("rp", "custo_rp"), ("p", "custo_p")):
                            if tr.get(sup):
                                custos[chave] = self._preco_comp(pid, reg, tr[sup])[0]
                        tdados = dict(base, tipo="transporte_trocar", para_composicao=tr.get("p") or tr.get("rp") or tr.get("ln"),
                                      **{k: v for k, v in custos.items() if v is not None})
                        tdados["de_tf"] = tdados.pop("de_composicao", None)
                        res.append(self._registrar(pid, "transporte", tdados,
                                                   f"Transporte trocado para {veic} ({alvo}{filtro}) [{cp.REGIMES[reg]}]"))
            else:
                raise ValueError("Ação inválida: use quantidade, retirar, incluir ou trocar.")
            res.append(self._registrar(pid, "tempo_fixo", dados, f"{desc} [{cp.REGIMES[reg]}]"))
        return dict(alteracoes=res)

    def buscar_insumo_base(self, pid: int, tipo: str, texto: str = "", limite: int = 40) -> list[dict]:
        """Insumos da tabela da base (EQ, MO, MAT) ou composições (D) para trocar numa composição."""
        reg = self.config(pid)["regime"]
        base = self._base_projeto(pid, reg)
        t = nz.sem_acento(texto or "").upper().split()
        if tipo == "D":
            fonte = {c: dict(descricao=d.get("descricao"), unidade=d.get("unidade")) for c, d in base["composicoes"].items()}
        else:
            fonte = (base.get("tabelas") or {}).get({"A": "EQ", "B": "MO", "C": "MAT"}[tipo], {})
            if not fonte:  # sem tabela sintética: usa os insumos presentes nas composições
                m = self.motor(pid, reg)
                chave = {"A": "equipamentos", "B": "mao_obra", "C": "materiais"}[tipo]
                fonte = {r["codigo"]: r for c in m.comps.values() for r in c[chave]}
        out = []
        for cod, d in fonte.items():
            if t and not all(x in nz.sem_acento(f"{cod} {d.get('descricao')}").upper() for x in t):
                continue
            out.append(dict(codigo=cod, descricao=d.get("descricao"), unidade=d.get("unidade"),
                            preco=d.get("preco", d.get("custo")), custo_op=d.get("custo_op"), custo_imp=d.get("custo_imp")))
            if len(out) >= limite:
                break
        return out

    def _dados_insumo(self, pid: int, reg: str, secao: str, codigo: str) -> dict:
        base = self._base_projeto(pid, reg)
        m = self.motor(pid, reg)
        if secao == "D":
            if codigo not in m.comps:
                raise ValueError(f"A composição auxiliar {codigo} não tem analítico na base.")
            c = m.comps[codigo]
            return dict(descricao=c["descricao"], unidade=c["unidade"])
        tab = (base.get("tabelas") or {}).get({"A": "EQ", "B": "MO", "C": "MAT"}[secao], {})
        d = tab.get(codigo)
        if not d:  # procura nas composições
            chave = {"A": "equipamentos", "B": "mao_obra", "C": "materiais"}[secao]
            d = next((r for c in m.comps.values() for r in c[chave] if r["codigo"] == codigo), None)
        if not d:
            raise ValueError(f"Insumo {codigo} não encontrado na base {base.get('uf')} {base.get('data_base')}.")
        if secao == "A":
            return dict(descricao=d.get("descricao"), custo_op=nz.numero(d.get("custo_op")), custo_imp=nz.numero(d.get("custo_imp")))
        if secao == "B":
            return dict(descricao=d.get("descricao"), custo=nz.numero(d.get("preco", d.get("custo"))), unidade=d.get("unidade"))
        return dict(descricao=d.get("descricao"), preco=nz.numero(d.get("preco")), unidade=d.get("unidade"))

    def trocar_insumo(self, pid: int, secao: str, de_codigo: str, para_codigo: str, composicao: str | None = None,
                      quant: float | None = None) -> dict:
        """Troca um insumo (A equipamento, B mão de obra, C material, D auxiliar) por outro da base."""
        secao = secao.upper()
        cfg = self.config(pid)
        res = []
        alvo = f"composição {composicao}" if composicao else "todas as composições"
        for reg in ("SD", "CD"):
            if not self.db.base_id(cfg["uf"], cfg["data_base"], reg):
                continue
            dados = dict(tipo="insumo_trocar", secao=secao, de_codigo=cp._cod(de_codigo), para_codigo=cp._cod(para_codigo),
                         composicao=cp._cod(composicao) if composicao else None, regime=reg,
                         quant=nz.numero(quant) if quant not in (None, "") else None,
                         **self._dados_insumo(pid, reg, secao, cp._cod(para_codigo)))
            res.append(self._registrar(pid, "troca de insumo", dados,
                                       f"Insumo {de_codigo} trocado por {para_codigo} – {dados.get('descricao')} ({alvo}) [{cp.REGIMES[reg]}]"))
        return dict(alteracoes=res)

    # ------------------------------------------------ composições próprias
    def proprias(self, pid: int | None = None) -> list[dict]:
        lst = self.db.proprias()
        if pid:
            for reg in ("SD", "CD"):
                try:
                    m = self.motor(pid, reg)
                except ValueError:
                    continue
                for p in lst:
                    if p["codigo"] in m.comps:
                        r = m.calcular(p["codigo"])
                        p[f"cd_{reg}"], p[f"preco_{reg}"] = r["custo_direto"], r["preco"]
                        p["avisos"] = m.avisos.get(p["codigo"], [])
        return lst

    def salvar_propria(self, c: dict) -> dict:
        cod = (c.get("codigo") or "").strip().upper()
        if not cod:
            raise ValueError("Informe o código da composição (ex.: PN01).")
        base_cod = cp._cod(cod)
        # não pode coincidir com composição oficial do SICRO (códigos numéricos); códigos próprios vindos de uma
        # planilha PATO (ex.: PN02) podem ser redefinidos — a composição própria prevalece sobre a da base
        if base_cod.isdigit():
            for b in self.db.bases():
                if base_cod in (self._base(b["id"]).get("composicoes") or {}) or base_cod in (self._base(b["id"]).get("sintetico") or {}):
                    raise ValueError(f"O código {cod} já existe no SICRO; use outro (ex.: PN01).")
                break
        dados = dict(c.get("dados") or {})
        for k in ("producao", "fic", "fit"):
            if k in dados:
                dados[k] = nz.numero(dados[k])
        self.db.salvar_propria(dict(codigo=base_cod, descricao=c.get("descricao"), unidade=c.get("unidade"),
                                    categoria=c.get("categoria") or "servico", dados=dados, origem=c.get("origem"),
                                    observacoes=c.get("observacoes")))
        return next(p for p in self.db.proprias() if p["codigo"] == base_cod)

    def excluir_propria(self, codigo: str) -> None:
        self.db.excluir_propria(cp._cod(codigo))

    def copiar_para_propria(self, pid: int, codigo: str, novo_codigo: str, descricao: str | None = None) -> dict:
        """Cria uma composição própria a partir de uma composição analítica (SICRO ou própria)."""
        reg = self.config(pid)["regime"]
        m = self.motor(pid, reg, sem_alteracoes=True)
        c = m.comps.get(cp._cod(codigo))
        if not c:
            raise KeyError(f"Composição {codigo} não tem analítico na base.")
        dados = dict(producao=c["producao"], fic=c.get("fic_ef", c["fic"]) if c.get("fic_base") == "execucao" else c["fic"],
                     fit=c["fit"], fic_base=c.get("fic_base") or "improdutivo",
                     equipamentos=[{k: e[k] for k in ("codigo", "quant", "util_op", "util_imp")} for e in c["equipamentos"]],
                     mao_obra=[{k: x[k] for k in ("codigo", "quant")} for x in c["mao_obra"]],
                     materiais=[{k: x[k] for k in ("codigo", "quant")} for x in c["materiais"]],
                     auxiliares=[{k: x.get(k) for k in ("codigo", "quant", "descricao", "unidade")} for x in c["auxiliares"]],
                     tempo_fixo=[{k: x[k] for k in ("insumo", "composicao", "quant")} for x in c["tempo_fixo"] if not x["retirado"]],
                     transportes=[{k: x[k] for k in ("insumo", "composicao", "quant", "custo_ln", "custo_rp", "custo_p",
                                                     "dmt_ln", "dmt_rp", "dmt_p", "descricao")} for x in c["transportes"]])
        if c.get("fic_base") != "execucao" and c.get("fic_ef") is not None:
            dados["fic"] = c["fic_ef"]
        cat = "tempo_fixo" if (c["descricao"] or "").startswith("Carga, manobra") else \
              "transporte" if (c["descricao"] or "").startswith("Transporte") else "servico"
        return self.salvar_propria(dict(codigo=novo_codigo, descricao=descricao or c["descricao"], unidade=c["unidade"],
                                        categoria=cat, dados=dados, origem=f"cópia de {codigo}"))

    def criar_exemplo_28m3(self) -> list[dict]:
        """PN01–PN04 conforme a 1ª Análise Técnica Final nº 32 – FGV/2024 (SEPN/CGCIT/DNIT): caminhão basculante
        com caçamba estanque de 28 m³ (cavalo mecânico + 2 caçambas de 14 m³; produção dobrada)."""
        ref = "1ª ATF nº 32 – FGV/2024 – SEPN/CGCIT/DNIT (BR-364/AC)"
        eq_tr = [dict(codigo="E9018", quant=1, util_op=1, util_imp=0), dict(codigo="A9345", quant=2, util_op=1, util_imp=0)]
        out = []
        # PN02 segue o PATO e a 5915320 do SICRO (FIC 0,0559 sobre o custo improdutivo)
        for cod, sup, prod, fic, fic_base in (("PN01", "rodovia em leito natural", 697.20, 0.03568, "execucao"),
                                              ("PN02", "rodovia em revestimento primário", 871.50, 0.0559, "improdutivo"),
                                              ("PN03", "rodovia pavimentada", 941.22, 0.0, "execucao")):
            out.append(self.salvar_propria(dict(
                codigo=cod, unidade="tkm", categoria="transporte",
                descricao=f"Transporte com caminhão basculante com caçamba estanque com capacidade de 28 m³ - {sup}",
                dados=dict(producao=prod, fic=fic, fit=0, fic_base=fic_base, equipamentos=eq_tr),
                origem=ref, observacoes="Troca de equipamentos (cavalo mecânico com dolly + 2 caçambas de 14 m³ = 28 m³); "
                                        "produção duplicada em relação à composição de 14 m³.")))
        out.append(self.salvar_propria(dict(
            codigo="PN04", unidade="t", categoria="tempo_fixo",
            descricao="Carga, manobra e descarga de agregados ou solos em caminhão basculante de 28 m³ - carga com "
                      "carregadeira de 3,40 m³ e descarga livre",
            dados=dict(producao=914.32, fic=0, fit=0, fic_base="execucao",
                       equipamentos=[dict(codigo="E9018", quant=1, util_op=0.77, util_imp=0.23),
                                     dict(codigo="A9345", quant=6, util_op=0.77, util_imp=0.23),
                                     dict(codigo="E9584", quant=2, util_op=1, util_imp=0)]),
            origem=ref, observacoes="Baseada na 5914351 (basculante 14 m³): caminhão substituído por caçamba e duplicadas "
                                    "as quantidades de equipamentos e a produção, mantendo o tempo de produção.")))
        return out

    def alterar_dmt(self, pid: int, composicao: str | None, insumo: str | None, dmt_ln=None, dmt_rp=None, dmt_p=None) -> dict:
        vals = {k: nz.numero(v) for k, v in (("dmt_ln", dmt_ln), ("dmt_rp", dmt_rp), ("dmt_p", dmt_p)) if v not in (None, "")}
        dados = dict(tipo="transporte_dmt", composicao=cp._cod(composicao) if composicao else None, insumo=insumo,
                     regime="AMBOS", **vals)
        alvo = f"composição {composicao}" if composicao else "todas as composições"
        return self._registrar(pid, "transporte", dados, f"DMT do transporte {insumo or ''} ({alvo}): "
                               + ", ".join(f"{k} = {v}" for k, v in vals.items()))

    def alteracoes(self, pid: int) -> list[dict]:
        return self.db.alteracoes_custo(pid)

    def ativar(self, pid: int, aid: int, ativa: bool) -> dict:
        self.db.ativar_alteracao_custo(aid, ativa)
        self.db.registrar(pid, "custos: alteração " + ("reativada" if ativa else "desfeita"), str(aid))
        return dict(id=aid, ativa=ativa)

    # ----------------------------------------------------- data-base
    def comparar_data_base(self, pid: int, uf: str, data_base: str) -> dict:
        cfg = self.config(pid)
        linhas = []
        for s in self.db.servicos(pid):
            d = dict(codigo=s["codigo"])
            for reg in ("SD", "CD"):
                try:
                    ma = self.motor(pid, reg)
                    mb = self.motor(pid, reg, uf, data_base)
                except ValueError:
                    continue
                pa = ma.calcular(s["codigo"])["preco"] if s["codigo"] in ma.comps else None
                pb = mb.calcular(s["codigo"])["preco"] if s["codigo"] in mb.comps else None
                d.update({f"atual_{reg}": pa, f"novo_{reg}": pb, "descricao": (mb.comps.get(s["codigo"]) or ma.comps.get(s["codigo"]) or {}).get("descricao")})
                if pa and pb:
                    d[f"var_{reg}"] = round((pb / pa - 1) * 100, 2)
            linhas.append(d)
        return dict(de=dict(uf=cfg["uf"], data_base=cfg["data_base"]), para=dict(uf=uf, data_base=data_base), servicos=linhas)


# --------------------------------------------------------------- exportação
def exportar_xlsx(cu: "Custos", pid: int, saida: str) -> str:
    """Planilha com os preços dos serviços escolhidos, as composições analíticas (formato SICRO) nos dois
    regimes, os insumos alterados, o BDI e o registro de alterações."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    azul = PatternFill("solid", fgColor="1F3A5F")
    cinza = PatternFill("solid", fgColor="E5E7EB")
    amarelo = PatternFill("solid", fgColor="FFF59D")
    branco = Font(bold=True, color="FFFFFF")
    neg = Font(bold=True)
    p = cu.sv.projeto(pid)
    cfg = cu.config(pid)
    dados = cu.servicos(pid)
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Serviços"
    ws["A1"] = "PREÇOS UNITÁRIOS DOS SERVIÇOS"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"{p.get('rodovia') or ''} {p.get('segmento') or ''}".strip()
    ws["A3"] = (f"Base SICRO {cfg['uf']} — data-base {cfg['data_base']} · BDI sem desoneração {cfg['bdi']['SD'] * 100:.2f}% · "
                f"com desoneração {cfg['bdi']['CD'] * 100:.2f}%")
    cab = ["Código", "Descrição", "Unid.", "Custo direto SD", "Preço SD", "Preço SD original", "Custo direto CD",
           "Preço CD", "Preço CD original", "Analítica"]
    for j, t in enumerate(cab, 1):
        c = ws.cell(5, j, t)
        c.fill, c.font = azul, branco
        c.alignment = Alignment(wrap_text=True)
    for i, s in enumerate(dados["servicos"], 6):
        vals = [s["codigo"], s.get("descricao"), s.get("unidade"), s.get("cd_SD"), s.get("preco_SD"),
                s.get("preco_original_SD"), s.get("cd_CD"), s.get("preco_CD"), s.get("preco_original_CD"),
                "sim" if s.get("analitica") else "não (preço sintético)"]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            if j in (4, 5, 6, 7, 8, 9):
                c.number_format = "#,##0.00"
        for reg, j in (("SD", 5), ("CD", 8)):
            if s.get(f"preco_{reg}") != s.get(f"preco_original_{reg}"):
                ws.cell(i, j).fill = amarelo
    for j, w in enumerate([12, 70, 7, 14, 12, 14, 14, 12, 14, 18], 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    for reg in ("SD", "CD"):
        try:
            m = cu.motor(pid, reg)
        except ValueError:
            continue
        ws = wb.create_sheet(f"CPU {reg}")
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 60
        for col in "CDEFGHIJK":
            ws.column_dimensions[col].width = 12
        r = 1
        cods, vistos = [s["codigo"] for s in dados["servicos"] if s["codigo"] in m.comps], set()

        def inclui(c):
            if c in vistos or c not in m.comps:
                return
            vistos.add(c)
            cods_ord.append(c)
            for a in m.comps[c]["auxiliares"]:
                inclui(a["codigo"])
        cods_ord: list[str] = []
        for c in cods:
            inclui(c)
        for cod in cods_ord:
            x = m.calcular(cod)
            ws.cell(r, 1, f"{x['codigo']} — {x['descricao']}").font = Font(bold=True, color="1F3A5F")
            ws.cell(r, 9, "Unidade:").font = neg
            ws.cell(r, 10, x["unidade"])
            r += 1
            ws.cell(r, 1, f"SICRO {m.base.get('uf')} {m.base.get('data_base')} — {REGIMES_TXT[reg]}")
            r += 1

            def titulo(t, cols):
                nonlocal r
                ws.cell(r, 1, t).font = neg
                for j, h in cols:
                    ws.cell(r, j, h).font = neg
                for j in range(1, 12):
                    ws.cell(r, j).fill = cinza
                r += 1
            titulo("A - EQUIPAMENTOS", [(6, "Quant."), (7, "Util. op."), (8, "Util. improd."), (9, "Custo op."),
                                        (10, "Custo improd."), (11, "Custo horário")])
            for e in x["equipamentos"]:
                for j, v in enumerate([e["codigo"], e["descricao"], None, None, None, e["quant"], e["util_op"],
                                       e["util_imp"], e["custo_op"], e["custo_imp"], e["custo"]], 1):
                    ws.cell(r, j, v)
                r += 1
            titulo("B - MÃO DE OBRA", [(8, "Quant."), (9, "Unid."), (10, "Custo horário"), (11, "Total")])
            for e in x["mao_obra"]:
                for j, v in enumerate([e["codigo"], e["descricao"], None, None, None, None, None, e["quant"],
                                       e["unidade"], e["custo"], e["total"]], 1):
                    ws.cell(r, j, v)
                r += 1
            for rot, val in (("Custo horário de execução", x["exec_horario"]), ("Produção da equipe", x["producao"]),
                             ("Custo unitário de execução", x["exec_unitario"]), ("Custo FIC", x["fic"]),
                             ("Custo FIT", x["fit"])):
                ws.cell(r, 2, rot)
                ws.cell(r, 11, val)
                r += 1
            for sec, chave in (("C - MATERIAIS", "materiais"), ("D - ATIVIDADES AUXILIARES", "auxiliares"),
                               ("E - TEMPO FIXO", "tempo_fixo")):
                titulo(sec, [(8, "Quant."), (9, "Unid."), (10, "Preço"), (11, "Custo")])
                for e in x[chave]:
                    cod_ = f"{e['insumo']} {e['composicao']}" if chave == "tempo_fixo" else e["codigo"]
                    desc = e["descricao"] + (" (retirado)" if e.get("retirado") else "")
                    for j, v in enumerate([cod_, desc, None, None, None, None, None, e["quant"], e.get("unidade"),
                                           e["preco"], e["total"]], 1):
                        ws.cell(r, j, v)
                    r += 1
            titulo("F - TRANSPORTES", [(3, "Quant."), (5, "Custo LN"), (6, "Custo RP"), (7, "Custo P"), (8, "DMT LN"),
                                       (9, "DMT RP"), (10, "DMT P"), (11, "Custo")])
            for e in x["transportes"]:
                for j, v in enumerate([f"{e['insumo']} {e['composicao']}", e["descricao"], e["quant"], e.get("unidade"),
                                       e["custo_ln"], e["custo_rp"], e["custo_p"], e["dmt_ln"], e["dmt_rp"], e["dmt_p"],
                                       e["total"]], 1):
                    ws.cell(r, j, v)
                r += 1
            for rot, val in (("SUBTOTAL", x["subtotal"]), ("CUSTO UNITÁRIO DIRETO TOTAL", x["custo_direto"]),
                             (f"BDI ({m.bdi * 100:.2f}%)", x["bdi"]), ("PREÇO UNITÁRIO TOTAL", x["preco"])):
                ws.cell(r, 1, rot).font = neg
                c = ws.cell(r, 11, val)
                c.font, c.number_format = neg, "#,##0.00"
                r += 1
            r += 2

    ws = wb.create_sheet("Insumos alterados")
    for j, t in enumerate(["Regime", "Tipo", "Código", "Descrição", "Unid.", "Original", "Atual", "Usos"], 1):
        c = ws.cell(1, j, t)
        c.fill, c.font = azul, branco
    r = 2
    for reg in ("SD", "CD"):
        try:
            ins = cu.insumos(pid, reg)
        except ValueError:
            continue
        for i in ins:
            if not i["alterado"]:
                continue
            orig = f"{i['original_op']} / {i['original_imp']}" if i["tipo"] == "EQ" else i["original"]
            atual = f"{i['custo_op']} / {i['custo_imp']}" if i["tipo"] == "EQ" else i["preco"]
            for j, v in enumerate([REGIMES_TXT[reg], i["tipo"], i["codigo"], i["descricao"], i["unidade"], orig, atual,
                                   i["usos"]], 1):
                ws.cell(r, j, v)
            r += 1
    ws.column_dimensions["D"].width = 60

    ws = wb.create_sheet("BDI e alterações")
    ws["A1"] = "BDI"
    ws["A1"].font = neg
    ws["A2"], ws["B2"] = "Sem desoneração", cfg["bdi"]["SD"]
    ws["A3"], ws["B3"] = "Com desoneração", cfg["bdi"]["CD"]
    ws["B2"].number_format = ws["B3"].number_format = "0.00%"
    ws["A5"] = "Alterações de custo (ativas)"
    ws["A5"].font = neg
    for i, a in enumerate([a for a in cu.alteracoes(pid) if a["ativa"]], 6):
        ws.cell(i, 1, a["quando"])
        ws.cell(i, 2, a["descricao"])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 110
    wb.save(saida)
    return saida


REGIMES_TXT = {"SD": "sem desoneração", "CD": "com desoneração"}
