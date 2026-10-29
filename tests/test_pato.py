"""Testes do sistema PATO. Rodar: python -m pytest -q   (o teste com LibreOffice é pulado se não houver)."""
import os
from pathlib import Path

import openpyxl
import pytest

from pato_app import normalizar as nz
from pato_app.calculo import medidas, resumo
from pato_app.importador import importar_xlsx
from pato_app.recalculo import soffice
from pato_app.servico import MODELO_PADRAO, Servico
from pato_app.xlsx_patch import XlsxPatcher, shift_formula

AQUI = Path(__file__).parent
BR174 = AQUI / "dados" / "INVENTARIO_BR174.xlsx"


@pytest.fixture(scope="module")
def pato():
    return importar_xlsx(str(MODELO_PADRAO))


@pytest.fixture(scope="module")
def br174():
    return importar_xlsx(str(BR174))


def _valor(res, prefixo):
    return next(x["valor"] for x in res["resumo_invent"] if x["descricao"].startswith(prefixo))


def test_normalizacao():
    assert nz.numero("1.234,56") == 1234.56
    assert nz.numero("85,65") == 85.65
    assert nz.numero("85+650") == 85.65
    assert nz.coordenada("-14.591364°", "lat") == -14.591364
    assert abs(nz.coordenada("14°35'28.9\"S", "lat") + 14.591361) < 1e-5
    assert nz.coordenada(8370697.5, "lat") is None  # UTM rejeitado
    assert nz.lado("Esquerdo") == "E" and nz.lado("LD") == "D" and nz.lado("D/E") == "E/D"
    assert nz.condicao("Satisfatório") == "Satisfatória"
    assert nz.condicao("DEMOLIR E RECOMPOR") == "Demolir e Recompor"


def test_formula_shift():
    f = "SUM(CAD_MFC!$L$21:$L$243)+CAD_MFC!X21+'CAD_MFC'!A300"
    assert shift_formula(f, "CAD_MFC", "PATO", 243, 5) == "SUM(CAD_MFC!$L$21:$L$248)+CAD_MFC!X21+'CAD_MFC'!A305"
    assert shift_formula('IF(A1="L250",B250,0)', "X", "X", 100, 2) == 'IF(A1="L250",B252,0)'
    assert shift_formula("VLOOKUP(B4,[1]onerado!$A:$H,7,FALSE)", "onerado", "BD", 1, 3).endswith("$A:$H,7,FALSE)")


def test_importa_pato_e_reproduz_totais(pato):
    """As quantidades calculadas batem com os valores das abas CAD_*/RESUMO_INVENT/PATO."""
    assert pato["projeto"]["km_ini"] == 85 and pato["projeto"]["km_fim"] == 212.5
    res = resumo(pato["projeto"], pato["itens"])
    assert _valor(res, "Faixa de pista") == pytest.approx(255)
    assert _valor(res, "Acostamento — Concreto") == pytest.approx(24)
    assert _valor(res, "Acostamento — Tratamento") == pytest.approx(231)
    assert _valor(res, "Meio-fio") == pytest.approx(143100)
    assert _valor(res, "Sarjeta") == pytest.approx(1950)
    assert _valor(res, "Descidas") == pytest.approx(22158)
    assert _valor(res, "Volume de bueiro") == pytest.approx(6910.2004, rel=1e-6)
    assert _valor(res, "Capina") == pytest.approx(47970)
    assert _valor(res, "Caiação") == pytest.approx(39594.9)
    assert _valor(res, "Cerca") == pytest.approx(166.59)
    assert _valor(res, "Roçada com roçadeira") == pytest.approx(4.564)
    assert _valor(res, "Roçada mecanizada") == pytest.approx(68.922)
    assert _valor(res, "Defensas") == pytest.approx(12640)
    pato_q = {x["item"]: x["valor"] for x in res["pato"]}
    esperado = {"01.02.05.07": 272550, "01.02.06.02": 3.6435, "01.02.06.03": 3619.725, "01.02.06.06": 87.55,
                "01.02.06.07": 11.2, "01.02.08.01": 703368.75, "01.02.08.02": 46.89125, "01.02.08.03": 72356.25,
                "01.02.08.04": 48237.5}
    for k, v in esperado.items():
        assert pato_q[k] == pytest.approx(v), k


def test_importa_layout_diferente(br174):
    """Inventário com colunas em outras posições (BR-174) é lido pelo nome do cabeçalho."""
    por = {}
    for it in br174["itens"]:
        por[it["categoria"]] = por.get(it["categoria"], 0) + 1
    assert por["SEG"] == 5 and por["PISTA"] == 5 and por["MFC"] == 61 and por["OAC"] == 79
    assert por["ROCADA"] == 96 and por["PLACA"] == 651 and por["MARCA"] == 669
    assert br174["projeto"]["rodovia"] == "BR-174/MT"
    roc = next(i for i in br174["itens"] if i["categoria"] == "ROCADA")
    assert roc["km_ini"] == 366.63 and roc["km_fim"] == 367.07  # usa o 2º par de km
    cerca = next(i for i in br174["itens"] if i["categoria"] == "CERCA")
    assert cerca["lat_ini"] == pytest.approx(-14.591364)  # '-14.591364°'
    marca = next(i for i in br174["itens"] if i["categoria"] == "MARCA")
    assert -20 < marca["lat_ini"] < -10  # ignora as colunas UTM
    # medidas conferem com as colunas calculadas da própria planilha
    for it in br174["itens"]:
        if it["categoria"] == "OAC" and it["dados"].get("volume"):
            assert medidas(it)["volume"] == pytest.approx(it["dados"]["volume"], rel=1e-6)


def test_patcher_preserva_e_insere(tmp_path):
    p = XlsxPatcher(str(MODELO_PADRAO))
    p.insert_rows("CAD_MFC", 243, 10)
    assert p.get("CAD_MFC", "L260") == "=SUM(L21:L253)"
    assert p.get("RESUMO_INVENT", "P16") == "=CAD_MFC!L260"
    p.set("CAPA", "D74", "BR-999/MT")
    out = tmp_path / "x.xlsx"
    p.save(str(out))
    q = XlsxPatcher(str(out))
    assert q.get("CAPA", "D74") == "BR-999/MT"
    assert "xl/calcChain.xml" not in q._raw
    # desenhos preservados
    assert sum(1 for n in q._raw if n.startswith("xl/drawings/")) == sum(1 for n in p._raw if n.startswith("xl/drawings/"))


@pytest.fixture()
def sv(tmp_path):
    return Servico(pasta=tmp_path)


def test_fluxo_servico(sv, tmp_path):
    r = sv.importar_arquivo(str(BR174))
    pid = r["projeto"]["id"]
    assert r["importados"] == 2500
    novo = sv.adicionar_itens(pid, [{"km_ini": 400, "km_fim": 400.5, "lado": "Esquerdo", "condicao": "recompor",
                                     "tipo": "MFC03", "largura": "0,25"}], "MFC")[0]
    assert novo["lado"] == "E" and novo["dados"]["largura"] == 0.25
    it = sv.atualizar_item(novo["id"], {"km_fim": 401})
    assert it["medidas"]["ext_m"] == pytest.approx(1000)
    out = sv.alterar_trecho(pid, 400, 450, recortar=True)
    assert out["removidos"] > 0
    itens = sv.itens(pid)
    assert all(400 - 1e-6 <= min(i["km_ini"], i.get("km_fim") or i["km_ini"]) and
               max(i["km_ini"], i.get("km_fim") or i["km_ini"]) <= 450 + 1e-6 for i in itens if i["km_ini"] is not None)
    g = sv.gerar_aceiro(pid, 30, 0.2)
    assert g["gerados"] == len(sv.itens(pid, "CERCA"))
    ex = sv.exportar_pato(pid)
    assert os.path.exists(ex["arquivo"])
    wb = openpyxl.load_workbook(ex["arquivo"], read_only=True)
    assert wb["CAD_SEG"]["D23"].value == 400
    assert wb["CAPA"]["D78"].value == "=450.0-400.0"
    wb.close()
    inv = sv.exportar_inventario(pid)
    assert os.path.exists(inv["arquivo"])
    assert sv.geojson(pid)["features"]
    assert sv.kml(pid).startswith("<?xml")


@pytest.mark.skipif(not soffice(), reason="LibreOffice não instalado")
def test_roundtrip_pato_total_inalterado(tmp_path, pato):
    """Reexportar o inventário do próprio PATO e recalcular mantém o total do orçamento."""
    from pato_app.exportador import preencher_modelo_pato
    from pato_app.recalculo import recalcular
    out, rec = tmp_path / "rt.xlsx", tmp_path / "rt_r.xlsx"
    preencher_modelo_pato(pato["projeto"], pato["itens"], str(MODELO_PADRAO), str(out))
    recalcular(str(out), str(rec))
    wb = openpyxl.load_workbook(rec, data_only=True, read_only=True)
    assert wb["RESUMO_PATO"]["D26"].value == pytest.approx(26682348.76, abs=0.01)
    assert wb["RESUMO_PATO"]["F26"].value == pytest.approx(27994808.64, abs=0.01)
    wb.close()


# ------------------------------------------------- nomenclatura (Álbum DNIT) e correções
from pato_app import correcao, padrao_dnit  # noqa: E402


def _pad(cat, tipo, **dados):
    it = {"categoria": cat, "tipo": tipo, "dados": dict(dados), "origem": "X!1"}
    return it, padrao_dnit.padronizar_item(it)


def test_padronizacao_album():
    it, log = _pad("MFC", "MFC03")
    assert it["tipo"] == "MFC 03" and log[0]["situacao"] == "padronizado"
    it, log = _pad("EDA", "EDA-03B")
    assert it["tipo"] == "EDA 03 B"
    it, log = _pad("EDA", "EDA01")
    assert it["tipo"] == "EDA01" and log[0]["situacao"] == "fora do padrão — verificar"
    it, log = _pad("STC", "STC-125/25", largura=1.25)
    assert it["tipo"] == "STC 125-25"
    it, log = _pad("STC", "STC 01", largura=0.73)
    assert it["tipo"] == "STC 73-15" and log[0]["situacao"] == "padronizado — conferir"
    it, log = _pad("DAR", "DAR01")
    assert log[0]["situacao"] == "fora do padrão — verificar" and "DAR 40-20" in log[0]["regra"]
    it, log = _pad("OAC", "Bueiro duplo tubular de concreto", dimensao="1,00")
    assert it["tipo"] == "BDTC" and it["dados"]["dimensao"] == 1.0 and it["dados"]["linhas"] == 2
    it, log = _pad("OAC", "BSTC Ø 0,40")
    assert it["tipo"] == "BSTC" and any("não consta no álbum" in e["regra"] for e in log)
    it, log = _pad("OAC", "4xBSCC", dimensao=1.5, linhas=4)
    assert it["tipo"] == "BSCC" and it["dados"]["linhas"] == 4


def test_correcoes_br158(pato):
    itens = [dict(i, id=n) for n, i in enumerate(pato["itens"])]
    an = correcao.analisar(pato["projeto"], itens)
    fix = {(a["origem"], a["campo"]): a["para"] for a in an["acoes"] if a["situacao"] == "corrigido"}
    assert fix[("CAD_OAC!80", "lat_ini")] == -11.0662          # sinal
    assert fix[("CAD_ACEIRO!6", "lat_fim")] == -10.356          # dígito
    assert fix[("CAD_ACEIRO!27", "lon_ini")] == -51.53637       # dígitos invertidos
    assert fix[("CAD_ALAMB!21", "snv")] == "158BMT0182"        # falta 'B'
    assert an["resumo"]["verificar em campo"]["itens"] == 0


def test_correcoes_br174_e_snv(br174):
    itens = [dict(i, id=n) for n, i in enumerate(br174["itens"])]
    an = correcao.analisar(br174["projeto"], itens)
    assert an["snv"]["concordancia"] < 0.1  # cadastro de segmentos em outra numeração
    sit = {(a["origem"], a["campo"]): a for a in an["acoes"]}
    assert sit[("CADASTRO TACHA, TACHÕES!26", "lat_ini")]["situacao"] == "corrigido"
    assert sit[("CADASTRO TACHA, TACHÕES!403", "item")]["situacao"] == "removido"
    assert sit[("CADASTRO TACHA, TACHÕES!389", "km_fim")]["para"] == 503.46
    assert sit[("CADASTRO DE ÁREA ROÇADA!73", "snv")]["para"] == "174BMT0110 / 174BMT0112"
    # regra do km informado: divergência de 1–2 km também é reposicionada (não fica para verificar)
    assert sit[("CADASTRO DE SARJETA!58", "lat_ini")]["situacao"] == "corrigido"
    assert not any(a["situacao"] == "verificar em campo" and a["campo"].startswith("lat") for a in an["acoes"])
    # tolerância menor → mais pontos reposicionados
    an2 = correcao.analisar(br174["projeto"], itens, tol_km=0.3)
    assert sum(1 for a in an2["acoes"] if a["campo"].startswith("lat")) > sum(1 for a in an["acoes"] if a["campo"].startswith("lat"))


def test_snv_oficial_csv(tmp_path, br174):
    csv_ = tmp_path / "snv.csv"
    csv_.write_text("BR;UF;Código;Local de Início;Local de Fim;km inicial;km final\n"
                    "174;MT;174BMT0110;A;B;366,5;394,0\n174;MT;174BMT0112;B;C;394,0;402,3\n"
                    "174;MT;174BMT0115;C;D;402,3;440,8\n174;MT;174BMT0120;D;E;440,8;500,4\n"
                    "174;MT;174BMT0125;E;F;500,4;502,5\n174;MT;174BMT0130;F;G;502,5;510,0\n"
                    "070;MT;070BMT0010;X;Y;0;10\n", encoding="utf-8")
    tab = correcao.ler_snv_oficial(str(csv_), "BR-174/MT")
    assert [t["codigo"] for t in tab][:2] == ["174BMT0110", "174BMT0112"] and len(tab) == 6
    itens = [dict(i, id=n) for n, i in enumerate(br174["itens"])]
    an = correcao.analisar(br174["projeto"], itens, tab, "202604A")
    assert an["snv"]["fonte_referencia"].startswith("SNV oficial")
    assert any("174BMT0100" in o for o in an["snv"]["observacoes"])


def test_revisao_e_ajustes(tmp_path, sv):
    r = sv.importar_arquivo(str(BR174))
    pid = r["projeto"]["id"]
    assert r["ajustes_automaticos"] > 100 and r["sugestoes"] > 0
    # certo → aplicado automaticamente e registrado em Ajustes
    mfc = sv.itens(pid, "MFC")[0]
    assert mfc["tipo"].startswith("MFC 0")
    tacha = next(i for i in sv.itens(pid, "TACHA") if i["origem"] == "CADASTRO TACHA, TACHÕES!26")
    assert tacha["lat_ini"] == pytest.approx(-14.7247)
    aplicadas = sv.sugestoes(pid, "aplicada")
    assert all(s_["automatica"] for s_ in aplicadas)
    assert not any(i["origem"] == "CADASTRO TACHA, TACHÕES!403" for i in sv.itens(pid, "TACHA"))
    # desfazer um ajuste automático volta o valor e a sugestão para a Revisão
    aj = next(s_ for s_ in aplicadas if s_["item_id"] == mfc["id"])
    sv.desfazer_sugestao(aj["id"])
    assert sv.db.item(mfc["id"])["tipo"] == "MFC03" and sv.db.sugestao(aj["id"])["status"] == "pendente"
    # pendentes já vêm com a sugestão mais assertiva preenchida
    pend = sv.sugestoes(pid, "pendente")
    d1 = next(s_ for s_ in pend if s_["atual"].get("tipo") == "DAR01")
    d2 = next(s_ for s_ in pend if s_["atual"].get("tipo") == "DAR02")
    assert d1["mudancas"]["tipo"] == "DAR 40-20" and d2["mudancas"]["tipo"] == "DAR 60-30"
    e1 = next(s_ for s_ in pend if s_["atual"].get("tipo") == "EDA01")
    assert e1["mudancas"]["tipo"] == "EDA 01 B"          # variante mais usual na referência
    bstm = next(s_ for s_ in pend if s_["atual"].get("tipo") == "BSTM")
    assert not bstm["mudancas"]                           # metálico: sem troca automática
    # aplicar sugestão, aplicar com valor próprio, dispensar
    sv.aplicar_sugestao(d1["id"])
    assert sv.db.item(d1["item_id"])["tipo"] == "DAR 40-20"
    sv.aplicar_sugestao(e1["id"], mudancas={"tipo": "EDA 01 A"})
    assert sv.db.item(e1["item_id"])["tipo"] == "EDA 01 A"
    sv.dispensar_sugestao(bstm["id"])
    sv.gerar_sugestoes(pid)
    assert not any(s_["item_id"] == bstm["item_id"] for s_ in sv.sugestoes(pid, "pendente"))
    manuais = [s_ for s_ in sv.sugestoes(pid, "aplicada") if not s_["automatica"]]
    assert {s_["id"] for s_ in manuais} >= {d1["id"], e1["id"]}
    # planilha original com valores (sem marcações)
    out = sv.exportar_original(pid)
    wb = openpyxl.load_workbook(out["arquivo"])
    c = wb["CADASTRO TACHA, TACHÕES"]["F26"]
    assert c.value == pytest.approx(-14.7247) and c.fill.fgColor.rgb in (None, "00000000")
    wb.close()
