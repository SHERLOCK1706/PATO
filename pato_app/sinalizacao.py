"""Sinalização de obras (abas DURACAO_SERVICOS, TEMPO_PROJETO_TIPO e UN_DIA do PATO) — Res. DNIT nº 8/2023, art. 44 e Anexo XIII.

1) Duração: horas de cada serviço que ocupa pista/acostamento = TRUNC(quantidade ÷ produção horária da equipe; 3);
   manutenção rotineira (MER) = extensão × h/km. Cada serviço tem o projeto-tipo de sinalização do Anexo XIII.
2) Por projeto-tipo: dias = horas ÷ 7,333 h/dia; horas de operação de bandeirola (5213850) = TRUNC(horas × operadores; 3).
3) Dispositivos (un.dia) = ARRED(quantidade do dispositivo no projeto-tipo × dias; 0), somados por código SICRO.
Calculado em separado para o grupo por desempenho (entra no GD) e para os demais serviços de conservação (itens 01.06).
"""
from __future__ import annotations

import math
import re

from . import normalizar as nz

HORAS_DIA = 7.333
TIPOS = {1: dict(rotulo="Fixo no acostamento", operadores=0), 3: dict(rotulo="Fixo na pista com operação PARE e SIGA", operadores=2),
         16: dict(rotulo="Móvel no acostamento", operadores=1), 17: dict(rotulo="Móvel na pista com operação PARE e SIGA", operadores=2)}
# dispositivos por projeto-tipo (planilha BR-158/MT, Anexo XIII) e código SICRO em que são pagos
DISPOSITIVOS = [
    dict(chave="R-1", descricao="Parada obrigatória (PARE e SIGA)", codigo="5212558", qtd={1: 0, 3: 2, 16: 0, 17: 2}),
    dict(chave="R-7", descricao="Proibido ultrapassar", codigo="5212557", qtd={1: 0, 3: 2, 16: 0, 17: 0}),
    dict(chave="R-19", descricao="Velocidade máxima permitida", codigo="5212557", qtd={1: 2, 3: 2, 16: 0, 17: 2}),
    dict(chave="A-24", descricao="Obras ou serviço", codigo="5212560", qtd={1: 1, 3: 0, 16: 0, 17: 0}),
    dict(chave="A15 + DIST", descricao="Parada obrigatória à frente com indicação de distância", codigo="5212560", qtd={1: 0, 3: 2, 16: 0, 17: 2}),
    dict(chave="A24 + DIST", descricao="Obras ou serviço com indicação de distância", codigo="5212560", qtd={1: 2, 3: 2, 16: 2, 17: 2}),
    dict(chave="DESVIO", descricao="Desvio (1,25 × 0,40 m)", codigo="5212560", qtd={1: 0, 3: 1, 16: 0, 17: 0}),
    dict(chave="FIM DAS OBRAS", descricao="Fim das obras (1,25 × 0,80 m)", codigo="5212556", qtd={1: 2, 3: 2, 16: 0, 17: 0}),
    dict(chave="SEM ACOST", descricao="Trecho sem acostamento (1,75 × 1,00 m)", codigo="5212556", qtd={1: 1, 3: 0, 16: 1, 17: 0}),
    dict(chave="CONE", descricao="Cone plástico para canalização de trânsito", codigo="5213835", qtd={1: 50, 3: 70, 16: 50, 17: 40}),
    dict(chave="TAMBOR", descricao="Cilindro canalizador de tráfego com base quadrada", codigo="5213838", qtd={1: 5, 3: 8, 16: 0, 17: 0}),
    dict(chave="CAVAL. ART.", descricao="Cavalete articulado", codigo="5213383", qtd={1: 1, 3: 2, 16: 0, 17: 2}),
]
CODIGO_BANDEIROLA = "5213850"
GRUPOS = {"desempenho": "Grupo por desempenho", "conservacao": "Serviços de conservação rodoviária"}


def trunc3(v: float) -> float:
    return math.trunc(float(f"{v * 1000:.12g}")) / 1000


def arred0(v: float) -> int:
    return int(math.floor(abs(v) + 0.5)) * (1 if v >= 0 else -1)


def padrao() -> dict:
    return dict(horas_dia=HORAS_DIA, tipos={str(k): dict(v) for k, v in TIPOS.items()},
                dispositivos=[dict(d, qtd={str(k): v for k, v in d["qtd"].items()}) for d in DISPOSITIVOS], servicos=[])


def calcular(dados: dict, quantidade_fn, producao_fn, extensao: float) -> dict:
    """quantidade_fn(servico) -> quantidade do serviço no PATO/desempenho; producao_fn(codigo) -> produção horária."""
    d = dados or padrao()
    hd = nz.numero(d.get("horas_dia")) or HORAS_DIA
    tipos = {int(k): v for k, v in (d.get("tipos") or {str(k): v for k, v in TIPOS.items()}).items()}
    servs, horas = [], {}
    for s in d.get("servicos") or []:
        g, t = s.get("grupo") or "conservacao", int(nz.numero(s.get("tipo")) or 3)
        if s.get("codigo") == "MER" or s.get("h_por_km") not in (None, ""):
            q = extensao
            prod = None
            h = (nz.numero(s.get("h_por_km")) or 0.0) * extensao
        else:
            q = quantidade_fn(s) or 0.0
            prod = producao_fn(s["codigo"]) if producao_fn else None
            prod = prod or nz.numero(s.get("producao_ref"))
            h = trunc3(q / prod) if prod else 0.0
        horas[(g, t)] = horas.get((g, t), 0.0) + h
        servs.append(dict(s, grupo=g, tipo=t, quantidade=q, producao=prod, horas=h))
    por_grupo = {}
    for g in GRUPOS:
        linhas_t, tot_op = [], 0.0
        for t, info in sorted(tipos.items()):
            h = round(horas.get((g, t), 0.0), 6)
            dias = h / hd
            op = trunc3(h * (nz.numero(info.get("operadores")) or 0))
            tot_op += op
            linhas_t.append(dict(tipo=t, rotulo=info.get("rotulo"), horas=h, dias=dias, operadores=info.get("operadores"), horas_operacao=op))
        disp, por_cod = [], {}
        for dv in d.get("dispositivos") or []:
            q = sum(arred0((nz.numero((dv.get("qtd") or {}).get(str(lt["tipo"]))) or 0) * lt["dias"]) for lt in linhas_t)
            disp.append(dict(dv, un_dia=q))
            if dv.get("codigo"):
                por_cod[dv["codigo"]] = por_cod.get(dv["codigo"], 0) + q
        por_cod[CODIGO_BANDEIROLA] = round(tot_op, 3)
        por_grupo[g] = dict(tipos=linhas_t, dispositivos=disp, codigos=por_cod)
    return dict(servicos=servs, grupos=por_grupo, horas_dia=hd)


# ----------------------------------------------------------- importação
def importar_planilha(caminho: str) -> dict | None:
    """Lê DURACAO_SERVICOS (serviços, projeto-tipo, produção), TEMPO_PROJETO_TIPO (operadores, h/dia) e UN_DIA (dispositivos)."""
    import openpyxl
    wv = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    wf = openpyxl.load_workbook(caminho, read_only=True)
    try:
        if not {"DURACAO_SERVICOS", "UN_DIA"} <= set(wv.sheetnames):
            return None
        d = padrao()
        vr = [tuple(r) + (None,) * 12 for r in wv["DURACAO_SERVICOS"].iter_rows(values_only=True)]
        fr = [tuple(r) + (None,) * 12 for r in wf["DURACAO_SERVICOS"].iter_rows(values_only=True)]
        grupo = "desempenho"
        for v, f in zip(vr, fr):
            rot = nz.sem_acento(str(v[1] or "")).upper()
            if "SERVICOS DE CONSERVACAO" in rot:
                grupo = "conservacao"
                continue
            if "DEMAIS SERVICOS" in rot:
                grupo = "dsm"
                continue
            cod = nz.texto(v[1])
            if not cod or grupo == "dsm" or nz.numero(v[8]) is None:
                continue
            s = dict(codigo=cod.strip().upper(), descricao=nz.texto(v[2]), unidade=nz.texto(v[4]), grupo=grupo,
                     tipo=int(nz.numero(v[8])), producao_ref=nz.numero(v[5]), quantidade_ref=nz.numero(v[3]))
            fq = str(f[3] or "")
            m = re.search(r"PATO!\$?[A-Z]+\$?(\d+):\$?[A-Z]+\$?(\d+)", fq)
            if m:
                s["fonte"] = "pato"
                s["linhas_pato"] = [int(m.group(1)), int(m.group(2))]
            elif "DESEMPENHO" in fq.upper():
                s["fonte"] = "desempenho"
            if s["codigo"] == "MER" or (isinstance(f[6], str) and "F" in f[6] and "D" in f[6] and "/" not in f[6]):
                s["h_por_km"] = nz.numero(v[5])
            d["servicos"].append(s)
        if "TEMPO_PROJETO_TIPO" in wv.sheetnames:
            ft = [tuple(r) + (None,) * 8 for r in wf["TEMPO_PROJETO_TIPO"].iter_rows(values_only=True)]
            vt = [tuple(r) + (None,) * 8 for r in wv["TEMPO_PROJETO_TIPO"].iter_rows(values_only=True)]
            for v, f in zip(vt, ft):
                t = nz.numero(v[1])
                if t is not None and int(t) in TIPOS and nz.numero(v[5]) is not None:
                    d["tipos"][str(int(t))]["operadores"] = nz.numero(v[5])
                m = re.search(r"/\s*([\d.]+)\s*$", str(f[4] or ""))
                if m:
                    d["horas_dia"] = float(m.group(1))
        # tabela de dispositivos por projeto-tipo (linhas após "DEMAIS SERVIÇOS" na aba UN_DIA) e o código SICRO de cada um
        ur = [tuple(r) + (None,) * 12 for r in wv["UN_DIA"].iter_rows(values_only=True)]
        uf = [tuple(r) + (None,) * 12 for r in wf["UN_DIA"].iter_rows(values_only=True)]
        cab = next((r for r in ur if all(isinstance(r[c], (int, float)) for c in (3, 4, 5, 6))), None)
        tipos_cols = [int(cab[c]) for c in (3, 4, 5, 6)] if cab else [1, 3, 16, 17]
        codigo_de = {}
        for v, f in zip(ur, uf):  # L32 = H33 + H34 → os dispositivos das linhas 33 e 34 são pagos no código K32
            if v[10] and isinstance(f[11], str) and "H" in f[11].upper():
                for n in re.findall(r"H(\d+)", f[11].upper()):
                    codigo_de[int(n)] = str(v[10]).strip()
        linha_disp = {}
        for i, v in enumerate(ur, start=1):
            if nz.texto(v[1]) and codigo_de.get(i):
                linha_disp[nz.texto(v[1]).strip().upper()] = codigo_de[i]
        tabela = []
        for v in ur[45:]:
            ch = nz.texto(v[1])
            if ch and all(nz.numero(v[c]) is not None for c in (3, 4, 5, 6)):
                tabela.append(dict(chave=ch.strip().upper(), descricao=nz.texto(v[2]),
                                   codigo=linha_disp.get(ch.strip().upper()),
                                   qtd={str(t): nz.numero(v[c]) for t, c in zip(tipos_cols, (3, 4, 5, 6))}))
        if tabela:
            d["dispositivos"] = tabela
        return d
    finally:
        wv.close()
        wf.close()
