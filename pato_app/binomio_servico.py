"""Binômio no sistema: índices de reajuste, materiais cotados, betuminosos (ANP) e anexos (cotações, ICMS)."""
from __future__ import annotations

import re
import shutil
import threading
import uuid
import json
from pathlib import Path

from . import binomio as bn
from . import composicoes as cp
from . import normalizar as nz

TIPO_MATERIAL = "cotacao"
TIPO_BETUMINOSO = "betuminoso"
TIPOS_ANEXO = {"cotacao": "Cotação", "icms": "ICMS (legislação/alíquota)", "anp": "Tabela ANP", "outro": "Outro"}


_TRAVA = threading.Lock()


def _num(v):
    return nz.numero(v) if v not in (None, "") else None


class Binomios:
    def __init__(self, sv):
        self.sv = sv
        self.db = sv.db
        self.pasta_anexos = Path(sv.pasta) / "anexos"

    # ------------------------------------------------------------ índices
    def garantir_indices(self) -> None:
        """Índices de reajustamento: os da planilha PATO de exemplo e a última publicação do DNIT incluída no sistema
        (modelos/indices_reajuste.json); meses que ainda não estão no banco são acrescentados."""
        if getattr(self, "_indices_ok", False):
            return
        with _TRAVA:
            if not self.db.indices():
                from .servico import MODELO_PADRAO
                try:
                    self._gravar_series(bn.ler_indices(str(MODELO_PADRAO)), "Planilha PATO de exemplo (INDICES_REAJ)")
                except Exception:  # noqa: BLE001 — sem índices o binômio funciona sem reajuste
                    pass
            arq = Path(__file__).parent / "modelos" / "indices_reajuste.json"
            if arq.exists():
                inc = json.loads(arq.read_text(encoding="utf-8"))
                tem = self.db.indices()
                novos = [(n, m, v) for n, ser in inc["series"].items() for m, v in ser.items() if m not in (tem.get(n) or {})]
                if novos:
                    self.db.gravar_indices(novos, inc.get("fonte"))
            self._indices_ok = True

    def _gravar_series(self, series: dict, fonte: str) -> int:
        return self.db.gravar_indices([(n, m, v) for n, s in series.items() for m, v in s.items()], fonte)

    def indices(self) -> dict:
        self.garantir_indices()
        series = self.db.indices()
        meses = sorted({m for s in series.values() for m in s})
        return dict(nomes=list(series), meses=meses, series=series,
                    padrao=dict(petreos=bn.INDICE_PETREOS, asfalto=bn.INDICE_ASFALTO))

    def importar_indices(self, caminho: str) -> dict:
        series = bn.ler_indices_pdf(caminho) if caminho.lower().endswith(".pdf") else bn.ler_indices(caminho)
        if not series:
            raise ValueError("Nenhum índice encontrado na planilha.")
        n = self._gravar_series(series, Path(caminho).name)
        return dict(gravados=n, series=len(series), **self.indices())

    def gravar_indice(self, nome: str, mes: str, valor) -> dict:
        m = bn.mes(mes)
        if not nome or not m or not re.fullmatch(r"\d{4}-\d{2}", m):
            raise ValueError("Informe o índice e o mês (AAAA-MM).")
        self.db.gravar_indices([(nome.strip(), m, _num(valor))], "manual")
        return self.indices()

    # -------------------------------------------------------- referências
    def _ref(self, pid: int) -> dict:
        from .custos import calcular_bdi
        cfg = self.sv.custos.config(pid)
        bdi = cfg["bdi"]["SD"]
        exato = calcular_bdi(cfg["parcelas"]["SD"])
        if abs(exato - bdi) < 0.00005:  # a planilha BINOMIO_ANP usa o BDI sem arredondar (BDI!E25)
            bdi = exato
        return dict(data_base=cfg["data_base"], bdi=bdi, regime=cfg["regime"], cfg=cfg)

    def _tabela_mat(self, pid: int) -> dict:
        try:
            base = self.sv.custos._base_projeto(pid, self.sv.custos.config(pid)["regime"])
        except ValueError:
            return {}
        return (base.get("tabelas") or {}).get("MAT") or {}

    def preco_sicro(self, pid: int, codigo: str | None, tab: dict | None = None) -> dict | None:
        if not codigo:
            return None
        tab = self._tabela_mat(pid) if tab is None else tab
        d = tab.get(cp._cod(codigo))
        if not d or d.get("preco") is None:
            return None
        return dict(preco=d["preco"], unidade=d.get("unidade"), descricao=d.get("descricao"))

    # --------------------------------------------------------- listagem
    def listar(self, pid: int) -> dict:
        ref = self._ref(pid)
        ind = self.db.indices() or (self.garantir_indices() or self.db.indices())
        tab = self._tabela_mat(pid)
        materiais, produtos = [], []
        anexos = self.db.anexos(pid)
        for b in self.db.binomios(pid):
            d = b["dados"]
            ax = [a for a in anexos if a["binomio_id"] == b["id"]]
            if b["tipo"] == TIPO_MATERIAL:
                ps = self.preco_sicro(pid, d.get("codigo"), tab)
                r = bn.calcular_material(d, ind, ref["data_base"], ps)
                r["sicro"] = ps
                materiais.append(dict(r, id=b["id"], anexos=ax))
            else:
                bdi = _num(d.get("bdi"))
                r = bn.calcular_betuminoso(d, ind, ref["data_base"], bdi if bdi is not None else ref["bdi"])
                r["sicro"] = self.preco_sicro(pid, d.get("codigo"), tab)
                produtos.append(dict(r, id=b["id"], anexos=ax, bdi_projeto=bdi is None))
        return dict(data_base=ref["data_base"], bdi=ref["bdi"], materiais=materiais, produtos=produtos,
                    resumo=self._resumo(materiais, produtos),
                    grupos=bn.GRUPOS, tipos_anexo=TIPOS_ANEXO, indices=list(ind),
                    anexos_gerais=[a for a in anexos if not a["binomio_id"]])

    @staticmethod
    def _assinatura(r: dict) -> dict | None:
        """O que foi conferido: vencedor e total (se mudar a cotação, a conferência deixa de valer)."""
        v = r.get("vencedor")
        if not v:
            return None
        return dict(fornecedor=v.get("fornecedor") or f"{v.get('local')}/{v.get('uf')}", total=round(v.get("total") or 0, 2))

    def _resumo(self, materiais: list[dict], produtos: list[dict]) -> list[dict]:
        """RESUMO das cotações: o mais vantajoso de cada material, o preço do SICRO, a conferência e se já foi adotado."""
        out = []
        for r in materiais + produtos:
            v = r.get("vencedor") or {}
            ass = self._assinatura(r)
            conf = r.get("conferido") or None
            sic = next((x for x in r.get("linhas") or [] if x.get("sicro")), None)
            n_cot = len([x for x in r.get("linhas") or [] if not x.get("sicro")]) if "nome" in r else len(r.get("linhas") or [])
            ap = r.get("aplicado") or {}
            avisos = []
            if "nome" in r and n_cot < 3:
                avisos.append(f"{n_cot} cotação(ões): a IN DNIT 10/2024 pede no mínimo 3 (ou justificativa).")
            if "nome" in r and n_cot and not r.get("anexos"):
                avisos.append("Sem o PDF da cotação anexado.")
            if not r.get("codigo"):
                avisos.append("Sem código do insumo SICRO: informe para adotar o preço na composição.")
            if "nome" in r and v and not v.get("sicro") and v.get("unidade") == "m3" and not (
                    nz.numero(next((c.get("densidade") for c in r.get("cotacoes") or [] if c.get("fornecedor") == v.get("fornecedor")), None))
                    or nz.numero(r.get("densidade"))):
                avisos.append("Cotação em m³ sem densidade: adotado 1,0 t/m³ — informe a densidade do material.")
            un = "t"
            if "nome" in r:
                un_s = ((r.get("sicro") or {}).get("unidade") or "t").replace("³", "3").lower()
                un = "m³" if un_s.startswith("m") else "t"
                preco = v.get("preco_m3_reaj" if un == "m³" else "preco_t_reaj")
            else:
                preco = v.get("aq_icms")
            out.append(dict(id=r["id"], material=r.get("nome") or r.get("produto"), codigo=r.get("codigo"),
                            tipo="material" if "nome" in r else "betuminoso", cotacoes=n_cot,
                            vencedor=ass["fornecedor"] if ass else None, sicro_vence=bool(v.get("sicro")),
                            total=v.get("total"), total_sicro=sic.get("total") if sic else None,
                            preco_adotar=preco, unidade=un, dmt_pav=v.get("dmt_pav"), dmt_npav=v.get("dmt_npav"),
                            conferido=bool(conf and ass and conf.get("vencedor") == ass), conferido_por=(conf or {}).get("por"),
                            conferencia_vencida=bool(conf and ass and conf.get("vencedor") != ass),
                            adotado=bool(ap and ass and ap.get("fornecedor") in (ass["fornecedor"], v.get("local"))
                                         and abs((ap.get("preco") or 0) - (preco or 0)) < 0.005),
                            avisos=avisos))
        return out

    def conferir(self, pid: int, bid: int, conferido: bool = True, por: str | None = None) -> dict:
        """Marca a conferência do resumo (vencedor e total no momento da conferência)."""
        r = self.item(pid, bid)
        b = self.db.binomio(bid)
        b["dados"]["conferido"] = dict(vencedor=self._assinatura(r), por=nz.texto(por)) if conferido else None
        self.db.salvar_binomio(pid, b["tipo"], b["dados"], bid)
        self.db.registrar(pid, "binômio: conferência", f"{r.get('nome') or r.get('produto')}: {'conferido' if conferido else 'desmarcado'}")
        return self.item(pid, bid)

    def aplicar_conferidos(self, pid: int) -> dict:
        """Adota nas composições o preço dos materiais conferidos e ainda não adotados."""
        msgs, n = [], 0
        for ln in self.listar(pid)["resumo"]:
            if ln["conferido"] and not ln["adotado"]:
                if not ln["codigo"]:
                    msgs.append(f"{ln['material']}: sem código SICRO — informe para adotar o preço.")
                    continue
                r = self.aplicar_vencedor(pid, ln["id"])
                msgs += r["mensagens"]
                n += 1
        if not n and not msgs:
            msgs.append("Nenhum material conferido pendente de adoção.")
        return dict(aplicados=n, mensagens=msgs, resumo=self.listar(pid)["resumo"])

    # ------------------------------------------------ cadastro de fornecedores
    @staticmethod
    def _endereco(end: str | None) -> dict:
        """Rodovia/km, cidade e UF do endereço ("Rodovia BR-174, km 224 - Pontes e Lacerda - MT")."""
        u = nz.sem_acento(end or "").upper()
        out = {}
        m = re.search(r"\bBR\s*[-–]?\s*(\d{3})\D{0,12}?\bKM\s*(\d+(?:[.,]\d+)?)", u)
        if m:
            out.update(br=m.group(1), km=nz.numero(m.group(2)))
        m = re.search(r"[-–,]\s*([A-Za-zÀ-ú' .]{3,40}?)\s*[-–/]\s*([A-Z]{2})\s*$", (end or "").strip())
        if m:
            out.update(cidade=m.group(1).strip(), uf=m.group(2))
        return out

    def fornecedores(self, pid: int) -> dict:
        """Cadastro de fornecedores do projeto, com a posição: ponto do KMZ (croqui) ou rodovia/km do endereço (SNV),
        o km mais próximo na rodovia do projeto e a distância em linha reta até o canteiro."""
        from . import snv
        from .croqui_auto import chave
        lst = (self.db.dados_projeto(pid, "fornecedores") or {}).get("lista") or []
        croqui = self.sv.transporte.dados_croqui(pid)
        kmz = {chave(o.get("nome")): o for o in croqui.get("ocorrencias") or [] if o.get("lat") is not None}
        br_p, uf_p = snv.br_uf(self.sv.projeto(pid).get("rodovia"))
        lc = self.sv.orcamento.locais_croqui(pid) if hasattr(self.sv, "orcamento") else {}
        km_can = (lc or {}).get("km_canteiro")
        pos_can = snv.posicao(br_p, uf_p, km_can) if br_p and km_can is not None else None
        out = []
        for f in lst:
            f = dict(f)
            k = chave(f.get("nome"))
            pk = kmz.get(k) or next((o for n, o in kmz.items() if k and (k in n or n in k)), None)
            if pk:
                f.update(lat=pk["lat"], lon=pk["lon"], fonte_posicao="KMZ")
            elif f.get("br") and f.get("km") is not None:
                pos = snv.posicao(f["br"], f.get("uf") or uf_p, f["km"])
                if pos:
                    f.update(lat=pos[0], lon=pos[1], fonte_posicao=f"endereço BR-{f['br']} km {f['km']:g} (SNV)")
            if f.get("lat") is not None:
                via = snv.localizar(None, f.get("uf") or uf_p, f["lat"], f["lon"])
                if f.get("br") and via and via["br"] != str(f["br"]).zfill(3):  # trecho coincidente: a BR do endereço
                    mesma = snv.localizar(f["br"], f.get("uf") or uf_p, f["lat"], f["lon"])
                    if mesma and mesma["distancia_km"] <= via["distancia_km"] + 0.5:
                        via = mesma
                if via:
                    f["rodovia_proxima"] = f"BR-{via['br']} km {via['km']:g} ({via['distancia_km']:g} km do eixo)"
                if br_p:
                    no_proj = snv.localizar(br_p, uf_p, f["lat"], f["lon"])
                    if no_proj:
                        f.update(km_projeto=no_proj["km"], distancia_eixo_km=no_proj["distancia_km"])
                if pos_can:
                    f["distancia_canteiro_km"] = snv.distancia((f["lat"], f["lon"]), pos_can)
            out.append(f)
        return dict(lista=out, km_canteiro=km_can,
                    nota="Distância ao canteiro em linha reta (estimativa): confira a DMT pela rota e separe pavimentada / não pavimentada.")

    def salvar_fornecedores(self, pid: int, lista: list[dict], acrescentar: bool = False) -> dict:
        from .croqui_auto import chave
        atual = (self.db.dados_projeto(pid, "fornecedores") or {}).get("lista") or [] if acrescentar else []
        por_nome = {chave(f.get("nome")): f for f in atual}
        for f in lista:
            nome = nz.texto(f.get("nome"))
            if not nome:
                continue
            e = self._endereco(f.get("endereco"))
            novo = dict(nome=nome, endereco=nz.texto(f.get("endereco")), insumo=nz.texto(f.get("insumo")),
                        cnpj=nz.texto(f.get("cnpj")), contato=nz.texto(f.get("contato")),
                        br=nz.texto(f.get("br")) or e.get("br"), km=_num(f.get("km")) if _num(f.get("km")) is not None else e.get("km"),
                        cidade=nz.texto(f.get("cidade")) or e.get("cidade"), uf=nz.texto(f.get("uf")) or e.get("uf"))
            por_nome[chave(nome)] = dict(por_nome.get(chave(nome)) or {}, **{k: v for k, v in novo.items() if v is not None})
        self.db.gravar_dados_projeto(pid, "fornecedores", dict(lista=list(por_nome.values())))
        self.db.registrar(pid, "binômio: fornecedores", f"{len(por_nome)} fornecedores no cadastro")
        return self.fornecedores(pid)

    def importar_fornecedores(self, pid: int, caminho: str) -> dict:
        """Planilha de fornecedores (colunas FORNECEDOR, ENDEREÇO, INSUMO; opcionais CNPJ, CONTATO)."""
        import openpyxl
        wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
        lista = []
        for ws in wb.worksheets:
            cols = None
            for row in ws.iter_rows(values_only=True):
                cel = [nz.sem_acento(str(v or "")).upper().strip() for v in row]
                if cols is None:
                    if any(c.startswith("FORNECEDOR") or c in ("EMPRESA", "RAZAO SOCIAL") for c in cel):
                        busca = lambda *ns: next((i for i, c in enumerate(cel) if any(c.startswith(n) for n in ns)), None)
                        cols = dict(nome=busca("FORNECEDOR", "EMPRESA", "RAZAO"), endereco=busca("ENDERECO", "LOCAL"),
                                    insumo=busca("INSUMO", "MATERIA", "PRODUTO"), cnpj=busca("CNPJ"), contato=busca("CONTATO", "TELEFONE", "E-MAIL"))
                    continue
                f = {k: row[i] for k, i in cols.items() if i is not None and i < len(row)}
                if f.get("nome"):
                    lista.append({k: str(v).strip() if v is not None else None for k, v in f.items()})
        if not lista:
            raise ValueError("Não encontrei a coluna FORNECEDOR na planilha.")
        return self.salvar_fornecedores(pid, lista, acrescentar=True)

    def _posicao_fornecedor(self, pid: int, nome: str | None, rodovia: dict | None = None) -> dict | None:
        from .croqui_auto import chave
        k = chave(nome)
        cad = self.fornecedores(pid)["lista"]
        f = next((f for f in cad if k and (chave(f["nome"]) == k or chave(f["nome"]) in k or k in chave(f["nome"]))), None)
        if f is None and rodovia:  # não cadastrado: posição pelo endereço da própria cotação
            from . import snv
            br_p, uf_p = snv.br_uf(self.sv.projeto(pid).get("rodovia"))
            pos = snv.posicao(rodovia["br"], uf_p, rodovia["km"])
            if pos:
                f = dict(nome=nome, lat=pos[0], lon=pos[1], fonte_posicao=f"endereço BR-{rodovia['br']} km {rodovia['km']:g} (SNV)")
                lc = self.sv.orcamento.locais_croqui(pid)
                pc = snv.posicao(br_p, uf_p, lc["km_canteiro"]) if br_p and lc.get("km_canteiro") is not None else None
                if pc:
                    f["distancia_canteiro_km"] = snv.distancia(pos, pc)
        return f

    # --------------------------------------------------- cotação por PDF (OCR)
    def ler_cotacao(self, pid: int, caminho: str, nome_arquivo: str | None = None) -> dict:
        """Lê a cotação (texto do PDF ou OCR) e sugere o material do binômio de cada item — só pré-preenche."""
        from . import cotacao_ocr
        from .croqui_auto import chave
        conhecidos = [f["nome"] for f in (self.db.dados_projeto(pid, "fornecedores") or {}).get("lista") or []]
        conhecidos += [o["nome"] for o in self.sv.transporte.dados_croqui(pid).get("ocorrencias") or [] if o.get("lat") is not None]
        r = cotacao_ocr.ler(caminho, conhecidos, nome_arquivo)
        pf = self._posicao_fornecedor(pid, r.get("fornecedor"), r.get("rodovia"))
        r["posicao"] = pf
        dist = r.get("distancia_km")
        if dist is None and pf and pf.get("distancia_canteiro_km") is not None:
            dist = pf["distancia_canteiro_km"]
            r["avisos"].append(f"DMT estimada em linha reta do fornecedor ({pf.get('fonte_posicao')}) ao canteiro: "
                               f"{dist:g} km — ajuste pela rota real.")
        if pf and (pf.get("distancia_canteiro_km") or 0) > 150:
            r["avisos"].append(f"Fornecedor a {pf['distancia_canteiro_km']:g} km do canteiro (linha reta): confira se a cotação é deste trecho.")
        if dist is None:
            r["avisos"].append("Distância não informada na cotação: localize o fornecedor (KMZ ou cadastro) ou digite a DMT.")
        mats = [dict(id=b["id"], nome=b["dados"].get("nome"), codigo=b["dados"].get("codigo"),
                     densidade=b["dados"].get("densidade")) for b in self.db.binomios(pid, TIPO_MATERIAL)]
        tab = None
        for it in r["itens"]:
            k = chave(it["material"])
            tk = set(k.split())  # todas as palavras do item no nome do material ("AREIA LAVADA" → AREIA MÉDIA LAVADA)
            cand = [m for m in mats if k and (tk <= set(chave(m["nome"]).split()) or chave(m["nome"]) in k)]
            m = min(cand, key=lambda m: abs(len(chave(m["nome"]).split()) - len(tk))) if cand else None
            it["binomio_id"] = m["id"] if m else None
            it["densidade"] = m["densidade"] if m else None
            if not m:  # novo material: sugere o insumo do SICRO pela descrição
                tab = tab if tab is not None else self._tabela_mat(pid)
                it["codigo"] = next((c for c, d in tab.items() if k and k in chave(d.get("descricao"))), None)
            cif = r.get("frete") == "CIF"
            it["dmt_pav"] = 0 if cif else dist
            it["dmt_npav"] = 0
        r["materiais"] = mats
        if r.get("distancia_km") is not None and r.get("frete") != "CIF":
            r["avisos"].append("A distância lida foi posta como DMT pavimentada — confira e separe o trecho não pavimentado.")
        if not r["itens"]:
            r["avisos"].append("Nenhum item com preço reconhecido — preencha manualmente.")
        return r

    def adicionar_cotacao(self, pid: int, caminho: str, nome_arquivo: str, dados: dict) -> dict:
        """Grava os itens conferidos da cotação: acrescenta a cotação ao material (ou cria o material) e anexa o arquivo."""
        feitos = []
        for it in dados.get("itens") or []:
            if not it.get("incluir", True) or it.get("preco") in (None, ""):
                continue
            cot = dict(fornecedor=dados.get("fornecedor"), cnpj=dados.get("cnpj"), data=dados.get("data"),
                       frete=dados.get("frete"), local=dados.get("local"), preco=it.get("preco"),
                       unidade=it.get("unidade"), densidade=it.get("densidade"),
                       dmt_pav=it.get("dmt_pav"), dmt_npav=it.get("dmt_npav"), anexo=nome_arquivo)
            bid = it.get("binomio_id") or None
            if bid:
                d = dict(self.db.binomio(int(bid))["dados"])
                d["cotacoes"] = [c for c in d.get("cotacoes") or []
                                 if nz.texto(c.get("fornecedor")) != nz.texto(cot["fornecedor"])] + [cot]
                r = self.salvar_material(pid, d, int(bid))
            else:
                r = self.salvar_material(pid, dict(nome=it.get("material") or "Material", codigo=it.get("codigo"),
                                                   grupo="petreo", densidade=it.get("densidade"), cotacoes=[cot]))
            self.anexar(pid, caminho, nome_arquivo, r["id"], "cotacao", dados.get("fornecedor"), "cotação lida automaticamente e conferida")
            feitos.append(dict(id=r["id"], material=r["nome"], vencedor=(r.get("vencedor") or {}).get("fornecedor")))
        if not feitos:
            raise ValueError("Nenhum item com preço para gravar.")
        return dict(gravados=feitos, binomio=self.listar(pid))

    def item(self, pid: int, bid: int) -> dict:
        lst = self.listar(pid)
        r = next((x for x in lst["materiais"] + lst["produtos"] if x["id"] == bid), None)
        if not r:
            raise KeyError(f"Binômio {bid} não encontrado.")
        return r

    # ---------------------------------------------------------- gravação
    @staticmethod
    def _limpar_material(d: dict) -> dict:
        grupo = d.get("grupo") if d.get("grupo") in bn.GRUPOS else "petreo"
        cots = []
        for c in d.get("cotacoes") or []:
            if not (c.get("fornecedor") or c.get("preco") not in (None, "") or c.get("sicro")):
                continue
            un = (c.get("unidade") or "t").lower().replace("³", "3")
            cots.append(dict(fornecedor=nz.texto(c.get("fornecedor")) or ("SICRO" if c.get("sicro") else None),
                             sicro=bool(c.get("sicro")), preco=_num(c.get("preco")), unidade="m3" if un.startswith("m") else "t",
                             densidade=_num(c.get("densidade")), data=bn.mes(c.get("data")),
                             dmt_pav=_num(c.get("dmt_pav")), dmt_npav=_num(c.get("dmt_npav")),
                             local=nz.texto(c.get("local")), contato=nz.texto(c.get("contato")),
                             cnpj=nz.texto(c.get("cnpj")), frete=c.get("frete") if c.get("frete") in ("CIF", "FOB") else None,
                             anexo=c.get("anexo")))
        return dict(nome=nz.texto(d.get("nome")) or "Material", codigo=cp._cod(d["codigo"]) if d.get("codigo") else None,
                    grupo=grupo, indice=d.get("indice") or bn.INDICE_PETREOS, densidade=_num(d.get("densidade")),
                    custo_pav=_num(d.get("custo_pav")), custo_npav=_num(d.get("custo_npav")), cotacoes=cots,
                    obs=nz.texto(d.get("obs")), aplicado=d.get("aplicado"), conferido=d.get("conferido"))

    @staticmethod
    def _limpar_produto(d: dict) -> dict:
        origens = []
        for o in d.get("origens") or []:
            if not (o.get("local") or o.get("preco_t") not in (None, "")):
                continue
            icms = _num(o.get("icms"))
            origens.append(dict(nome=nz.texto(o.get("nome")) or f"R-{len(origens) + 1}", uf=nz.texto(o.get("uf")),
                                local=nz.texto(o.get("local")), preco_t=_num(o.get("preco_t")),
                                d_fluvial=_num(o.get("d_fluvial")), d_pav=_num(o.get("d_pav")), d_rp=_num(o.get("d_rp")),
                                d_ln=_num(o.get("d_ln")), icms=icms / 100 if icms and icms > 1 else icms,
                                pedagio=_num(o.get("pedagio")), balsa=_num(o.get("balsa"))))
        out = dict(produto=nz.texto(d.get("produto")) or "Material betuminoso",
                   codigo=cp._cod(d["codigo"]) if d.get("codigo") else None, indice=d.get("indice") or bn.INDICE_ASFALTO,
                   origens=origens, obs=nz.texto(d.get("obs")), aplicado=d.get("aplicado"), conferido=d.get("conferido"),
                   fluvial_sempre=bool(d.get("fluvial_sempre", True)))
        for k, padrao in (("pis", 0.0065), ("cofins", 0.03), ("bdi_transporte", 0.15), ("bdi_transporte_cd", 0.1867),
                          ("ton_viagem", 60)):
            v = _num(d.get(k))
            if v is not None and k != "ton_viagem" and v > 1:
                v /= 100
            out[k] = padrao if v is None else v
        bdi = _num(d.get("bdi"))
        out["bdi"] = None if bdi is None else (bdi / 100 if bdi > 1 else bdi)
        for k in ("indice_atual", "base_rodoviario", "base_fluvial"):
            if _num(d.get(k)) is not None:
                out[k] = _num(d.get(k))
        return out

    def salvar_material(self, pid: int, dados: dict, bid: int | None = None) -> dict:
        d = self._limpar_material(dados)
        antes = (self.db.binomio(bid) or {}).get("dados", {}) if bid else {}
        for k in ("aplicado", "conferido"):
            if not d.get(k):
                d[k] = antes.get(k)
        bid = self.db.salvar_binomio(pid, TIPO_MATERIAL, d, bid)
        self.db.registrar(pid, "binômio: material", d["nome"])
        return self.item(pid, bid)

    def salvar_produto(self, pid: int, dados: dict, bid: int | None = None) -> dict:
        d = self._limpar_produto(dados)
        antes = (self.db.binomio(bid) or {}).get("dados", {}) if bid else {}
        for k in ("aplicado", "conferido"):
            if not d.get(k):
                d[k] = antes.get(k)
        bid = self.db.salvar_binomio(pid, TIPO_BETUMINOSO, d, bid)
        self.db.registrar(pid, "binômio: betuminoso", d["produto"])
        return self.item(pid, bid)

    def excluir(self, pid: int, bid: int) -> None:
        for a in self.db.anexos(pid, bid):
            self.excluir_anexo(a["id"])
        self.db.excluir_binomio(bid)

    def importar_de_planilha(self, pid: int, caminho: str, substituir: bool = False) -> dict:
        d = bn.importar_planilha(caminho)
        if not d["materiais"] and not d["produtos"]:
            raise ValueError("A planilha não tem as abas BINOMIO_COTACOES / BINOMIO_ANP.")
        if substituir:
            for b in self.db.binomios(pid):
                self.excluir(pid, b["id"])
        for m in d["materiais"]:
            self.db.salvar_binomio(pid, TIPO_MATERIAL, self._limpar_material(m))
        for p in d["produtos"]:
            self.db.salvar_binomio(pid, TIPO_BETUMINOSO, self._limpar_produto(p))
        try:  # índices da mesma planilha (se houver)
            self._gravar_series(bn.ler_indices(caminho), Path(caminho).name)
        except ValueError:
            pass
        self.db.registrar(pid, "binômio: importado", f"{Path(caminho).name}: {len(d['materiais'])} materiais, "
                                                     f"{len(d['produtos'])} betuminosos")
        return dict(self.listar(pid), importados=dict(materiais=len(d["materiais"]), produtos=len(d["produtos"])))

    # ---------------------------------------------- aplicar o vencedor
    def aplicar_vencedor(self, pid: int, bid: int, aplicar_dmt: bool = True, codigo: str | None = None) -> dict:
        r = self.item(pid, bid)
        v = r.get("vencedor")
        if not v:
            raise ValueError("Não há cotação válida para escolher o vencedor.")
        cod = cp._cod(codigo) if codigo else r.get("codigo")
        if not cod:
            raise ValueError("Informe o código do insumo do SICRO (ex.: M0191) para aplicar o preço nas composições.")
        cu = self.sv.custos
        feitas, msgs = [], []
        if "produto" in r:  # betuminoso: custo de aquisição com ICMS (o BDI é aplicado na composição)
            preco = v["aq_icms"]
            feitas.append(cu.alterar_insumo(pid, "MAT", cod, {"preco": preco}, "AMBOS",
                                            f"binômio ANP — {v.get('local')}/{v.get('uf')} (aquisição c/ ICMS, R$/t)"))
            msgs.append(f"{cod}: R$ {preco:.2f}/t (aquisição ANP com ICMS – {v.get('local')}).")
            forn = v.get("local")
        else:
            un_sicro = ((r.get("sicro") or {}).get("unidade") or "t").replace("³", "3").lower()
            preco = v["preco_m3_reaj"] if un_sicro.startswith("m") else v["preco_t_reaj"]
            forn = v.get("fornecedor")
            if v.get("sicro"):
                msgs.append(f"O preço do SICRO ({cod}) já é o mais vantajoso; preço mantido.")
            else:
                feitas.append(cu.alterar_insumo(pid, "MAT", cod, {"preco": preco}, "AMBOS",
                                                f"binômio — {forn} (cotação reajustada)"))
                msgs.append(f"{cod}: R$ {preco:.2f}/{'m³' if un_sicro.startswith('m') else 't'} ({forn}).")
            dp, dn = v.get("dmt_pav") or 0, v.get("dmt_npav") or 0
            if aplicar_dmt and (dp or dn):
                feitas.append(cu.alterar_dmt(pid, None, cod, dmt_rp=dn, dmt_p=dp))
                msgs.append(f"DMT do transporte de {cod}: pavimentada {dp} km, não pavimentada {dn} km.")
        b = self.db.binomio(bid)
        b["dados"]["aplicado"] = dict(fornecedor=forn, preco=preco, codigo=cod, alteracoes=[f["id"] for f in feitas])
        self.db.salvar_binomio(pid, b["tipo"], b["dados"], bid)
        return dict(mensagens=msgs, alteracoes=feitas, item=self.item(pid, bid))

    def exportar(self, pid: int) -> dict:
        p = self.sv.projeto(pid)
        dest = self.sv._nome(p, "binomio", "xlsx")
        exportar_xlsx(self.listar(pid), p, str(dest))
        return dict(arquivo=str(dest), nome=dest.name)

    # ------------------------------------------------------------ anexos
    def anexar(self, pid: int, origem: str, nome: str, binomio_id: int | None = None, tipo: str = "cotacao",
               rotulo: str | None = None, descricao: str | None = None) -> dict:
        if binomio_id and not self.db.binomio(binomio_id):
            raise KeyError(f"Binômio {binomio_id} não encontrado.")
        pasta = self.pasta_anexos / str(pid)
        pasta.mkdir(parents=True, exist_ok=True)
        seguro = re.sub(r"[^\w.\- ]+", "_", Path(nome).name).strip() or "anexo"
        dest = pasta / f"{uuid.uuid4().hex[:8]}_{seguro}"
        shutil.copyfile(origem, dest)
        aid = self.db.inserir_anexo(dict(projeto_id=pid, binomio_id=binomio_id, tipo=tipo if tipo in TIPOS_ANEXO else "outro",
                                         rotulo=rotulo, nome=seguro, caminho=str(dest), descricao=descricao,
                                         tamanho=dest.stat().st_size))
        self.db.registrar(pid, "binômio: anexo", f"{seguro} ({TIPOS_ANEXO.get(tipo, tipo)})")
        return self.db.anexo(aid)

    def anexo(self, aid: int) -> dict:
        a = self.db.anexo(aid)
        if not a or not Path(a["caminho"]).exists():
            raise KeyError("Anexo não encontrado.")
        return a

    def excluir_anexo(self, aid: int) -> None:
        a = self.db.anexo(aid)
        if a:
            Path(a["caminho"]).unlink(missing_ok=True)
            self.db.excluir_anexo(aid)


# ------------------------------------------------------------------ exportação
def exportar_xlsx(lst: dict, projeto: dict, saida: str) -> str:
    """Planilha do binômio no modelo da planilha de referência (BINOMIO_COTACOES, RESUMO_COTACOES, BINOMIO_ANP),
    com fórmulas vivas para conferência."""
    import datetime as dt

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L

    wb = Workbook()
    fina = Side(style="thin", color="999999")
    borda = Border(left=fina, right=fina, top=fina, bottom=fina)
    cab = PatternFill("solid", fgColor="1F4E79")
    venc = PatternFill("solid", fgColor="C8E6C9")
    branco = Font(bold=True, color="FFFFFF")
    neg = Font(bold=True)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def data(m):
        try:
            return dt.date(int(m[:4]), int(m[5:7]), 1)
        except (TypeError, ValueError):
            return None

    def titulo(ws, linha, texto, ate):
        ws.cell(linha, 2, texto).font = Font(bold=True, size=12)
        ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=ate)

    def cabecalho(ws, linha, cols, col0=2):
        for j, t in enumerate(cols):
            c = ws.cell(linha, col0 + j, t)
            c.fill, c.font, c.alignment, c.border = cab, branco, centro, borda
        ws.row_dimensions[linha].height = 42

    db_ = lst["data_base"]
    ident = [("Projeto", projeto.get("nome")), ("Rodovia", projeto.get("rodovia")), ("Trecho", projeto.get("trecho")),
             ("Segmento", f"km {projeto.get('km_ini')} ao km {projeto.get('km_fim')}"), ("Data-base", db_)]

    # ---------------------------------------------------------- cotações
    ws = wb.active
    ws.title = "BINOMIO_COTACOES"
    for i, (k, v) in enumerate(ident):
        ws.cell(2 + i, 2, k).font = neg
        ws.cell(2 + i, 3, v)
    titulo(ws, 14, "VERIFICAÇÃO DO BINÔMIO - AQUISIÇÃO E TRANSPORTE DE MATERIAIS PÉTREOS E OUTROS MATERIAIS COTADOS", 23)
    cabecalho(ws, 15, ["COTAÇÃO", "CÓDIGO SICRO", "FORNECEDOR", "PREÇO UNIT (T)", "DENSIDADE", "PREÇO UNIT (M3)", "DATA DA COTAÇÃO",
                       "ÍNDICE DE REAJ.", "DATA-BASE", "ÍNDICE DE REAJ.", "FATOR", "PREÇO UNIT (T)", "PREÇO UNIT (M3)",
                       "DMT (PAV)", "DMT (ÑPAV)", "CUSTO UNIT. (PAV)", "CUSTO UNIT. (ÑPAV)", "TOTAL TRANSPORTE (T)",
                       "TOTAL BINÔMIO AQUISIÇÃO E TRANSPORTE", "ÍNDICE", "GRUPO", "VANTAJOSIDADE", "VENCEDOR"])
    for j, t in enumerate(["", "", "", "li", "t/m³", "li", "mês/ano", "li", "mês/ano", "lo", "K", "lo", "lo", "KM", "KM",
                           "T.KM", "T.KM", "", "", "", "", "", ""]):
        ws.cell(16, 2 + j, t).alignment = centro
    lin = 17
    blocos = []
    for m in lst["materiais"]:
        if not m["linhas"]:
            continue
        ini = lin
        for x in m["linhas"]:
            r = lin
            ws.cell(r, 2, m["nome"] if r == ini else None)
            ws.cell(r, 3, m.get("codigo"))
            ws.cell(r, 4, x["fornecedor"])
            ws.cell(r, 6, x["densidade"])
            cotado = x["preco"] if x.get("preco") is not None and not x.get("sicro") else None
            if x["unidade"] == "t":
                ws.cell(r, 5, cotado if cotado is not None else x["preco_t"])
                ws.cell(r, 7, f"=E{r}*F{r}")
            else:
                ws.cell(r, 7, cotado if cotado is not None else x["preco_m3"])
                ws.cell(r, 5, f"=G{r}/F{r}")
            ws.cell(r, 8, data(x["data"])).number_format = "mmm/yy"
            ws.cell(r, 9, x["indice_cotacao"])
            ws.cell(r, 10, data(db_)).number_format = "mmm/yy"
            ws.cell(r, 11, x["indice_base"])
            ws.cell(r, 12, f'=IF(OR(I{r}="",K{r}=""),0,K{r}/I{r}-1)')
            ws.cell(r, 13, f"=TRUNC(E{r}*(1+L{r}),2)")
            ws.cell(r, 14, f"=TRUNC(G{r}*(1+L{r}),2)")
            ws.cell(r, 15, x.get("dmt_pav") or 0)
            ws.cell(r, 16, x.get("dmt_npav") or 0)
            ws.cell(r, 17, m.get("custo_pav") or 0)
            ws.cell(r, 18, m.get("custo_npav") or 0)
            ws.cell(r, 19, f"=O{r}*Q{r}+P{r}*R{r}")
            ws.cell(r, 20, f"=S{r}+M{r}")
            ws.cell(r, 21, m["indice_nome"] if r == ini else None)
            ws.cell(r, 22, GRUPOS_TXT.get(m.get("grupo"), m.get("grupo")) if r == ini else None)
            for c in range(2, 25):
                ws.cell(r, c).border = borda
                if x["vencedor"]:
                    ws.cell(r, c).fill = venc
            lin += 1
        fim = lin - 1
        ws.cell(ini, 23, f"=SMALL(T{ini}:T{fim},1)").font = neg
        ws.cell(ini, 24, f"=INDEX(D{ini}:D{fim},MATCH(W{ini},T{ini}:T{fim},0))").font = neg
        blocos.append((m, ini, fim))
        lin += 1
    for c, w in zip(range(2, 25), [26, 10, 22, 11, 9, 11, 10, 10, 10, 10, 9, 11, 11, 8, 8, 9, 9, 11, 13, 20, 16, 13, 22]):
        ws.column_dimensions[L(c)].width = w
    for r in ws.iter_rows(min_row=17, max_row=lin, min_col=5, max_col=20):
        for c in r:
            if c.column in (5, 7, 13, 14, 19, 20, 23):
                c.number_format = "#,##0.00"
            elif c.column == 12:
                c.number_format = "0.0000%"
    ws.freeze_panes = "D17"

    # ------------------------------------------------------------ resumo
    wr = wb.create_sheet("RESUMO_COTACOES")
    titulo(wr, 2, "RESUMO DO BINÔMIO - AQUISIÇÃO E TRANSPORTE DE MATERIAIS PÉTREOS E OUTROS MATERIAIS COTADOS", 9)
    wr.cell(3, 2, f"Data-base: {db_}")
    cabecalho(wr, 5, ["MATERIAL", "CÓDIGO SICRO", "TOTAL BINÔMIO (R$/T)", "ORIGEM", "PREÇO UNITÁRIO DE AQUISIÇÃO (R$/M³)",
                      "PREÇO UNITÁRIO DE AQUISIÇÃO (R$/T)", "DMT EM RODOVIA PAVIMENTADA (KM)", "DMT EM RODOVIA NÃO PAVIMENTADA (KM)"])
    for i, (m, a, b) in enumerate(blocos):
        r = 6 + i
        rng = lambda col: f"BINOMIO_COTACOES!${col}${a}:${col}${b}"  # noqa: E731
        idx = f"MATCH(D{r},{rng('T')},0)"
        vals = [m["nome"], m.get("codigo"), f"=BINOMIO_COTACOES!W{a}", f"=INDEX({rng('D')},{idx})",
                f"=INDEX({rng('N')},{idx})", f"=INDEX({rng('M')},{idx})", f"=INDEX({rng('O')},{idx})", f"=INDEX({rng('P')},{idx})"]
        for j, v in enumerate(vals):
            c = wr.cell(r, 2 + j, v)
            c.border = borda
            if j in (2, 4, 5):
                c.number_format = "#,##0.00"
    for c, w in zip(range(2, 10), [30, 12, 14, 24, 16, 16, 16, 16]):
        wr.column_dimensions[L(c)].width = w

    # -------------------------------------------------------------- ANP
    wa = wb.create_sheet("BINOMIO_ANP")
    titulo(wa, 2, "CUSTO DO MATERIAL BETUMINOSO - BINÔMIO DE AQUISIÇÃO E TRANSPORTE", 24)
    wa.cell(3, 2, "Portaria DNIT nº 1977/2017 e Portaria complementar nº 434/2017 · Transporte rodoviário = 26,939 + 0,253·D1 + 0,299·D2 + "
                  "0,412·D3 · Transporte fluvial = 20,7256 + 0,1603·D · Kr: Pavimentação jul/2014 = 270,237 · Kf: nov/2016 = 303,752")
    r = 5
    for p in lst["produtos"]:
        if not p["linhas"]:
            continue
        wa.cell(r, 2, p["produto"]).font = Font(bold=True, size=11)
        wa.cell(r, 6, "Índice (data-base)")
        wa.cell(r, 8, p.get("indice_atual"))
        wa.cell(r, 9, "Kr")
        wa.cell(r, 10, f"=TRUNC((H{r}-270.237)/270.237,3)")
        wa.cell(r, 11, "Kf")
        wa.cell(r, 12, f"=TRUNC((H{r}-303.752)/303.752,3)")
        wa.cell(r, 13, "BDI transp.")
        wa.cell(r, 14, p.get("bdi_transporte"))
        wa.cell(r, 15, "t/viagem")
        wa.cell(r, 16, p.get("ton_viagem"))
        k = r
        cabecalho(wa, r + 1, ["ORIGEM", "UF", "LOCAL", "FLU (D)", "PAV (D1)", "RP (D2)", "LN (D3)", "ICMS", "COFINS", "PIS/PASEP",
                              "BDI", "AQUISIÇÃO ANP (R$/T)", "PEDÁGIO (viagem)", "BALSA (viagem)", "TRANSPORTE S/ ICMS",
                              "AQUISIÇÃO C/ ICMS", "TRANSPORTE C/ ICMS", "PEDÁGIO (R$/T)", "BALSA (R$/T)", "AQUISIÇÃO C/ BDI",
                              "TRANSPORTE C/ BDI", "TOTAL (R$/T)", "VENCEDOR"])
        ini = r + 2
        for i, o in enumerate(p["linhas"]):
            x = ini + i
            vals = [o.get("nome"), o.get("uf"), o.get("local"), o.get("d_fluvial") or 0, o.get("d_pav") or 0, o.get("d_rp") or 0,
                    o.get("d_ln") or 0, o.get("icms") or 0, p.get("cofins"), p.get("pis"), p.get("bdi"), o.get("aq_sem_icms"),
                    o.get("pedagio") or 0, o.get("balsa") or 0]
            for j, v in enumerate(vals):
                wa.cell(x, 2 + j, v)
            flu = f"(20.7256+0.1603*E{x})*(1+$L${k})"
            if not p.get("fluvial_sempre", True):
                flu = f"IF(E{x}>0,{flu},0)"
            wa.cell(x, 16, f"=TRUNC((26.939+0.253*F{x}+0.299*G{x}+0.412*H{x})*(1+$J${k})+{flu},2)")
            wa.cell(x, 17, f"=TRUNC(M{x}/(1-(I{x}-((K{x}+J{x})*I{x}))),2)")
            wa.cell(x, 18, f"=TRUNC(P{x}/(1-(I{x}+J{x}+K{x}-((K{x}+J{x})*I{x}))),2)")
            wa.cell(x, 19, f"=N{x}/$P${k}")
            wa.cell(x, 20, f"=O{x}/$P${k}")
            wa.cell(x, 21, f"=TRUNC(Q{x}*(1+L{x}),2)")
            wa.cell(x, 22, f"=TRUNC((R{x}+S{x}+T{x})*(1+$N${k}),2)")
            wa.cell(x, 23, f"=U{x}+V{x}")
            fim = ini + len(p["linhas"]) - 1
            wa.cell(x, 24, f'=IF(W{x}=MIN($W${ini}:$W${fim}),"VENCEDOR","")')
            for c in range(2, 25):
                wa.cell(x, c).border = borda
                if o["vencedor"]:
                    wa.cell(x, c).fill = venc
                if c in (9, 10, 11, 12):
                    wa.cell(x, c).number_format = "0.00%"
                elif c >= 13:
                    wa.cell(x, c).number_format = "#,##0.00"
        r = ini + len(p["linhas"]) + 1
    for c, w in zip(range(2, 25), [9, 6, 20, 8, 8, 8, 8, 8, 8, 8, 9, 13, 11, 10, 12, 13, 12, 10, 10, 13, 13, 13, 11]):
        wa.column_dimensions[L(c)].width = w

    # ---------------------------------------------------- índices e anexos
    wi = wb.create_sheet("INDICES_USADOS")
    cabecalho(wi, 1, ["Item", "Índice", "Mês", "Valor"], 1)
    lin = 2
    vistos = set()
    for m in lst["materiais"]:
        for x in m["linhas"]:
            for mes_, v in ((x["data"], x["indice_cotacao"]), (db_, x["indice_base"])):
                if (m["indice_nome"], mes_) not in vistos and v is not None:
                    vistos.add((m["indice_nome"], mes_))
                    for j, val in enumerate([m["nome"], m["indice_nome"], mes_, v]):
                        wi.cell(lin, 1 + j, val)
                    lin += 1
    for p in lst["produtos"]:
        for j, val in enumerate([p["produto"], p.get("indice") or bn.INDICE_ASFALTO, db_, p.get("indice_atual")]):
            wi.cell(lin, 1 + j, val)
        lin += 1
    for c, w in zip("ABCD", [40, 34, 10, 12]):
        wi.column_dimensions[c].width = w

    wx = wb.create_sheet("ANEXOS")
    cabecalho(wx, 1, ["Item", "Tipo", "Fornecedor/origem", "Arquivo", "Anexado em"], 1)
    lin = 2
    for it in [(m["nome"], m["anexos"]) for m in lst["materiais"]] + [(p["produto"], p["anexos"]) for p in lst["produtos"]] + \
            [("Projeto", lst["anexos_gerais"])]:
        for a in it[1]:
            for j, val in enumerate([it[0], TIPOS_ANEXO.get(a["tipo"], a["tipo"]), a.get("rotulo"), a["nome"], a.get("criado_em")]):
                wx.cell(lin, 1 + j, val)
            lin += 1
    for c, w in zip("ABCDE", [40, 24, 24, 44, 20]):
        wx.column_dimensions[c].width = w
    wb.save(saida)
    return saida


GRUPOS_TXT = bn.GRUPOS
