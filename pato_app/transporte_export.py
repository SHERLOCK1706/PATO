"""Planilha da Fase 3 no modelo do PATO: CROQUI (desenho em células + DMT), VMDf, FIT e TRANSPORTES_CONSERVA,
com fórmulas vivas (FIT e momentos) para conferência."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from . import transporte as tp

FINA = Side(style="thin", color="999999")
BORDA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
CAB = PatternFill("solid", fgColor="1F4E79")
BRANCO = Font(bold=True, color="FFFFFF")
NEG = Font(bold=True)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
PISTA = PatternFill("solid", fgColor="595959")
URBANO = PatternFill("solid", fgColor="F4B183")
OCORR = PatternFill("solid", fgColor="A9D18E")


def _cab(ws, linha, cols, col0=2):
    for j, t in enumerate(cols):
        c = ws.cell(linha, col0 + j, t)
        c.fill, c.font, c.alignment, c.border = CAB, BRANCO, CENTRO, BORDA
    ws.row_dimensions[linha].height = 36


def _ident(ws, p, titulo, ate=12):
    ws.cell(2, 2, titulo).font = Font(bold=True, size=13)
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=ate)
    for i, (k, v) in enumerate((("Rodovia", p.get("rodovia")), ("Trecho", p.get("trecho")), ("Subtrecho", p.get("subtrecho")),
                                ("Segmento", f"km {p.get('km_ini')} ao km {p.get('km_fim')}"), ("SNV", p.get("snv")))):
        ws.cell(4 + i, 2, k).font = NEG
        ws.cell(4 + i, 3, v)


def exportar_xlsx(tr, pid: int, saida: str) -> str:
    p = tr.sv.projeto(pid)
    cro = tr.croqui(pid)
    d = cro["dados"]
    q = tr.quadro(pid)
    wb = Workbook()

    # ------------------------------------------------------------------ CROQUI
    ws = wb.active
    ws.title = "CROQUI"
    _ident(ws, p, "CROQUI DE OCORRÊNCIA DE MATERIAIS E DMT's", 14)
    a, b = cro["km_ini"], cro["km_fim"]
    eventos = [(a, "INÍCIO DO TRECHO", "marco", None)]
    for o in d.get("ocorrencias", []):
        if o.get("km") not in (None, ""):
            eventos.append((float(o["km"]), f"{o['nome']} ({tp.TIPOS.get(o.get('tipo'), o.get('tipo'))})", "ocorr", o))
    for t in d.get("travessias", []):
        eventos.append((float(t["km_ini"]), f"Travessia urbana {t.get('nome') or ''} — início", "urb_i", t))
        eventos.append((float(t["km_fim"]), f"Travessia urbana {t.get('nome') or ''} — fim", "urb_f", t))
    for lo in d.get("localidades", []):
        if lo.get("km") not in (None, ""):
            eventos.append((float(lo["km"]), lo.get("nome") or "", "local", lo))
    eventos.append((b, "FIM DO TRECHO", "marco", None))
    eventos.sort(key=lambda e: e[0])
    _cab(ws, 11, ["KM", "", "PISTA", "", "OCORRÊNCIA / REFERÊNCIA", "ACESSO LN (km)", "ACESSO RP (km)", "ACESSO P (km)"])
    lin = 12
    urb = 0
    for km, txt, tipo, o in eventos:
        if tipo == "urb_i":
            urb += 1
        ws.cell(lin, 2, km).number_format = "0.000"
        c = ws.cell(lin, 4)
        c.fill = URBANO if (urb > 0 or tipo == "urb_f") else PISTA
        ws.cell(lin, 6, txt).font = NEG if tipo in ("marco", "ocorr") else Font(italic=tipo != "ocorr")
        if tipo == "ocorr":
            lado = 5 if (o.get("lado") or "D").upper().startswith("E") else 3
            ws.cell(lin, lado).fill = OCORR
            for j, k in enumerate(("acesso_ln", "acesso_rp", "acesso_p")):
                ws.cell(lin, 7 + j, o.get(k) or 0)
        if tipo == "urb_f":
            urb -= 1
        lin += 1
        # trecho de pista até o próximo evento
        ws.cell(lin, 4).fill = URBANO if urb > 0 else PISTA
        lin += 1
    fixas = [o for o in d.get("ocorrencias", []) if o.get("km") in (None, "")]
    if fixas:
        lin += 1
        ws.cell(lin, 2, "Ocorrências com distância fixa até a pista").font = NEG
        lin += 1
        _cab(ws, lin, ["OCORRÊNCIA", "TIPO", "LN (km)", "RP (km)", "P (km)"])
        for o in fixas:
            lin += 1
            for j, v in enumerate([o["nome"], tp.TIPOS.get(o.get("tipo"), o.get("tipo")), o.get("dist_ln") or 0,
                                   o.get("dist_rp") or 0, (b - a) / 2 if o.get("meia_extensao") else (o.get("dist_p") or 0)]):
                ws.cell(lin, 2 + j, v).border = BORDA
    for c, w in zip(range(2, 10), [10, 3, 8, 3, 48, 14, 14, 14]):
        ws.column_dimensions[L(c)].width = w

    # DMT
    wd = wb.create_sheet("DMT")
    _ident(wd, p, "CÁLCULO DAS DISTÂNCIAS MÉDIAS DE TRANSPORTE", 14)
    wd.cell(10, 2, "DMT média até a pista = ((x − a)² + (b − x)²) / (2·(b − a)); Pr = P − Pu"
                   + (" · parcela urbana pelo método da planilha de referência" if d.get("metodo") != "exato"
                      else " · parcela urbana exata (destinos uniformes no segmento)"))
    _cab(wd, 12, ["ORIGEM", "DESTINO", "LN", "RP", "P", "Pr", "Pu", "DMT TOTAL", "MÉTODO", "FIT"])
    for i, r in enumerate(cro["rotas"]):
        x = 13 + i
        for j, v in enumerate([r["origem"], r["destino"], r["ln"], r["rp"], r["p"], r["pr"], r["pu"], r["dmt"],
                               r.get("metodo"), "X" if r.get("fit", 1) != 1 else ""]):
            c = wd.cell(x, 2 + j, v)
            c.border = BORDA
            if 4 <= 2 + j <= 9:
                c.number_format = "0.000"
    for c, w in zip(range(2, 12), [20, 20, 9, 9, 9, 9, 9, 11, 20, 6]):
        wd.column_dimensions[L(c)].width = w

    # ------------------------------------------------------------------ VMDf
    wv = wb.create_sheet("VMDf")
    _ident(wv, p, "VOLUME MÉDIO DIÁRIO POR FAIXA", 12)
    wv.cell(10, 2, f"Fonte: modelagem do VMDa — DNIT/PNCT ({tr.croqui(pid)['fonte_vmd']})")
    _cab(wv, 12, ["SEGMENTO", "SNV", "ORIGEM (km)", "DESTINO (km)", "EXT. (km)", "VMDa", "FAIXAS/SENTIDO",
                  "VMDf CRESCENTE", "VMDf DECRESCENTE", "FONTE"])
    for i, s in enumerate(d.get("vmd", [])):
        x = 13 + i
        vals = [s.get("segmento"), s.get("snv"), s.get("km_ini"), s.get("km_fim"), f"=E{x}-D{x}", s.get("vmda"),
                s.get("faixas_sentido") or 1, tp.vmdf_segmento(s, "C"), tp.vmdf_segmento(s, "D"), s.get("fonte")]
        for j, v in enumerate(vals):
            wv.cell(x, 2 + j, v).border = BORDA
    for c, w in zip(range(2, 12), [10, 14, 11, 11, 9, 10, 10, 12, 12, 34]):
        wv.column_dimensions[L(c)].width = w

    # ------------------------------------------------------------------- FIT
    wf = wb.create_sheet("FIT")
    _ident(wf, p, "FATOR DE INTERFERÊNCIA DE TRÁFEGO - FIT", 12)
    wf.cell(10, 2, "A1 = 70 + 0,0035·VMDf (s) · A2 = 60·DMTu (s) · FIT = 1 + (A1 + A2) / (DMTp·60) "
                   "— Manual de Custos de Infraestrutura de Transportes (DNIT), nova metodologia")
    _cab(wf, 12, ["ORIGEM", "DESTINO", "SENTIDO", "DMTp (km)", "DMTu (km)", "VMDf", "A1", "A2", "FIT"])
    linha_fit = {}
    x = 13
    for r in cro["rotas"]:
        if r.get("fit", 1) == 1:
            continue
        vals = [r["origem"], r["destino"], r.get("sentido"), r["p"], r["pu"], r.get("vmdf"), f"=70+0.0035*G{x}",
                f"=60*F{x}", f"=1+((H{x}+I{x})/(E{x}*60))"]
        for j, v in enumerate(vals):
            wf.cell(x, 2 + j, v).border = BORDA
        wf.cell(x, 10).number_format = "0.0000"
        linha_fit[(r["origem"], r["destino"])] = x
        x += 1
    for c, w in zip(range(2, 11), [20, 20, 14, 11, 11, 9, 9, 9, 10]):
        wf.column_dimensions[L(c)].width = w

    # ---------------------------------------------------- TRANSPORTES_CONSERVA
    wt = wb.create_sheet("TRANSPORTES_CONSERVA")
    _ident(wt, p, "QUADRO DE TRANSPORTES - SERVIÇOS DE CONSERVAÇÃO RODOVIÁRIA", 15)
    lin = 11
    totais_cel = []
    por_comp: dict[str, list] = {}
    for ln in q["linhas"]:
        for cod, mom in (ln.get("momentos") or {}).items():
            por_comp.setdefault(cod, []).append((ln, cod))
    descricoes = {t["codigo"]: t.get("descricao") for t in q["totais"]}
    for cod in sorted(por_comp):
        sup = tp.superficie(descricoes.get(cod) or "")
        wt.cell(lin, 2, cod).font = NEG
        wt.cell(lin, 3, descricoes.get(cod)).font = NEG
        lin += 1
        _cab(wt, lin, ["CÓDIGO", "DISCRIMINAÇÃO", "MATERIAL", "QUANT. ANUAL", "F. UTILIZAÇÃO (t/un)", "PESO (t)",
                       "ORIGEM", "DESTINO", "DMT (km)", "FIT", "MOMENTO (tkm)"])
        ini = lin + 1
        for ln, _ in por_comp[cod]:
            lin += 1
            dist = {"p": ln.get("dmt_p"), "rp": ln.get("dmt_rp"), "ln": ln.get("dmt_ln")}[sup]
            fx = linha_fit.get((ln.get("origem"), ln.get("destino"))) if sup == "p" else None
            fit = f"=FIT!J{fx}" if fx and ln.get("fit", 1) != 1 else 1
            vals = [ln.get("servico"), ln.get("servico_desc"), ln.get("material"), ln.get("quantidade_anual") or 0,
                    ln.get("fator"), f"=E{lin}*F{lin}", ln.get("origem"), ln.get("destino"), dist, fit,
                    f"=TRUNC(G{lin}*J{lin}*K{lin},3)"]
            for j, v in enumerate(vals):
                c = wt.cell(lin, 2 + j, v)
                c.border = BORDA
            wt.cell(lin, 11).number_format = "0.0000"
            wt.cell(lin, 12).number_format = "#,##0.000"
        lin += 1
        wt.cell(lin, 2, "TOTAL:").font = NEG
        c = wt.cell(lin, 12, f"=SUM(L{ini}:L{lin - 1})")
        c.font, c.number_format = NEG, "#,##0.000"
        totais_cel.append((cod, descricoes.get(cod), f"TRANSPORTES_CONSERVA!L{lin}"))
        lin += 2
    for c, w in zip(range(2, 13), [10, 46, 22, 13, 12, 11, 14, 16, 10, 9, 14]):
        wt.column_dimensions[L(c)].width = w

    # resumo
    wr = wb.create_sheet("RESUMO_TRANSPORTES")
    _ident(wr, p, "MOMENTOS DE TRANSPORTE (quantidades dos serviços 01.05 do PATO)", 6)
    _cab(wr, 11, ["CÓDIGO", "SERVIÇO DE TRANSPORTE", "MOMENTO (tkm)"])
    for i, (cod, desc, ref) in enumerate(totais_cel):
        for j, v in enumerate([cod, desc, f"={ref}"]):
            c = wr.cell(12 + i, 2 + j, v)
            c.border = BORDA
        wr.cell(12 + i, 4).number_format = "#,##0.000"
    wr.column_dimensions["B"].width = 12
    wr.column_dimensions["C"].width = 90
    wr.column_dimensions["D"].width = 16
    wb.save(saida)
    return saida
