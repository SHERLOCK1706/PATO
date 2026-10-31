"""Planilha PATO (Resolução DNIT nº 8/2023, Anexo XIX) — quantidades × preços, nos dois regimes.

Cada item: quantidade = TRUNC|ARRED(base × NE; 3), em que a base é
  • medida do inventário (vínculo vivo com o inventário do projeto),
  • outro item (quantidade ou base dele — ex.: usinagem = tapa-buraco × 2,4 t/m³),
  • momento do quadro de transportes (t·km),
  • valor informado (com a memória de cálculo da planilha de origem),
ou quantidade direta (ex.: itens do levantamento de condição do pavimento – IP).

Preço unitário (com BDI) por regime: composição SICRO/própria (motor, com as alterações do projeto); lista sintética;
aquisição/transporte de betuminosos (AM/TM — binômio ANP); grupos de desempenho (GD — parcela mensal);
preços informados (administração local, canteiro, instalação industrial, mobilização).
Parcial = TRUNC(quantidade × preço; 2), como na planilha PATO.
"""
from __future__ import annotations

import math
import re
from pathlib import Path

import openpyxl

from . import composicoes as cp
from . import normalizar as nz

# seções do Anexo XIX (grupos de tarefas) — usadas no resumo, no % dos grupos (Anexo XX) e no cronograma
SECOES = {
    "01.01": "Grupo de serviços por desempenho",
    "01.02": "Conservação corretiva rotineira",
    "01.03": "Conservação preventiva periódica",
    "01.04": "Conservação de emergência",
    "01.05": "Serviços de transportes",
    "01.06": "Sinalização de obras temporárias",
    "01.07": "Serviços diversos (administração local, canteiro, mobilização)",
    "01.08": "Serviços incluídos pela lista (classificar no grupo correto)",
    "02": "Demais serviços de manutenção – DSM",
}
LIMITE_DSM = 0.25  # Res. 8/2023, art. 6º, III


def trunc(v: float, casas: int = 2) -> float:
    f = 10 ** casas
    return math.trunc(float(f"{float(f'{v:.15g}') * f:.15g}")) / f


def arred(v: float, casas: int = 3) -> float:
    # ARRED do Excel (meio para longe do zero)
    f = 10 ** casas
    x = float(f"{v * f:.15g}")
    return math.floor(abs(x) + 0.5) * (1 if x >= 0 else -1) / f


def _txt(v) -> str:
    return "" if v is None else str(v).strip()


_PRODUTO_ASF = re.compile(r"(CAP\s*\d+\s*/\s*\d+|RR-?\s?1C-?E|RR-?\s?2C-?E|RR-?\s?1C|RR-?\s?2C|RC-?\s?1C-?E|RM-?\s?1C|EAI|ADP|CM-?\s?30|"
                          r"AMP|ASFALTO BORRACHA)", re.I)


def chave_asfalto(desc: str | None) -> str | None:
    m = _PRODUTO_ASF.search(nz.sem_acento(desc or "").upper())
    return re.sub(r"[\s-]", "", m.group(1)).upper() if m else None


# --------------------------------------------------------------- importação
def importar_referencia(caminho: str) -> dict:
    """Lê as abas PATO, DESEMPENHO e BD de uma planilha PATO (fórmulas e valores)."""
    wf = openpyxl.load_workbook(caminho, data_only=False)
    wv = openpyxl.load_workbook(caminho, data_only=True)
    try:
        if "PATO" not in wf.sheetnames:
            raise ValueError("A planilha não tem a aba PATO.")
        ws, vs = wf["PATO"], wv["PATO"]
        try:
            from . import ip as _ip
            dados_ip = _ip.importar_planilha(caminho)
        except ValueError:
            dados_ip = None
        itens: list[dict] = []
        por_linha: dict[int, dict] = {}
        fim = ws.max_row
        for r in range(1, ws.max_row + 1):
            if "PREÇO TOTAL DOS DEMAIS" in nz.sem_acento(_txt(vs.cell(r, 3).value)).upper().replace("PRECO", "PREÇO"):
                fim = r
        for r in range(12, fim + 1):
            item, cod, desc = _txt(vs.cell(r, 3).value), vs.cell(r, 4).value, vs.cell(r, 5).value
            if not item or not re.match(r"^\d{2}(\.\d{2})*", item):
                continue
            base = dict(id=f"L{r}", item=item, descricao=_txt(desc))
            if cod in (None, ""):
                base["tipo"] = "titulo"
                itens.append(base)
                por_linha[r] = base
                continue
            F, H, J = ws.cell(r, 6).value, ws.cell(r, 8).value, ws.cell(r, 10).value
            fv, hv, jv = vs.cell(r, 6).value, vs.cell(r, 8).value, vs.cell(r, 10).value
            x = dict(base, tipo="servico", codigo=cp._cod(str(cod)), origem=_txt(vs.cell(r, 2).value),
                     unidade_base=_txt(vs.cell(r, 7).value), unidade=_txt(vs.cell(r, 11).value),
                     ne=nz.numero(hv) if hv not in (None, "") else 1.0, ne_expr=F_txt(H),
                     arred="arred" if isinstance(J, str) and "ROUND" in J.upper() else "trunc",
                     preco_ref=dict(SD=nz.numero(vs.cell(r, 12).value), CD=nz.numero(vs.cell(r, 14).value)),
                     quantidade_ref=nz.numero(jv), _F=F, _fv=fv, _J=J, _jv=jv, _linha=r)
            itens.append(x)
            por_linha[r] = x
        # base de cada item
        for x in itens:
            if x["tipo"] != "servico":
                continue
            F, fv, J = x.pop("_F"), x.pop("_fv"), x.pop("_J")
            jv = x.pop("_jv")
            x.pop("_linha")
            fs = F if isinstance(F, str) else None
            m = re.fullmatch(r"=\$?([JF])\$?(\d+)", (fs or "").replace(" ", ""))
            bip = None
            if dados_ip:
                from . import ip as _ip
                bip = _ip.base_do_pato(J if isinstance(J, str) else None, dados_ip["parametros"]) or \
                    _ip.base_do_pato(fs, dados_ip["parametros"])
            if bip:
                bip["valor"] = nz.numero(jv) if bip["campo"] == "pato" else nz.numero(fv)
                x["base"] = bip
                x["base_ref_valor"] = nz.numero(fv)
            elif isinstance(J, str) and not re.search(r"(TRUNC|ROUND)\(\s*\$?[FH]\$?\d+\s*\*\s*\$?[FH]\$?\d+", J.upper()):
                x["base"] = dict(tipo="direta", valor=nz.numero(jv), expr=F_txt(J), fonte=_fonte(J))
                x["base_ref_valor"] = nz.numero(fv)
            elif m and int(m.group(2)) in por_linha:
                x["base"] = dict(tipo="item", ref=por_linha[int(m.group(2))]["id"], campo="quantidade" if m.group(1) == "J" else "base")
            elif fs and "UN_DIA!" in fs.upper():
                x["base"] = dict(tipo="undia", grupo="conservacao", codigo=x["codigo"], valor=nz.numero(fv), expr=F_txt(fs),
                                 fonte="sinalização de obras (un.dia)")
            elif fs and "TRANSPORTES_CONSERVA!" in fs.upper():
                x["base"] = dict(tipo="transporte", codigo=x["codigo"], valor=nz.numero(fv))
            elif fs and fs.startswith("="):
                x["base"] = dict(tipo="valor", valor=nz.numero(fv), expr=F_txt(fs), fonte=_fonte(fs))
            else:
                x["base"] = dict(tipo="valor", valor=nz.numero(F) or 0.0, expr=None, fonte="informado")
        desempenho = _ler_desempenho(wf, wv)
        precos = _ler_bd(wv)
        prazo = None
        try:
            prazo = nz.numero(wv["PATO"]["R7"].value)
        except Exception:  # noqa: BLE001
            pass
        totais = {}
        for r in range(1, vs.max_row + 1):
            t = re.sub(r"\s+", " ", nz.sem_acento(_txt(vs.cell(r, 3).value)).upper())
            if t.startswith("TOTAL GERAL DO ORCAMENTO"):
                totais = dict(SD=nz.numero(vs.cell(r, 13).value), CD=nz.numero(vs.cell(r, 15).value))
        try:
            from . import ne_rocada as _ner
            dados_ner = _ner.importar_planilha(caminho)
        except Exception:  # noqa: BLE001
            dados_ner = None
        try:
            from . import sinalizacao as _sz
            dados_sz = _sz.importar_planilha(caminho)
        except Exception:  # noqa: BLE001 — modelo sem as abas de sinalização
            dados_sz = None
        return dict(itens=itens, desempenho=desempenho, precos_informados=precos, prazo_anos=prazo,
                    totais_ref=totais, origem=Path(caminho).name, ip=dados_ip, sinalizacao=dados_sz,
                    ne_rocada=dados_ner)
    finally:
        wf.close()
        wv.close()


def F_txt(v) -> str | None:  # noqa: N802 — memória de cálculo (fórmula da planilha de origem)
    return v if isinstance(v, str) and v.startswith("=") else None


def _fonte(f: str) -> str:
    u = (f or "").upper()
    for chave, rot in (("CALCULO_IP", "levantamento IP"), ("RESUMO_INVENT", "inventário"), ("CAD_", "inventário"),
                       ("UN_DIA", "sinalização de obras (un.dia)"), ("NE_ROCADA", "NE roçada"),
                       ("TRANSPORTES_CONSERVA", "quadro de transportes")):
        if chave in u:
            return rot
    return "memória de cálculo"


def _ler_desempenho(wf, wv) -> dict:
    if "DESEMPENHO" not in wf.sheetnames:
        return {}
    ws, vs = wf["DESEMPENHO"], wv["DESEMPENHO"]
    grupos: dict[str, dict] = {}
    atual = sub = None
    for r in range(1, ws.max_row + 1):
        b = _txt(vs.cell(r, 2).value)
        if re.match(r"^GD\d*\s*-", b) or (b.startswith("GD") and " - " in b):
            cod = b.split("-")[0].strip()
            atual = grupos.setdefault(cod, dict(codigo=cod, descricao=b.split("-", 1)[1].strip(), linhas=[]))
            continue
        if re.fullmatch(r"GD\d+", b):
            sub = f"{b} – {_txt(vs.cell(r, 3).value)}"
            continue
        if atual is None or not re.fullmatch(r"\d+[A-Z]?", b):
            continue
        D, Hf = ws.cell(r, 4).value, ws.cell(r, 8).value
        dv, hv = vs.cell(r, 4).value, vs.cell(r, 8).value
        auxiliar = not isinstance(Hf, str) or "TRUNC(" not in Hf.upper()
        atual["linhas"].append(dict(
            codigo=b, subgrupo=sub, descricao=_txt(vs.cell(r, 3).value), unidade=_txt(vs.cell(r, 9).value),
            qtd_total=dict(tipo="valor", valor=nz.numero(dv), expr=F_txt(D), fonte=_fonte(D or "")) if not auxiliar else None,
            ne=nz.numero(vs.cell(r, 5).value), freq=nz.numero(vs.cell(r, 6).value), desc_qtd=_txt(vs.cell(r, 7).value),
            qtd_km=nz.numero(hv) if auxiliar else None, qtd_km_expr=F_txt(Hf) if auxiliar else None,
            undia=bool(auxiliar and isinstance(Hf, str) and "UN_DIA" in Hf.upper()),
            freq_ne_rocada=bool(isinstance(ws.cell(r, 6).value, str) and "NE_ROCADA" in ws.cell(r, 6).value.upper())))
    return grupos


def _ler_bd(wv) -> dict:
    """Preços da aba BD que não vêm de composição SICRO (ADM, CAN, MOB…), por regime."""
    if "BD" not in wv.sheetnames:
        return {}
    out = {}
    for r in wv["BD"].iter_rows(min_row=4, values_only=True):
        r = tuple(r) + (None,) * 6
        cod = _txt(r[1])
        if cod and not cod[:1].isdigit() and r[4] is not None and not cod.startswith(("AM", "TM", "PN", "GD")):
            out[cod] = dict(descricao=_txt(r[2]), unidade=_txt(r[3]), SD=nz.numero(r[4]), CD=nz.numero(r[5]),
                            fonte="planilha PATO (BD)")
        elif cod[:2] in ("AM", "TM") and r[4] is not None:
            out[cod] = dict(descricao=_txt(r[2]), unidade=_txt(r[3]), SD=nz.numero(r[4]), CD=nz.numero(r[5]),
                            fonte="planilha PATO (RESUMO_ANP)", asfalto=True)
    return out


# --------------------------------------------------------------- cálculo
class Calculo:
    """Resolve quantidades e preços de um PATO (dados do projeto + motores + binômio + transporte)."""

    def __init__(self, dados: dict, medidas: dict[str, float], momentos: dict[str, float], preco_fn, extensao: float,
                 ip_calc: dict | None = None, producao_fn=None, anos_contrato: float | None = None,
                 sinal_dados: dict | None = None):
        self.d = dados
        self.medidas = medidas
        self.momentos = momentos
        self.preco_fn = preco_fn  # (codigo, regime) -> (preço, fonte) | (None, motivo)
        self.ext = extensao
        self.idx = {x["id"]: x for x in dados.get("itens", [])}
        self.ip = {c["chave"]: c for c in (ip_calc or {}).get("calculo", [])}
        self.producao_fn = producao_fn  # código -> produção horária da equipe (composição da base)
        self.anos = anos_contrato or nz.numero(dados.get("prazo_anos")) or 1.0
        self.sinal_dados = sinal_dados
        self.ne_rocada = None  # NE da roçada (DNIT 182/2018) — frequência das linhas de roçada/capina do GD
        self._sinal = None
        self._q: dict[str, tuple] = {}

    # ------------------------------------------------ sinalização de obras
    def gd_qtd_km(self, cod: str) -> float | None:
        """Quantidade por km de um serviço dos grupos por desempenho (aba DESEMPENHO, coluna H)."""
        for g in (self.d.get("desempenho") or {}).values():
            for ln in g["linhas"]:
                if ln["codigo"] == cod and ln.get("qtd_total"):
                    qt = ln["qtd_total"]
                    tot = (self.medidas[qt["medida"]] * (nz.numero(qt.get("fator")) or 1.0)
                           if qt.get("tipo") == "inventario" and self.medidas.get(qt.get("medida")) is not None
                           else nz.numero(qt.get("valor")) or 0.0)
                    return trunc(tot * (nz.numero(ln.get("ne")) or 1.0) * self.freq(ln) / (self.ext or 1), 3)
        return None

    def freq(self, ln: dict) -> float:
        liga = ln.get("freq_ne_rocada") or ("182/2018" in (ln.get("desc_qtd") or ""))  # projetos importados antes da regra
        if liga and self.ne_rocada is not None:
            return self.ne_rocada
        return nz.numero(ln.get("freq")) or 1.0

    def sinalizacao(self) -> dict | None:
        if self._sinal is None and self.sinal_dados:
            from . import sinalizacao as sz

            def qtd(sv_):
                if sv_.get("fonte") == "desempenho":
                    v = self.gd_qtd_km(sv_["codigo"])
                    return v if v is not None else nz.numero(sv_.get("quantidade_ref"))
                if sv_.get("fonte") == "pato":
                    lim = sv_.get("linhas_pato")
                    tot, achou = 0.0, False
                    for x in self.d.get("itens", []):
                        if x.get("tipo") != "servico" or x.get("codigo") != sv_["codigo"]:
                            continue
                        if (x.get("base") or {}).get("tipo") == "undia":
                            continue
                        mrow = re.fullmatch(r"L(\d+)", x["id"])
                        if lim and mrow and not (lim[0] <= int(mrow.group(1)) <= lim[1]):
                            continue
                        tot += self.quantidade(x)[0]
                        achou = True
                    return tot if achou else 0.0
                return nz.numero(sv_.get("quantidade_ref")) or 0.0
            self._sinal = sz.calcular(self.sinal_dados, qtd, self.producao_fn, self.ext)
        return self._sinal

    def base(self, x: dict) -> tuple[float, str]:
        b = x.get("base") or {}
        t = b.get("tipo")
        if t == "item":
            ref = self.idx.get(b.get("ref"))
            if not ref:
                return 0.0, "item de referência removido"
            q, bs, _ = self.quantidade(ref)
            return (q if b.get("campo", "quantidade") == "quantidade" else bs), f"{'quantidade' if b.get('campo', 'quantidade') == 'quantidade' else 'base'} do item {ref['item']}"
        if t == "inventario":
            v = self.medidas.get(b.get("medida"))
            if v is None:
                return nz.numero(b.get("valor")) or 0.0, f"medida “{b.get('medida')}” sem inventário — valor anterior"
            fator = nz.numero(b.get("fator")) if b.get("fator") not in (None, "") else 1.0
            return v * fator, f"inventário: {b.get('medida')}" + (f" × {fator:g}" if fator != 1 else "")
        if t == "transporte":
            v = self.momentos.get(b.get("codigo") or x.get("codigo"))
            if v is None:
                return nz.numero(b.get("valor")) or 0.0, "momento da planilha (gere o quadro de transportes)"
            return v, "quadro de transportes"
        if t == "producao":
            prod = self.producao_fn(x["codigo"]) if self.producao_fn else None
            fonte_p = "composição da base"
            if prod is None:
                prod, fonte_p = nz.numero(b.get("producao_ref")) or 0.0, "produção da planilha de origem"
            h, dias = nz.numero(b.get("horas")) or 8.0, nz.numero(b.get("dias")) or 1.0
            return prod * h * dias, f"produção {prod:g}/h ({fonte_p}) × {h:g} h × {dias:g} dia(s)"
        if t == "extensao":
            f = nz.numero(b.get("fator")) or 0.0
            return self.ext * f, f"extensão {self.ext:g} km × {f:g}"
        if t == "contrato":
            f = nz.numero(b.get("fator")) or 1.0
            return f / self.anos, f"{f:g} ÷ {self.anos:g} ano(s) de contrato"
        if t == "undia":
            sg = self.sinalizacao()
            if not sg:
                return nz.numero(b.get("valor")) or 0.0, "sinalização de obras não cadastrada — valor anterior"
            v = sg["grupos"][b.get("grupo", "conservacao")]["codigos"].get(b.get("codigo") or x["codigo"], 0)
            return float(v), "sinalização de obras (un.dia, Anexo XIII)"
        if t == "ip":
            if not self.ip:
                return nz.numero(b.get("valor")) or 0.0, "levantamento IP não lançado — valor anterior"
            v = sum((self.ip.get(c) or {}).get(b.get("campo", "pato"), 0.0) or 0.0 for c in b.get("chaves") or [])
            fator = nz.numero(b.get("fator")) or 1.0
            rot = " + ".join(b.get("chaves") or [])
            return v * fator, f"levantamento IP: {rot}" + (f" × {fator:g}" if fator != 1 else "")
        if t == "meses":
            return nz.numero(b.get("valor")) or 0.0, "prazo (meses)"
        return nz.numero(b.get("valor")) or 0.0, b.get("fonte") or "informado"

    def quantidade(self, x: dict) -> tuple[float, float, str]:
        if x["id"] in self._q:
            return self._q[x["id"]]
        self._q[x["id"]] = (0.0, 0.0, "referência circular")
        b = x.get("base") or {}
        if b.get("tipo") == "ip" and b.get("campo", "pato") == "pato":
            q, fonte = self.base(x)
            res = (round(q, 6), nz.numero(x.get("base_ref_valor")) or 0.0, fonte)
            self._q[x["id"]] = res
            return res
        if b.get("tipo") == "direta":
            q = nz.numero(b.get("valor")) or 0.0
            bs = nz.numero(x.get("base_ref_valor")) or 0.0
            res = (q, bs, b.get("fonte") or "quantidade direta")
        else:
            bs, fonte = self.base(x)
            ne = nz.numero(x.get("ne")) if x.get("ne") not in (None, "") else 1.0
            q = (arred if x.get("arred") == "arred" else lambda v, c=3: trunc(v, c))(bs * ne, 3)
            res = (q, bs, fonte)
        self._q[x["id"]] = res
        return res

    def preco_gd(self, cod: str, regime: str) -> tuple[float | None, str, list]:
        g = (self.d.get("desempenho") or {}).get(cod)
        if not g:
            return None, "grupo de desempenho sem serviços", []
        linhas, soma = [], 0.0
        for ln in g["linhas"]:
            if ln.get("undia") and self.sinalizacao():
                qkm = float(self.sinalizacao()["grupos"]["desempenho"]["codigos"].get(ln["codigo"], 0))
                tot = None
            elif ln.get("qtd_km") is not None and not ln.get("qtd_total"):
                qkm = nz.numero(ln["qtd_km"]) or 0.0
                tot = None
            else:
                qt = ln.get("qtd_total") or {}
                if qt.get("tipo") == "inventario" and self.medidas.get(qt.get("medida")) is not None:
                    tot = self.medidas[qt["medida"]] * (nz.numero(qt.get("fator")) or 1.0)
                else:
                    tot = nz.numero(qt.get("valor")) or 0.0
                qkm = trunc(tot * (nz.numero(ln.get("ne")) or 1.0) * self.freq(ln) / (self.ext or 1), 3)
            p, fonte = self.preco_fn(ln["codigo"], regime)
            v = qkm * (p or 0.0)
            soma += v
            linhas.append(dict(ln, qtd_total_calc=tot, qtd_km_calc=qkm, preco=p, preco_fonte=fonte, valor_km_ano=v))
        km_mes = trunc(soma / 12, 2)
        return trunc(km_mes * self.ext, 2), f"desempenho: R$ {km_mes:.2f}/km/mês × {self.ext:g} km", linhas

    def preco(self, x: dict, regime: str) -> tuple[float | None, str]:
        cod = x["codigo"]
        man = (x.get("preco_manual") or {}).get(regime)
        if man not in (None, ""):
            return nz.numero(man), "informado no item"
        if cod.startswith("GD"):
            p, f, _ = self.preco_gd(cod, regime)
            return p, f
        p, f = self.preco_fn(cod, regime)
        if p is None and (x.get("preco_ref") or {}).get(regime) is not None:
            return x["preco_ref"][regime], "preço da planilha de origem (sem composição na base)"
        return p, f

    def calcular(self) -> dict:
        out, tot = [], {"SD": 0.0, "CD": 0.0}
        secoes: dict[str, dict] = {}
        nomes = {x["item"]: (x.get("descricao") or "").strip() for x in self.d.get("itens", []) if x["tipo"] == "titulo"}
        for x in self.d.get("itens", []):
            if x["tipo"] == "titulo":
                out.append(dict(x))
                continue
            q, bs, fonte = self.quantidade(x)
            r = dict(x, quantidade=q, base_valor=bs, base_fonte=fonte)
            for reg in ("SD", "CD"):
                p, pf = self.preco(x, reg)
                r[f"preco_{reg}"] = p
                r[f"preco_fonte_{reg}"] = pf
                r[f"parcial_{reg}"] = trunc(q * p, 2) if p is not None else None
                tot[reg] += r[f"parcial_{reg}"] or 0.0
            sec = secao_do_item(x["item"])
            s = secoes.setdefault(sec, dict(secao=sec, nome=SECOES.get(sec) or nomes.get(sec) or sec, SD=0.0, CD=0.0, n=0))
            s["SD"] += r["parcial_SD"] or 0.0
            s["CD"] += r["parcial_CD"] or 0.0
            s["n"] += 1
            out.append(r)
        tot = {k: round(v, 2) for k, v in tot.items()}
        for s in secoes.values():
            s["SD"], s["CD"] = round(s["SD"], 2), round(s["CD"], 2)
            s["pct_SD"] = s["SD"] / tot["SD"] if tot["SD"] else 0
            s["pct_CD"] = s["CD"] / tot["CD"] if tot["CD"] else 0
        dsm = secoes.get("02", {})
        conserv = {r: tot[r] - (dsm.get(r) or 0) for r in ("SD", "CD")}
        avisos = []
        for reg in ("SD", "CD"):
            if conserv[reg] and (dsm.get(reg) or 0) > LIMITE_DSM * conserv[reg]:
                avisos.append(f"DSM acima de 25% da conservação ({REG_TXT[reg]}) — Res. 8/2023, art. 6º, III.")
        sem_preco = sorted({r["codigo"] for r in out if r["tipo"] == "servico" and (r.get("preco_SD") is None or r.get("preco_CD") is None)})
        return dict(itens=out, totais=tot, secoes=sorted(secoes.values(), key=lambda s: s["secao"]),
                    vantajosa="SD" if tot["SD"] <= tot["CD"] else "CD", conservacao=conserv, avisos=avisos,
                    sem_preco=sem_preco)


REG_TXT = {"SD": "sem desoneração", "CD": "com desoneração"}


def secao_do_item(item: str) -> str:
    if item.startswith("02"):
        return "02"
    return item[:5] if len(item) >= 5 else item


def curva_abc(itens: list[dict], regime: str, faixa_a: float = 0.80, faixa_b: float = 0.95) -> list[dict]:
    """Anexo XXI — agrupa por código e classifica: A até 80% acumulado, B até 95%, C o restante."""
    agr: dict[str, dict] = {}
    for r in itens:
        if r.get("tipo") != "servico" or not r.get(f"parcial_{regime}"):
            continue
        a = agr.setdefault(r["codigo"], dict(codigo=r["codigo"], descricao=r.get("descricao"), unidade=r.get("unidade"),
                                             quantidade=0.0, preco=r.get(f"preco_{regime}"), total=0.0,
                                             secao=secao_do_item(r["item"])))
        a["quantidade"] += r["quantidade"]
        a["total"] += r[f"parcial_{regime}"]
    lst = sorted(agr.values(), key=lambda a: -a["total"])
    soma = sum(a["total"] for a in lst) or 1
    acum = 0.0
    for a in lst:
        a["pct"] = a["total"] / soma
        acum += a["pct"]
        a["acumulado"] = acum
        a["faixa"] = "A" if acum - a["pct"] < faixa_a else "B" if acum - a["pct"] < faixa_b else "C"
    return lst


def cronograma(calc: dict, regime: str, meses: int) -> dict:
    """Cronograma físico-financeiro: desempenho, administração local e itens mensais em parcelas iguais; mobilização
    e canteiro no 1º mês; desmobilização no último; demais serviços distribuídos igualmente (a ajustar)."""
    meses = max(1, int(meses or 12))
    linhas, por_mes = [], [0.0] * meses
    for s in calc["secoes"]:
        v = s[regime]
        dist = [0.0] * meses
        if s["secao"] == "01.07":
            # itens da seção: ADM mensal; CAN/INSIND/MOB no 1º mês; DESMOB no último
            for r in calc["itens"]:
                if r.get("tipo") != "servico" or secao_do_item(r["item"]) != "01.07":
                    continue
                p = r.get(f"parcial_{regime}") or 0.0
                cod = r["codigo"].upper()
                if cod.startswith("DESMOB"):
                    dist[-1] += p
                elif cod.startswith(("MOB", "CAN", "INSIND")):
                    dist[0] += p
                else:
                    for i in range(meses):
                        dist[i] += p / meses
        else:
            for i in range(meses):
                dist[i] = v / meses
        for i in range(meses):
            por_mes[i] += dist[i]
        linhas.append(dict(secao=s["secao"], nome=s["nome"], total=v, meses=[round(x, 2) for x in dist]))
    acum, ac = [], 0.0
    tot = sum(por_mes) or 1
    for v in por_mes:
        ac += v
        acum.append(ac / tot)
    return dict(linhas=linhas, por_mes=[round(v, 2) for v in por_mes], acumulado=acum, meses=meses)
