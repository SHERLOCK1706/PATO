"""Itens do grupo 01.07 do PATO: administração local (ADM), canteiro de obras (CAN), instalação industrial (INSIND)
e mobilização/desmobilização (MOB/DESMOB) — SICRO, Manual de Custos Vol. 07 (canteiros), 08 (administração local)
e 09 (mobilização); Res. DNIT nº 8/2023. Replica as abas ADM_LOCAL, ADM_LOCAL_VARIAVEL, CANTEIRO, USINA e MOB do PATO,
com os preços de mão de obra e de equipamentos da base SICRO do projeto, o BDI do projeto e o prazo do contrato.

ADM local
  parcela fixa, vinculada e variável: MO mensalista (Σ TRUNC(q × salário; 2)) e veículos ((h prod × CHP + h improd × CHI) × q),
  multiplicados pelos meses (contrato; parcela variável = equipes de laboratório ELA = Σ ARRED(Q do serviço no contrato ÷ QE; 5));
  manutenção de canteiro (meses do contrato − 1); despesas diversas (%); preço = TRUNC(custo × (1 + BDI); 2); mensal = TRUNC(preço ÷ meses; 2).
Canteiro
  NMAX = TRUNC(NMO + funcionários das parcelas; 2); áreas: escritório 57,95 + 4,50·NPF, refeitório 1,55·50%·NMAX,
  alojamento 3,11·50%·(NMO + NPV), banheiros 0,77·(NMO + NPV) + áreas fixas; área equivalente = TRUNC(área × FEAC; 2);
  CCO = (k1·k2·k3·AC + AD·FEAD)·CMCC, k3 = 1 + 0,0008 × distância (km); PCO = TRUNC(CCO × (1 + BDI); 2).
Instalação industrial
  CCO = k1·k2·k3·AC·CMCC + custo direto da instalação da usina (composição SICRO 903810); PCO = TRUNC(CCO × (1 + BDI); 2).
Mobilização
  pessoal = ROUNDDOWN(Σ equipamentos × 50%) × passagem reajustada (TRUNC(passagem × índice data-base ÷ índice pesquisa; 2));
  equipamento: tempo de viagem K = DMT pav ÷ v pav + DMT npav ÷ v npav, descansos L = ROUNDDOWN(K ÷ 6), tempo de descanso
  M = L × 0,5 h; custo = TRUNC(q × FU × (K × CHP + M × CHI) × k; 0), com CHP/CHI do veículo transportador (tabela IMOB, Vol. 09).
  DESMOB = MOB.
"""
from __future__ import annotations

import json
import math
import re
from pathlib import Path

from . import normalizar as nz
from .pato_orcamento import arred, trunc

CODIGOS = {"ADM": "ADMINISTRAÇÃO LOCAL", "CAN": "INSTALAÇÃO DO CANTEIRO DE OBRAS E ACAMPAMENTOS",
           "INSIND": "INSTALAÇÃO INDUSTRIAL - USINA DE ASFALTO A QUENTE",
           "MOB": "MOBILIZAÇÃO DE PESSOAL E EQUIPAMENTOS DE GRANDE PORTE",
           "DESMOB": "DESMOBILIZAÇÃO DE PESSOAL E EQUIPAMENTOS DE GRANDE PORTE"}
REGIMES = ("SD", "CD")
INDICE_MOB = "Mobilização e Desmobilização"
_IMOB = None


def imob() -> dict:
    """Tabela do Vol. 09: equipamento → veículo de transporte rodoviário e fator de utilização (FU)."""
    global _IMOB
    if _IMOB is None:
        _IMOB = json.loads((Path(__file__).parent / "modelos" / "imob.json").read_text(encoding="utf-8"))["equipamentos"]
    return _IMOB


def _n(v, padrao=0.0) -> float:
    x = nz.numero(v)
    return padrao if x is None else x


# --------------------------------------------------------------- cálculo
def meses_parcela(regra, meses: float, ela: float) -> float:
    if regra in (None, "", "contrato"):
        return meses
    if regra == "contrato-1":
        return max(meses - 1, 0)
    if regra == "ela":
        return ela
    return _n(regra)


def calcular_ela(adm: dict, quantidade_fn) -> tuple[float, list[dict]]:
    linhas, tot = [], 0.0
    for e in adm.get("ela") or []:
        q = (quantidade_fn(e["codigo"]) if quantidade_fn else None)
        q = _n(e.get("quantidade_ref")) if q is None else q
        qe = _n(e.get("qe"))
        v = arred(q / qe, 5) if qe else 0.0
        tot += v
        linhas.append(dict(e, quantidade=q, ela=v))
    return round(tot, 5), linhas


def calcular_adm(adm: dict, mo_fn, eq_fn, bdi: dict, meses: float, quantidade_fn=None) -> dict:
    ela, ela_linhas = calcular_ela(adm, quantidade_fn)
    out = dict(ela=ela, ela_linhas=ela_linhas, parcelas=[], meses=meses)
    tot = {r: 0.0 for r in REGIMES}
    for p in adm.get("parcelas") or []:
        m = meses_parcela(p.get("meses"), meses, ela)
        lp = dict(chave=p.get("chave"), titulo=p.get("titulo"), meses=m, regra=p.get("meses"), mao_obra=[], veiculos=[])
        for reg in REGIMES:
            smo = sve = 0.0
            for i, x in enumerate(p.get("mao_obra") or []):
                pr = mo_fn(x["codigo"], reg)
                v = trunc(_n(x.get("quant")) * (pr or 0), 2)
                smo += v
                if reg == "SD":
                    lp["mao_obra"].append(dict(x, **{"preco_SD": pr, "total_SD": v}))
                else:
                    lp["mao_obra"][i].update(preco_CD=pr, total_CD=v)
            for i, x in enumerate(p.get("veiculos") or []):
                op, imp = eq_fn(x["codigo"], reg) or (None, None)
                v = (_n(x.get("h_prod")) * (op or 0) + _n(x.get("h_improd")) * (imp or 0)) * _n(x.get("quant"))
                sve += v
                if reg == "SD":
                    lp["veiculos"].append(dict(x, chp_SD=op, chi_SD=imp, total_SD=v))
                else:
                    lp["veiculos"][i].update(chp_CD=op, chi_CD=imp, total_CD=v)
            smo = round(smo, 2)
            lp[f"mo_mes_{reg}"], lp[f"veic_mes_{reg}"] = smo, sve
            lp[f"total_{reg}"] = trunc(smo * m, 2) + (trunc(sve * m, 2) if p.get("veiculos") else 0.0)
            tot[reg] += lp[f"total_{reg}"]
        out["parcelas"].append(lp)
    pct = _n(adm.get("despesas_pct"), 5.0)
    for reg in REGIMES:
        sub = round(tot[reg], 2)
        desp = trunc(sub * pct / 100, 2)
        custo = round(sub + desp, 2)
        preco = trunc(custo * (1 + bdi[reg]), 2)
        out[reg] = dict(subtotal=sub, despesas=desp, custo=custo, bdi=bdi[reg], preco=preco,
                        mensal=trunc(preco / meses, 2) if meses else None)
    out["despesas_pct"] = pct
    return out


def funcionarios(adm: dict) -> dict:
    s = {}
    for p in adm.get("parcelas") or []:
        s[p.get("chave")] = s.get(p.get("chave"), 0.0) + sum(_n(x.get("quant")) for x in p.get("mao_obra") or [])
    return dict(fixa=s.get("fixa", 0.0), vinculada=s.get("vinculada", 0.0), variavel=s.get("variavel", 0.0))


def k3(c: dict) -> float:
    return 1 + _n(c.get("k3_coef"), 0.0008) * _n(c.get("distancia_km"))


def calcular_canteiro(c: dict, adm: dict, bdi: dict) -> dict:
    f = funcionarios(adm)
    nmo = _n(c.get("nmo"))
    npf, npv = f["fixa"], f["variavel"]
    nmax = trunc(nmo + f["fixa"] + f["vinculada"] + f["variavel"], 2)
    areas = [dict(nome="Escritório e seção técnica", formula="57,95 + 4,50 × NPF", area=57.95 + 4.5 * npf,
                  feac=_n(c.get("feac_escritorio"), 0.7)),
             dict(nome="Refeitório e cozinha", formula="1,55 × (50% × NMAX)", area=1.55 * (0.5 * nmax),
                  feac=_n(c.get("feac_refeitorio"), 0.7)),
             dict(nome="Alojamentos", formula="3,11 × 50% × (NMO + NPV)", area=3.11 * 0.5 * (nmo + npv),
                  feac=_n(c.get("feac_alojamento"), 0.7)),
             dict(nome="Banheiros e vestiário", formula="0,77 × (NMO + NPV)", area=0.77 * (nmo + npv),
                  feac=_n(c.get("feac_banheiros"), 0.7))]
    areas += [dict(x, formula="área de referência", area=_n(x.get("area")), feac=_n(x.get("feac"), 0.5))
              for x in c.get("areas_fixas") or []]
    for a in areas:
        a["equivalente"] = trunc(a["area"] * a["feac"], 2)
    ac = round(sum(a["equivalente"] for a in areas), 2)
    ad = ac if c.get("ad") in (None, "") else _n(c.get("ad"))
    kk = (_n(c.get("k1"), 0.8), _n(c.get("k2"), 1.11), k3(c))
    out = dict(funcionarios=dict(f, nmo=nmo, nmax=nmax), areas=areas, area_total=round(sum(a["area"] for a in areas), 4),
               ac=ac, ad=ad, k1=kk[0], k2=kk[1], k3=kk[2], fead=_n(c.get("fead"), 0.05))
    for reg in REGIMES:
        cmcc = _n((c.get("cmcc") or {}).get(reg))
        cco = (kk[0] * kk[1] * kk[2] * ac + ad * out["fead"]) * cmcc
        out[reg] = dict(cmcc=cmcc, custo=cco, bdi=bdi[reg], preco=trunc(cco * (1 + bdi[reg]), 2))
    return out


def calcular_usina(u: dict, can: dict, custo_comp_fn, bdi: dict) -> dict:
    areas = [dict(x, area=_n(x.get("area")), feac=_n(x.get("feac"), 0.6)) for x in u.get("areas") or []]
    for a in areas:
        a["equivalente"] = trunc(a["area"] * a["feac"], 2)
    ac = round(sum(a["equivalente"] for a in areas), 2)
    kk = (_n(can.get("k1"), 0.8), _n(can.get("k2"), 1.11), k3(can))
    cod = u.get("codigo_instalacao")
    out = dict(areas=areas, ac=ac, k1=kk[0], k2=kk[1], k3=kk[2], codigo_instalacao=cod)
    for reg in REGIMES:
        cmcc = _n((can.get("cmcc") or {}).get(reg))
        inst = (custo_comp_fn(cod, reg) if cod and custo_comp_fn else None)
        inst = _n((u.get("instalacao_ref") or {}).get(reg)) if inst is None else inst
        cco = kk[0] * kk[1] * kk[2] * (ac * cmcc) + inst
        out[reg] = dict(cmcc=cmcc, instalacao=inst, custo=cco, bdi=bdi[reg], preco=trunc(cco * (1 + bdi[reg]), 2))
    return out


def fator_passagem(m: dict, indice_fn, data_base: str | None) -> tuple[float, dict]:
    nome = m.get("indice") or INDICE_MOB
    a = indice_fn(nome, m.get("mes_indice_pesquisa")) if indice_fn and m.get("mes_indice_pesquisa") else None
    b = indice_fn(nome, data_base) if indice_fn and data_base else None
    if a and b:
        return b / a - 1, dict(indice=nome, a=a, b=b, mes_a=m.get("mes_indice_pesquisa"), mes_b=data_base)
    return _n(m.get("fator_ref")), dict(indice=nome, a=None, b=None, mes_a=m.get("mes_indice_pesquisa"), mes_b=data_base)


def calcular_mob(m: dict, eq_fn, bdi: dict, indice_fn=None, data_base: str | None = None) -> dict:
    tab = imob()
    fator, ind = fator_passagem(m, indice_fn, data_base)
    passagem = trunc(_n(m.get("passagem")) * (1 + fator), 2)
    eqs = [x for x in m.get("equipamentos") or []]
    n_pessoal = math.floor(sum(_n(x.get("quant")) for x in eqs) * _n(m.get("pessoal_pct"), 0.5))
    hd = _n(m.get("horas_entre_descansos"), 6.0)
    td = _n(m.get("tempo_descanso"), 0.5)
    linhas, soma = [], {r: 0.0 for r in REGIMES}
    for x in eqs:
        info = tab.get(x["codigo"]) or {}
        tr = x.get("transporte") or info.get("transporte")
        fu = _n(x.get("fu"), _n(info.get("fu")))
        vp, vn = _n(x.get("vel_pav"), 60), _n(x.get("vel_npav"), 50)
        k = (_n(x.get("dmt_pav")) / vp if vp else 0) + (_n(x.get("dmt_npav")) / vn if vn else 0)
        nd = math.floor(k / hd) if hd else 0
        mm = nd * td
        ln = dict(x, transporte=tr, fu=fu, tempo_viagem=k, descansos=nd, tempo_descanso=mm,
                  descricao=x.get("descricao") or info.get("descricao"))
        for reg in REGIMES:
            op, imp = (eq_fn(tr, reg) if tr else None) or (None, None)
            v = math.trunc(float(f"{_n(x.get('quant')) * fu * ((k * (op or 0) + mm * (imp or 0)) * _n(x.get('k'), 1)):.15g}"))
            ln[f"chp_{reg}"], ln[f"chi_{reg}"], ln[f"custo_{reg}"] = op, imp, float(v)
            soma[reg] += v
        linhas.append(ln)
    out = dict(equipamentos=linhas, pessoal=dict(quantidade=n_pessoal, passagem=passagem, fator=fator, **ind,
                                                  itinerario=m.get("itinerario"), dmt=m.get("dmt_pessoal")))
    for reg in REGIMES:
        pes = trunc(n_pessoal * passagem, 2)
        custo = round(pes + soma[reg], 2)
        out[reg] = dict(pessoal=pes, equipamentos=soma[reg], custo=custo, bdi=bdi[reg], preco=trunc(custo * (1 + bdi[reg]), 2))
    return out


def calcular(d: dict, *, mo_fn, eq_fn, custo_comp_fn, bdi: dict, meses: float, quantidade_fn=None,
             indice_fn=None, data_base: str | None = None) -> dict:
    adm = calcular_adm(d.get("adm") or {}, mo_fn, eq_fn, bdi, meses, quantidade_fn)
    can = calcular_canteiro(d.get("canteiro") or {}, d.get("adm") or {}, bdi)
    usi = calcular_usina(d.get("usina") or {}, d.get("canteiro") or {}, custo_comp_fn, bdi)
    mob = calcular_mob(d.get("mob") or {}, eq_fn, bdi, indice_fn, data_base)
    precos = {}
    for reg in REGIMES:
        precos.setdefault("ADM", {})[reg] = adm[reg]["mensal"]
        precos.setdefault("CAN", {})[reg] = can[reg]["preco"]
        precos.setdefault("INSIND", {})[reg] = usi[reg]["preco"]
        precos.setdefault("MOB", {})[reg] = mob[reg]["preco"]
        precos.setdefault("DESMOB", {})[reg] = mob[reg]["preco"]
    return dict(adm=adm, canteiro=can, usina=usi, mob=mob, precos=precos)


def sugerir_equipamentos_mob(codigos_eq: dict[str, float]) -> list[dict]:
    """A partir dos equipamentos das composições do PATO ({código: maior quantidade}), os que exigem transporte rodoviário."""
    tab = imob()
    out = []
    for cod in sorted(codigos_eq):
        info = tab.get(cod)
        if info and info.get("transporte"):
            out.append(dict(codigo=cod, descricao=info.get("descricao"), quant=max(1, math.ceil(codigos_eq[cod] - 1e-9)),
                            transporte=info["transporte"], fu=info.get("fu")))
    return out


# ------------------------------------------------------------- importação
def _linhas(wf, wv, aba, ncol=24):
    fr = [tuple(r) + (None,) * ncol for r in wf[aba].iter_rows(values_only=True)]
    vr = [tuple(r) + (None,) * ncol for r in wv[aba].iter_rows(values_only=True)]
    return fr, vr


def _col(letra: str) -> int:
    return ord(letra) - 65


def _imp_adm(wf, wv) -> dict:
    fr, vr = _linhas(wf, wv, "ADM_LOCAL")
    secs = {1: "fixa", 2: "vinculada", 3: "variavel", 4: "manutencao"}
    titulos = {}
    parcelas: dict[int, dict] = {}
    sec, meses_ref = None, {}
    for f, v in zip(fr, vr):
        b = nz.texto(v[1]) or ""
        m = re.match(r"^(\d)\s*-\s*(.+)", b)
        if m:
            sec = int(m.group(1))
            titulos[sec] = m.group(2).strip()
            continue
        cod = nz.texto(v[2])
        if sec in secs and cod and re.match(r"^[PE]\d", cod.strip()):
            p = parcelas.setdefault(sec, dict(chave=secs[sec], titulo=titulos.get(sec), mao_obra=[], veiculos=[]))
            cod = cod.strip()
            if cod.startswith("P"):
                p["mao_obra"].append(dict(codigo=cod, descricao=nz.texto(v[3]), quant=_n(v[5])))
            else:
                p["veiculos"].append(dict(codigo=cod, descricao=nz.texto(v[3]), quant=_n(v[5]), h_prod=_n(v[6]), h_improd=_n(v[7])))
        if sec == 5:  # resumo: meses de cada parcela
            it = nz.texto(v[1]) or ""
            mm = re.match(r"^5\.(\d)\.\d", it)
            if mm and v[5] is not None:
                meses_ref.setdefault(int(mm.group(1)), (f[5], _n(v[5])))
        if sec == 6 and "DESPESAS" in nz.sem_acento(nz.texto(v[3]) or "").upper():
            despesas = _n(v[5], 5.0)
            parcelas.setdefault(0, {})["despesas"] = despesas
    desp = (parcelas.pop(0, None) or {}).get("despesas", 5.0)
    fixa_m = (meses_ref.get(1) or (None, 24.0))[1]
    out = []
    for k in sorted(parcelas):
        p = parcelas[k]
        f, val = meses_ref.get(k, (None, fixa_m))
        if isinstance(f, str) and "VARIAVEL" in f.upper():
            p["meses"] = "ela"
        elif abs(val - fixa_m) < 1e-9:
            p["meses"] = "contrato"
        elif abs(val - (fixa_m - 1)) < 1e-9:
            p["meses"] = "contrato-1"
        else:
            p["meses"] = val
        p["meses_ref"] = val
        out.append(p)
    ela = []
    if "ADM_LOCAL_VARIAVEL" in wv.sheetnames:
        fr2, vr2 = _linhas(wf, wv, "ADM_LOCAL_VARIAVEL")
        for f, v in zip(fr2, vr2):
            cod = nz.texto(v[2])
            if cod and re.match(r"^\d{5,}", cod.strip()) and _n(v[6]):
                ela.append(dict(codigo=cod.strip(), descricao=nz.texto(v[3]), unidade=nz.texto(v[4]), qe=_n(v[6]),
                                quantidade_ref=_n(v[5])))
    return dict(parcelas=out, despesas_pct=desp, ela=ela, meses_ref=fixa_m)


def _rotulo(v) -> str:
    return nz.sem_acento(nz.texto(v) or "").upper().strip()


def _imp_canteiro(wf, wv) -> dict:
    fr, vr = _linhas(wf, wv, "CANTEIRO")
    c = dict(areas_fixas=[], cmcc={})
    feac, em_fixas = {}, False
    for f, v in zip(fr, vr):
        r = _rotulo(v[2])
        if "AREAS FIXAS" in r:
            em_fixas = True
            continue
        if em_fixas:
            if r and nz.numero(v[7]) is not None:
                c["areas_fixas"].append(dict(nome=nz.texto(v[2]).strip(), area=_n(v[7])))
                continue
            if _rotulo(v[6]).startswith("AREA TOTAL") or not r and nz.numero(v[7]) is not None:
                em_fixas = False
        if r and nz.numero(v[5]) is not None and nz.numero(v[4]) is not None and _rotulo(v[3]) in ("M²", "M2"):
            feac[r] = _n(v[5])
        if "MAO DE OBRA ORDINARIA" in r:
            c["nmo"] = _n(v[7])
        if r.startswith("K1"):
            c["k1"] = _n(v[6])
        elif r.startswith("K2"):
            c["k2"] = _n(v[6])
        elif r.startswith("K3"):
            m = re.match(r"=\s*1\s*\+\s*([\d.]+)\s*\*\s*([\d.]+)", str(f[6] or ""))
            if m:
                c["k3_coef"], c["distancia_km"] = float(m.group(1)), float(m.group(2))
            else:
                c["k3_coef"], c["distancia_km"] = 0.0008, (_n(v[6], 1) - 1) / 0.0008
        elif r.startswith("FEAD"):
            c["fead"] = _n(v[6])
        elif r.startswith("CMCC SEM"):
            c["cmcc"]["SD"] = _n(v[6])
        elif r.startswith("CMCC COM"):
            c["cmcc"]["CD"] = _n(v[6])
    for nome, chave in (("ESCRITORIO", "feac_escritorio"), ("REFEITORIO", "feac_refeitorio"),
                        ("ALOJAMENTO", "feac_alojamento"), ("BANHEIRO", "feac_banheiros")):
        x = next((val for k, val in feac.items() if k.startswith(nome)), None)
        if x is not None:
            c[chave] = x
    for a in c["areas_fixas"]:
        x = feac.get(_rotulo(a["nome"]))
        if x is not None:
            a["feac"] = x
    return c


def _imp_usina(wf, wv) -> dict:
    fr, vr = _linhas(wf, wv, "USINA")
    u = dict(areas=[], instalacao_ref={})
    cob = False
    for f, v in zip(fr, vr):
        r = _rotulo(v[2])
        if "AREAS COBERTAS" in r:
            cob = True
            continue
        if cob:
            if r and nz.numero(v[4]) is not None and nz.numero(v[5]) is not None:
                u["areas"].append(dict(nome=nz.texto(v[2]).strip(), area=_n(v[4]), feac=_n(v[5])))
                continue
            if _rotulo(v[3]) == "TOTAL":
                cob = False
        cod = nz.texto(v[1])
        if cod and re.match(r"^\d{6,}$", cod.strip()):
            u["codigo_instalacao"] = cod.strip()
            u["descricao_instalacao"] = nz.texto(v[2])
            u["instalacao_ref"] = dict(SD=_n(v[4]), CD=_n(v[6]))
    return u


def _imp_mob(wf, wv) -> dict:
    fr, vr = _linhas(wf, wv, "MOB")
    m = dict(equipamentos=[], indice=INDICE_MOB)
    for f, v in zip(fr, vr):
        cod = nz.texto(v[1])
        if "MOTORISTAS" in _rotulo(v[2]):
            m["dmt_pessoal"] = _n(v[6])
            m["itinerario"] = nz.texto(v[8])
            mm = re.search(r"\*\s*([\d.]+)%", str(f[5] or ""))
            if mm:
                m["pessoal_pct"] = float(mm.group(1)) / 100
        if cod and re.match(r"^E\d{4}", cod.strip()) and nz.numero(v[5]) is not None:
            fl = str(f[11] or "")
            md = re.search(r"/\s*([\d.]+)", fl)
            if md:
                m["horas_entre_descansos"] = float(md.group(1))
            mt = re.search(r"\*\s*([\d.]+)", str(f[12] or ""))
            if mt:
                m["tempo_descanso"] = float(mt.group(1))
            m["equipamentos"].append(dict(codigo=cod.strip(), descricao=nz.texto(v[2]), transporte=nz.texto(v[3]),
                                          quant=_n(v[5]), dmt_pav=_n(v[6]), dmt_npav=_n(v[7]), vel_pav=_n(v[8], 60),
                                          vel_npav=_n(v[9], 50), k=_n(v[13], 1), fu=nz.numero(v[14])))
    # passagem: bloco "ÍNDICE DE REAJ." (a = mês do índice da pesquisa, b = data-base) e "Valor da passagem"
    from .binomio import mes
    for i, v in enumerate(vr):
        for j, x in enumerate(v):
            r = _rotulo(x)
            if r.startswith("VALOR DA PASSAGEM:"):
                m["passagem"] = next((nz.numero(y) for y in v[j + 1:] if nz.numero(y) is not None), None)
            elif r.startswith("DATA DA PE"):
                y = next((y for y in v[j + 1:] if y is not None), None)
                m["data_pesquisa"] = mes(y)
            elif r.startswith("VALOR DA PASSAGEM REAJ"):
                m["passagem_reaj_ref"] = next((nz.numero(y) for y in v[j + 1:] if nz.numero(y) is not None), None)
            elif r == "A" and i + 1 < len(vr):
                m["mes_indice_pesquisa"] = mes(vr[i + 1][j])
            elif r.startswith("R=A/B") and i + 2 < len(vr):
                m["fator_ref"] = nz.numero(vr[i + 2][j])
    return m


def importar_planilha(caminho: str) -> dict | None:
    import openpyxl
    wv = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    wf = openpyxl.load_workbook(caminho, read_only=True)
    try:
        if "ADM_LOCAL" not in wv.sheetnames:
            return None
        d = dict(adm=_imp_adm(wf, wv))
        if "CANTEIRO" in wv.sheetnames:
            d["canteiro"] = _imp_canteiro(wf, wv)
        if "USINA" in wv.sheetnames:
            d["usina"] = _imp_usina(wf, wv)
        if "MOB" in wv.sheetnames:
            d["mob"] = _imp_mob(wf, wv)
        return d
    finally:
        wv.close()
        wf.close()


# ======================================================================================================================
# SICRO — Manual de Custos de Infraestrutura de Transportes, 2ª edição (2025): Vol. 06 (canteiro), 07 (administração
# local) e 08 (mobilização), conforme os modelos ADM_Local_2025, CANT_Construcao_2025, USINA_2025 e MOB_2025.
# ======================================================================================================================
MODELOS = {"sicro2025": "SICRO — Manual de Custos 2ª edição (2025)", "br158": "Planilha de origem (PATO BR-158/MT)"}
H_MES = 182.49  # horas por mês (Tf do manejo florestal)
K3 = dict(pav=0.0008, rp=0.0009, leito=0.0014)  # Vol. 06, Tabela 36
VEL = dict(caminhao=dict(pav=60.0, rp=50.0, leito=40.0), cavalo=dict(pav=30.0, rp=25.0, leito=20.0))
MANEJO_PADRAO = ["5501700", "5501701", "4915768", "4915764", "4915765", "4915766", "4915767"]


def _mo(cod, q, desc=None):
    return dict(codigo=cod, quant=q, descricao=desc)


def _ve(cod, q, desc=None):
    return dict(codigo=cod, quant=q, h_prod=44.0, h_improd=138.49, descricao=desc)


def padrao_2025(base: dict | None = None) -> dict:
    """Parâmetros da 2ª edição. Aproveita da planilha de origem o que é do projeto: equipes de laboratório (QE), CMCC,
    distância aos centros fornecedores, mão de obra de pico, lista e DMT dos equipamentos e a passagem."""
    b = base or {}
    adm0, can0, mob0 = b.get("adm") or {}, b.get("canteiro") or {}, b.get("mob") or {}
    adm = dict(despesas_pct=5.0, casas_ela=2, ela=adm0.get("ela") or [],
               manejo=[dict(codigo=c) for c in MANEJO_PADRAO], horas_mes=H_MES, parcelas=[
        dict(chave="fixa", titulo="Parcela fixa", meses="contrato",
             mao_obra=[_mo("P9819", 0.25, "Engenheiro supervisor"), _mo("P9958", 1, "Motorista de ônibus"),
                       _mo("P9806", 1, "Auxiliar administrativo"), _mo("P9803", 1, "Almoxarife"), _mo("P9842", 1, "Faxineiro")],
             veiculos=[_ve("E9093", 0.25, "Veículo leve (engenheiro)"), _ve("E9239", 1, "Miniônibus 30 passageiros"),
                       _ve("E9093", 1, "Veículo leve (gerência administrativa)")]),
        dict(chave="vinculada", titulo="Parcela vinculada", meses="contrato",
             mao_obra=[_mo("P9916", 1, "Encarregado de conservação rodoviária"), _mo("P9875", 2, "Encarregado de turma")],
             veiculos=[_ve("E9093", 1, "Veículo leve")]),
        dict(chave="variavel", titulo="Parcela variável — laboratório de asfalto", meses="ela",
             mao_obra=[_mo("P9858", 1, "Laboratorista"), _mo("P9833", 2, "Auxiliar de laboratório")],
             veiculos=[_ve("E9125", 1, "Van furgão")]),
        dict(chave="manejo", titulo="Parcela variável — manejo florestal", meses="manejo",
             mao_obra=[_mo("P9947", 1, "Técnico florestal")], veiculos=[]),
        dict(chave="manutencao", titulo="Manutenção do canteiro de obras e acampamentos", meses="contrato-1",
             mao_obra=[_mo("P9952", 0.1, "Pedreiro"), _mo("P9954", 0.1, "Servente"), _mo("P9953", 0.1, "Eletricista")],
             veiculos=[], veiculos_horas=[dict(codigo="E9686", horas=5.5, descricao="Caminhão guindauto"),
                                          dict(codigo="E9669", horas=11.0, descricao="Caminhão tanque 8.000 l"),
                                          dict(codigo="E9524", horas=5.5, descricao="Motoniveladora")]),
    ])
    nmo = _n(can0.get("nmo"))
    can = dict(mo_pico=round(nmo / 1.33, 4) if nmo else 15.0, k1=0.8, k2=1.11, fead=0.05,
               dt=dict(pav=_n(can0.get("distancia_km"), 50.0), rp=0.0, leito=0.0), cmcc=dict(can0.get("cmcc") or {}),
               areas_fixas=[dict(nome="Almoxarifado", area=29.54, feac=0.5), dict(nome="Oficina", area=37.10, feac=0.5),
                            dict(nome="Guarita", area=5.29, feac=0.7)])
    usi = dict(codigo_instalacao=(b.get("usina") or {}).get("codigo_instalacao") or "903810",
               areas=[dict(nome="Laboratório", area=94.36, feac=0.6)], k1=0.8, k2=1.11, fead=0.05)
    eqs = []
    for e in mob0.get("equipamentos") or []:
        auto = (e.get("transporte") or "") == e.get("codigo")
        eqs.append(dict(codigo=e["codigo"], descricao=e.get("descricao"), quant=e.get("quant"), transporte=e.get("transporte"),
                        fu=e.get("fu"), dmt=dict(pav=_n(e.get("dmt_pav")), rp=_n(e.get("dmt_npav")), leito=0.0),
                        k=1.0 if auto else 2.0, pedagio_eixos=0, pedagio_valor=0.0))
    mob = dict(passagem=mob0.get("passagem"), mes_indice_pesquisa=mob0.get("mes_indice_pesquisa"), fator_ref=mob0.get("fator_ref"),
               indice=mob0.get("indice") or INDICE_MOB, itinerario=mob0.get("itinerario"), dmt_pessoal=mob0.get("dmt_pessoal"),
               descanso_carga_h=6.0, descanso_passageiros_h=4.0, tempo_descanso=0.5, equipamentos=eqs)
    return dict(adm=adm, canteiro=can, usina=usi, mob=mob)


def com_2025(d: dict) -> dict:
    """Dados do projeto com a metodologia da 2ª edição como padrão (mantém os da planilha de origem para a opção 'br158')."""
    d = dict(d or {})
    d.setdefault("modelo", "sicro2025")
    if not d.get("sicro2025"):
        d["sicro2025"] = padrao_2025(d)
    return d


def _tf(adm: dict, quantidade_fn, producao_fn) -> tuple[float, list[dict]]:
    linhas, tot = [], 0.0
    hm = _n(adm.get("horas_mes"), H_MES)
    for e in adm.get("manejo") or []:
        q = quantidade_fn(e["codigo"]) if quantidade_fn else None
        prod = _n(e.get("producao")) or (producao_fn(e["codigo"]) if producao_fn else None)
        if not q or not prod:
            linhas.append(dict(e, quantidade=q, producao=prod, tf=0.0))
            continue
        v = arred(q / (prod * hm), 2)
        tot += v
        linhas.append(dict(e, quantidade=q, producao=prod, tf=v))
    return round(tot, 2), linhas


def calcular_adm_2025(adm: dict, mo_fn, eq_fn, bdi: dict, meses: float, quantidade_fn=None, producao_fn=None) -> dict:
    casas = int(_n(adm.get("casas_ela"), 2))
    ela_l, ela = [], 0.0
    for e in adm.get("ela") or []:
        q = quantidade_fn(e["codigo"]) if quantidade_fn else None
        q = _n(e.get("quantidade_ref")) if q is None else q
        v = arred(q / _n(e.get("qe")), casas) if _n(e.get("qe")) else 0.0
        ela += v
        ela_l.append(dict(e, quantidade=q, ela=v))
    ela = round(ela, casas)
    tf, tf_l = _tf(adm, quantidade_fn, producao_fn)
    out = dict(ela=ela, ela_linhas=ela_l, tf=tf, tf_linhas=tf_l, parcelas=[], meses=meses, metodologia="sicro2025")
    tot = {r: 0.0 for r in REGIMES}
    for p in adm.get("parcelas") or []:
        regra = p.get("meses")
        m = tf if regra == "manejo" else meses_parcela(regra, meses, ela)
        lp = dict(chave=p.get("chave"), titulo=p.get("titulo"), meses=m, regra=regra, mao_obra=[], veiculos=[], veiculos_horas=[])
        for reg in REGIMES:
            smo = sve = 0.0
            for i, x in enumerate(p.get("mao_obra") or []):
                pr = mo_fn(x["codigo"], reg)
                v = trunc(_n(x.get("quant")) * (pr or 0), 2)
                smo += v
                (lp["mao_obra"].append(dict(x, preco_SD=pr, total_SD=v)) if reg == "SD" else lp["mao_obra"][i].update(preco_CD=pr, total_CD=v))
            for i, x in enumerate(p.get("veiculos") or []):
                op, imp = eq_fn(x["codigo"], reg) or (None, None)
                v = (_n(x.get("h_prod")) * (op or 0) + _n(x.get("h_improd")) * (imp or 0)) * _n(x.get("quant"))
                sve += v
                (lp["veiculos"].append(dict(x, chp_SD=op, chi_SD=imp, total_SD=v)) if reg == "SD" else lp["veiculos"][i].update(chp_CD=op, chi_CD=imp, total_CD=v))
            for i, x in enumerate(p.get("veiculos_horas") or []):
                op, _imp = eq_fn(x["codigo"], reg) or (None, None)
                v = (op or 0) * _n(x.get("horas"))  # equipamento por hora produtiva no mês (Tabela 48)
                sve += v
                (lp["veiculos_horas"].append(dict(x, chp_SD=op, total_SD=v)) if reg == "SD" else lp["veiculos_horas"][i].update(chp_CD=op, total_CD=v))
            lp[f"mo_mes_{reg}"], lp[f"veic_mes_{reg}"] = round(smo, 2), sve
            lp[f"total_{reg}"] = trunc((round(smo, 2) + sve) * m, 2)  # mão de obra + veículos × meses (resumo do modelo)
            tot[reg] += lp[f"total_{reg}"]
        out["parcelas"].append(lp)
    pct = _n(adm.get("despesas_pct"), 5.0)
    for reg in REGIMES:
        sub = round(tot[reg], 2)
        desp = trunc(sub * pct / 100, 2)
        custo = round(sub + desp, 2)
        preco = trunc(custo * (1 + bdi[reg]), 2)
        out[reg] = dict(subtotal=sub, despesas=desp, custo=custo, bdi=bdi[reg], preco=preco, mensal=trunc(preco / meses, 2) if meses else None)
    out["despesas_pct"] = pct
    return out


def _func_2025(adm: dict) -> dict:
    s = {}
    for p in adm.get("parcelas") or []:
        k = "variavel" if p.get("chave") in ("variavel", "manejo") else p.get("chave")
        s[k] = s.get(k, 0.0) + sum(_n(x.get("quant")) for x in p.get("mao_obra") or [])
    return dict(fixa=s.get("fixa", 0.0), vinculada=s.get("vinculada", 0.0), variavel=s.get("variavel", 0.0))


def k3_2025(dt: dict) -> float:
    return 1 + sum(K3[k] * _n((dt or {}).get(k)) for k in K3)


def calcular_canteiro_2025(c: dict, adm: dict, bdi: dict) -> dict:
    f = _func_2025(adm)
    nmo = 1.33 * _n(c.get("mo_pico"))
    npf, npv = f["fixa"], f["variavel"]
    nmax = trunc(nmo + f["fixa"] + f["vinculada"] + f["variavel"], 2)
    areas = [dict(nome="Escritório e seção técnica", formula="57,95 + 4,50 × NPF", area=57.95 + 4.5 * npf, feac=0.7),
             dict(nome="Refeitório e cozinha", formula="1,55 × (50% × NMAX)", area=1.55 * 0.5 * nmax, feac=0.7),
             dict(nome="Alojamentos", formula="3,11 × 50% × (NMO + NPV)", area=3.11 * 0.5 * (nmo + npv), feac=0.7),
             dict(nome="Banheiros e vestiário", formula="0,77 × (NMO + NPV)", area=0.77 * (nmo + npv), feac=0.7)]
    areas += [dict(x, formula="área de referência", area=_n(x.get("area")), feac=_n(x.get("feac"), 0.5)) for x in c.get("areas_fixas") or []]
    for a in areas:
        a["equivalente"] = trunc(a["area"] * a["feac"], 2)
    ac = round(sum(a["equivalente"] for a in areas), 2)
    kk = (_n(c.get("k1"), 0.8), _n(c.get("k2"), 1.11), k3_2025(c.get("dt")))
    out = dict(funcionarios=dict(f, nmo=nmo, nmax=nmax, mo_pico=_n(c.get("mo_pico"))), areas=areas, ac=ac, ad=ac, k1=kk[0], k2=kk[1],
               k3=kk[2], fead=_n(c.get("fead"), 0.05), dt=c.get("dt"), metodologia="sicro2025")
    for reg in REGIMES:
        cmcc = _n((c.get("cmcc") or {}).get(reg))
        cco = (kk[0] * kk[1] * kk[2] * ac + ac * out["fead"]) * cmcc
        out[reg] = dict(cmcc=cmcc, custo=cco, bdi=bdi[reg], preco=trunc(cco * (1 + bdi[reg]), 2))
    return out


def calcular_usina_2025(u: dict, can: dict, custo_comp_fn, bdi: dict) -> dict:
    areas = [dict(x, area=_n(x.get("area")), feac=_n(x.get("feac"), 0.6)) for x in u.get("areas") or []]
    for a in areas:
        a["equivalente"] = trunc(a["area"] * a["feac"], 2)
    ac = round(sum(a["equivalente"] for a in areas), 2)
    kk = (_n(u.get("k1"), 0.8), _n(u.get("k2"), 1.11), k3_2025(can.get("dt")))
    cod = u.get("codigo_instalacao")
    out = dict(areas=areas, ac=ac, ad=ac, k1=kk[0], k2=kk[1], k3=kk[2], fead=_n(u.get("fead"), 0.05), codigo_instalacao=cod,
               metodologia="sicro2025")
    for reg in REGIMES:
        cmcc = _n((can.get("cmcc") or {}).get(reg))
        inst = custo_comp_fn(cod, reg) if cod and custo_comp_fn else None
        inst = inst or 0.0
        cco = (kk[0] * kk[1] * kk[2] * ac + ac * out["fead"]) * cmcc + inst
        out[reg] = dict(cmcc=cmcc, instalacao=inst, custo=cco, bdi=bdi[reg], preco=trunc(cco * (1 + bdi[reg]), 2))
    return out


def _tipo_transportador(desc: str | None) -> str:
    t = nz.sem_acento(desc or "").upper()
    return "cavalo" if re.search(r"CAVALO|DOLLY|SEMIRREBOQUE|PRANCHA|REBOQUE", t) else "caminhao"


def calcular_mob_2025(m: dict, adm: dict, can: dict, eq_fn, desc_fn, bdi: dict, indice_fn=None, data_base=None) -> dict:
    fator, ind = fator_passagem(m, indice_fn, data_base)
    passagem = trunc(_n(m.get("passagem")) * (1 + fator), 2)
    f = _func_2025(adm)
    nmo = 1.33 * _n(can.get("mo_pico"))
    nadm = f["fixa"] + f["vinculada"] + 0.5 * f["variavel"]  # NADM = NPF-V + 50% × NPV
    ntf = 0.5 * nmo + nadm  # NTF = 50% × NMO + NADM
    n_pes = math.ceil(ntf - 1e-9) if m.get("pessoal_qtd") in (None, "") else _n(m.get("pessoal_qtd"))
    td_c, td_h = _n(m.get("descanso_carga_h"), 6.0), _n(m.get("tempo_descanso"), 0.5)
    linhas, soma = [], {r: 0.0 for r in REGIMES}
    for x in m.get("equipamentos") or []:
        tr = x.get("transporte")
        tipo = _tipo_transportador(desc_fn(tr) if desc_fn and tr else "")
        dmt = x.get("dmt") or {}
        tv = sum(_n(dmt.get(k)) / VEL[tipo][k] for k in ("pav", "rp", "leito"))
        td = math.floor(tv / td_c) * td_h if td_c else 0.0
        cc = _n(x.get("pedagio_eixos")) * _n(x.get("pedagio_valor"))
        fu, q, k = _n(x.get("fu"), 1.0), _n(x.get("quant")), _n(x.get("k"), 1.0)
        ln = dict(x, tipo_transportador=tipo, tempo_viagem=tv, tempo_descanso=td, pedagio=cc, dmt_total=sum(_n(dmt.get(k_)) for k_ in dmt))
        for reg in REGIMES:
            op, imp = (eq_fn(tr, reg) if tr else None) or (None, None)
            v = q * fu * (((tv * (op or 0)) + td * (imp or 0)) * k + cc)  # CMD = QT × FU × {[(TV×CHP + TD×CHI) × K] + CC}
            ln[f"chp_{reg}"], ln[f"chi_{reg}"], ln[f"custo_{reg}"] = op, imp, round(v, 2)
            soma[reg] += v
        linhas.append(ln)
    out = dict(equipamentos=linhas, metodologia="sicro2025",
               pessoal=dict(quantidade=n_pes, ntf=ntf, nadm=nadm, nmo=nmo, passagem=passagem, fator=fator, **ind,
                            itinerario=m.get("itinerario"), dmt=m.get("dmt_pessoal")))
    for reg in REGIMES:
        pes = trunc(n_pes * passagem, 2)
        custo = round(pes + soma[reg], 2)
        out[reg] = dict(pessoal=pes, equipamentos=round(soma[reg], 2), custo=custo, bdi=bdi[reg], preco=trunc(custo * (1 + bdi[reg]), 2))
    return out


def calcular_2025(d: dict, *, mo_fn, eq_fn, custo_comp_fn, bdi, meses, quantidade_fn=None, producao_fn=None, desc_fn=None,
                  indice_fn=None, data_base=None) -> dict:
    s = d.get("sicro2025") or padrao_2025(d)
    adm = calcular_adm_2025(s["adm"], mo_fn, eq_fn, bdi, meses, quantidade_fn, producao_fn)
    can = calcular_canteiro_2025(s["canteiro"], s["adm"], bdi)
    usi = calcular_usina_2025(s["usina"], s["canteiro"], custo_comp_fn, bdi)
    mob = calcular_mob_2025(s["mob"], s["adm"], s["canteiro"], eq_fn, desc_fn, bdi, indice_fn, data_base)
    precos = {c: {} for c in CODIGOS}
    for reg in REGIMES:
        precos["ADM"][reg] = adm[reg]["mensal"]
        precos["CAN"][reg] = can[reg]["preco"]
        precos["INSIND"][reg] = usi[reg]["preco"]
        precos["MOB"][reg] = precos["DESMOB"][reg] = mob[reg]["preco"]
    return dict(adm=adm, canteiro=can, usina=usi, mob=mob, precos=precos, metodologia="sicro2025")
