"""Nível de esforço da roçada (aba NE_ROCADA do PATO) — norma DNIT 182/2018-PRO e Res. DNIT nº 8/2023, art. 24, § 4º.

NERp = Σ (NERe_i / d_i) ÷ Σ (1 / d_i)  — média das estações pluviométricas mais próximas ponderada pelo inverso da distância.
NE da roçada (roçadas e capinas por ano) = NERp arredondado + acréscimo (planilha BR-158/MT: 5,10 → 6), ou o valor que
a unidade local justificar pela observação de campo (art. 24, § 4º).
"""
from __future__ import annotations

import json
import math
from pathlib import Path

from . import normalizar as nz

N_ESTACOES = 5  # a norma admite de 2 a 5 estações, necessariamente as mais próximas do ponto médio (qualquer UF)
_EST = None


def estacoes() -> list[dict]:
    """As 82 estações pluviométricas da norma DNIT 182/2018 (Anexo A, arquivo Est-Pluv_NERP do IPR)."""
    global _EST
    if _EST is None:
        _EST = json.loads((Path(__file__).parent / "modelos" / "estacoes_nerp.json").read_text(encoding="utf-8"))["estacoes"]
    return _EST


def distancia_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    """Distância geodésica (círculo máximo, raio médio 6.371 km) — equivale à régua do Google Earth."""
    p1, p2 = math.radians(lat1), math.radians(lat2)
    h = math.sin((p2 - p1) / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(math.radians(lon2 - lon1) / 2) ** 2
    return 2 * 6371.0088 * math.asin(math.sqrt(h))


def estacoes_proximas(lat: float, lon: float, n: int = N_ESTACOES) -> list[dict]:
    n = max(2, min(5, int(n or N_ESTACOES)))
    lst = sorted(({**e, "distancia": round(distancia_km(lat, lon, e["lat"], e["lon"]), 1)} for e in estacoes()),
                 key=lambda e: e["distancia"])
    return [dict(codigo=e["codigo"], uf=e["uf"], nome=e["nome"], municipio=e["municipio"], lat=e["lat"], lon=e["lon"],
                 distancia=e["distancia"], nere=e["nere"]) for e in lst[:n]]


def calcular(d: dict | None) -> dict | None:
    if not d or (not d.get("estacoes") and d.get("adotado") in (None, "")):
        return None
    num = den = 0.0
    linhas = []
    for e in d["estacoes"]:
        dist, ner = nz.numero(e.get("distancia")), nz.numero(e.get("nere"))
        if not dist or ner is None:
            continue
        inv = 1 / dist
        num += ner * inv
        den += inv
        linhas.append(dict(e, inverso=inv, ponderado=ner * inv))
    nerp = num / den if den else None
    acr = nz.numero(d.get("acrescimo")) if d.get("acrescimo") not in (None, "") else 1.0
    sugerido = (math.floor(nerp + 0.5) + acr) if nerp is not None else None
    adotado = nz.numero(d.get("adotado")) if d.get("adotado") not in (None, "") else sugerido
    return dict(estacoes=linhas, nerp=nerp, acrescimo=acr, sugerido=sugerido, adotado=adotado,
                justificativa=d.get("justificativa"), ponto_medio=d.get("ponto_medio"), fonte=d.get("fonte"))


def importar_planilha(caminho: str) -> dict | None:
    import openpyxl
    wv = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    try:
        if "NE_ROCADA" not in wv.sheetnames:
            return None
        est, nerp, adot = [], None, None
        for r in wv["NE_ROCADA"].iter_rows(values_only=True):
            r = tuple(r) + (None,) * 10
            rot = nz.sem_acento(str(r[1] or "")).upper()
            if isinstance(r[1], (int, float)) and nz.numero(r[6]) is not None and nz.numero(r[7]) is not None:
                est.append(dict(codigo=nz.texto(r[2]), uf=nz.texto(r[3]), nome=nz.texto(r[4]), municipio=nz.texto(r[5]),
                                distancia=nz.numero(r[6]), nere=nz.numero(r[7])))
            elif rot.startswith("NERP") and "ADOTAR" not in rot:
                nerp = nz.numero(r[7])
            elif rot.startswith("NE ROCADA"):
                adot = nz.numero(r[7])
        if not est:
            return None
        d = dict(estacoes=est, acrescimo=1.0, adotado=None, justificativa=None)
        c = calcular(d)
        if adot is not None and c and adot != c["sugerido"]:
            d["adotado"] = adot
        return d
    finally:
        wv.close()
