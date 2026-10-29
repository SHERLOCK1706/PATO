"""Leitura de inventários em Excel (modelo DNIT "Cadastro de ..." ou abas CAD_* do PATO).

As colunas são reconhecidas pelo texto do cabeçalho, não pela posição, então o
mesmo importador lê planilhas com layouts ligeiramente diferentes.
"""
from __future__ import annotations

import posixpath
import re
import zipfile
from typing import Any

import openpyxl
from lxml import etree
from openpyxl.utils import column_index_from_string, get_column_letter

from collections import Counter

from . import normalizar as nz
from . import padrao_dnit
from .catalogo import CAMPOS_NUM, CATEGORIAS, REGRAS_RE, categoria_da_aba

_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_NSR = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_MERGE_RE = re.compile(rb'<mergeCell ref="([A-Z]+)(\d+):([A-Z]+)(\d+)"')

ROTULOS_PROJETO = {
    "RODOVIA": "rodovia", "TRECHO": "trecho", "SUBTRECHO": "subtrecho", "SEGMENTO": "segmento",
    "EXTENSAO": "extensao_txt", "SNV VERSAO": "snv_versao", "DATA DO INVENTARIO": "data_inventario",
    "UNIDADE LOCAL": "unidade_local", "DATA-BASE": "data_base", "DATA BASE": "data_base",
    "SNV": "snv", "PRAZO": "prazo",
}


def _merges(path: str) -> dict[str, list[tuple[int, int, int, int]]]:
    """Células mescladas por aba (lidas direto do XML; o modo read_only não as expõe)."""
    out: dict[str, list] = {}
    with zipfile.ZipFile(path) as z:
        wb = etree.fromstring(z.read("xl/workbook.xml"))
        rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        alvo = {r.get("Id"): r.get("Target") for r in rels}
        for s in wb.iter(f"{_NS}sheet"):
            t = alvo.get(s.get(f"{_NSR}id"), "")
            t = t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join("xl", t))
            try:
                raw = z.read(t)
            except KeyError:
                continue
            out[s.get("name")] = [
                (int(m[1]), column_index_from_string(m[0].decode()), int(m[3]), column_index_from_string(m[2].decode()))
                for m in _MERGE_RE.findall(raw)
            ]
    return out


def _ler_linhas(ws, limite_vazias: int = 300) -> list[tuple]:
    linhas, vazias = [], 0
    for row in ws.iter_rows(values_only=True):
        if any(v not in (None, "") for v in row):
            vazias = 0
        else:
            vazias += 1
            if vazias > limite_vazias:
                break
        linhas.append(row)
    return linhas


def _cel(linhas, r: int, c: int):
    """r, c base 1."""
    if r < 1 or r > len(linhas):
        return None
    row = linhas[r - 1]
    return row[c - 1] if c - 1 < len(row) else None


def _achar_cabecalho(linhas) -> int | None:
    for alvo in ("SNV", "KM INICIAL"):
        for r in range(1, min(len(linhas), 80) + 1):
            for v in linhas[r - 1]:
                if isinstance(v, str) and nz.rotulo(v) == alvo:
                    return r
    return None


def _parece_dado(v) -> bool:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return True
    return isinstance(v, str) and nz.eh_snv(v)


def _mapear_colunas(cat: str, rotulos: dict[int, str]) -> tuple[dict[str, list[int]], dict[int, str]]:
    campos: dict[str, list[int]] = {}
    extras: dict[int, str] = {}
    for c in sorted(rotulos):
        lab = rotulos[c]
        for cats, rx, campo in REGRAS_RE:
            if cats is not None and cat not in cats:
                continue
            if rx.search(lab):
                campos.setdefault(campo, []).append(c)
                break
        else:
            extras[c] = lab.split(" > ")[-1]
    return campos, extras


def ler_aba(ws, nome: str, cat: str, merges, vocab: Counter | None = None) -> dict:
    linhas = _ler_linhas(ws)
    rel = dict(aba=nome, categoria=cat, nome_categoria=CATEGORIAS[cat]["nome"], linhas=0,
               colunas={}, nao_mapeadas=[], avisos=[], itens=[])
    hdr = _achar_cabecalho(linhas)
    if not hdr:
        rel["avisos"].append("Cabeçalho (SNV / KM INICIAL) não encontrado.")
        return rel
    ncols = max((len(r) for r in linhas[hdr - 2: hdr + 3]), default=0)
    ds = None
    for r in range(hdr + 1, min(hdr + 15, len(linhas)) + 1):
        if any(_parece_dado(_cel(linhas, r, c)) for c in range(1, ncols + 1)):
            ds = r
            break
    if ds is None:
        rel["avisos"].append("Nenhuma linha de dados.")
        return rel

    # rótulos combinando as linhas de cabeçalho (com mesclagens propagadas)
    hrows = list(range(max(1, hdr - 1), ds))
    grade = {(r, c): _cel(linhas, r, c) for r in hrows for c in range(1, ncols + 1)}
    for (r1, c1, r2, c2) in merges:
        if r1 in hrows:
            v = grade.get((r1, c1))
            for c in range(c1, min(c2, ncols) + 1):
                grade[(r1, c)] = v
    rotulos: dict[int, str] = {}
    for c in range(1, ncols + 1):
        partes = []
        for r in hrows:
            t = nz.rotulo(grade.get((r, c)))
            if t and (not partes or partes[-1] != t):
                partes.append(t)
        if partes:
            rotulos[c] = " > ".join(partes)
    # limite direito da tabela: coluna de Observações, ou duas colunas sem rótulo
    if rotulos:
        inicio = min(rotulos)
        fim, vazias = inicio, 0
        for c in range(inicio, ncols + 1):
            if c in rotulos and "RESUMO" in rotulos[c]:
                break
            if c in rotulos:
                fim, vazias = c, 0
                if "OBSERVA" in rotulos[c]:
                    break
            else:
                vazias += 1
                if vazias >= 2:
                    break
        rotulos = {c: v for c, v in rotulos.items() if inicio <= c <= fim}

    campos, extras = _mapear_colunas(cat, rotulos)
    rel["colunas"] = {k: [f"{get_column_letter(c)} · {rotulos[c].split(' > ')[-1]}" for c in v] for k, v in campos.items()}
    rel["mapa"] = {k: [get_column_letter(c) for c in v] for k, v in campos.items()}
    rel["nao_mapeadas"] = [f"{get_column_letter(c)} · {v}" for c, v in extras.items()]
    for obrig in ("km_ini",):
        if obrig not in campos:
            rel["avisos"].append("Coluna de km/estaca não identificada.")

    for r in range(ds, len(linhas) + 1):
        it = _ler_linha(linhas, r, campos, extras, cat, vocab)
        if it is None:
            continue
        it["origem"] = f"{nome}!{r}"
        rel["itens"].append(it)
    rel["linhas"] = len(rel["itens"])
    return rel


def _ler_linha(linhas, r, campos, extras, cat, vocab: Counter | None = None) -> dict | None:
    val: dict[str, Any] = {}
    for campo, cols in campos.items():
        brutos = [_cel(linhas, r, c) for c in cols]
        if campo in ("lat_ini", "lat_fim", "lon_ini", "lon_fim"):
            eixo = "lat" if campo.startswith("lat") else "lon"
            v = next((x for x in (nz.coordenada(b, eixo) for b in brutos) if x is not None), None)
        elif campo in ("km_ini", "km_fim"):
            nums = [nz.numero(b) for b in brutos]
            nums = [x for x in nums if x is not None]
            v = nums[-1] if nums else None
        elif campo in CAMPOS_NUM:
            v = next((x for x in (nz.numero(b) for b in brutos) if x is not None), None)
        else:
            v = next((x for x in (nz.texto(b) for b in brutos) if x is not None), None)
        if v is not None:
            val[campo] = v
    snv = val.get("snv")
    if snv and not nz.eh_snv(snv) and re.search(r"TOTAL|SOMA|CONSTA|:", nz.sem_acento(snv).upper()):
        return None
    if val.get("km_ini") is None and val.get("lat_ini") is None:
        return None
    if "snv" in campos and not snv and val.get("lat_ini") is None and cat != "ACEIRO":
        return None  # linha solta (anotação), sem SNV nem coordenada
    if cat == "ACEIRO" and not any(k in val for k in ("km_fim", "lat_fim")):
        return None
    item: dict[str, Any] = {"categoria": cat}
    for k in ("snv", "km_ini", "km_fim", "lat_ini", "lon_ini", "lat_fim", "lon_fim", "tipo", "obs"):
        if k in val:
            item[k] = val.pop(k)
    for campo, fn in (("tipo_pista", nz.tipo_pista), ("lado", nz.lado), ("condicao", nz.condicao)):
        if campo in val:
            bruto = val.pop(campo)
            item[campo] = fn(bruto)
            if vocab is not None and str(bruto) != str(item[campo]):
                vocab[(campo, str(bruto), str(item[campo]))] += 1
    extra = {}
    for c, lab in extras.items():
        v = nz.texto(_cel(linhas, r, c))
        if v is not None:
            extra[lab] = v
    if extra:
        val["extra"] = extra
    item["dados"] = val
    return item


def _projeto(wb, nomes_cat: dict[str, str]) -> dict:
    proj: dict[str, Any] = {}

    def varrer(ws, r0, r1):
        linhas = []
        for i, row in enumerate(ws.iter_rows(min_row=r0, max_row=r1, values_only=True)):
            linhas.append(row)
        for i, row in enumerate(linhas):
            for j, v in enumerate(row):
                if not isinstance(v, str):
                    continue
                lab = nz.rotulo(v)
                campo = ROTULOS_PROJETO.get(lab)
                if not campo or campo in proj:
                    continue
                valor = next((x for x in row[j + 1:] if nz.texto(x) is not None), None)
                if valor is None and i + 1 < len(linhas) and j < len(linhas[i + 1]):
                    valor = linhas[i + 1][j]
                if valor is not None and not (isinstance(valor, str) and nz.rotulo(valor) in ROTULOS_PROJETO):
                    proj[campo] = valor

    if "CAPA" in wb.sheetnames:
        ws = wb["CAPA"]
        varrer(ws, 70, 90)
        uf = list(ws.iter_rows(min_row=86, max_row=86, min_col=2, max_col=2, values_only=True))
        if uf and uf[0][0]:
            proj["estado"] = uf[0][0]
    seg = next((n for n, c in nomes_cat.items() if c == "SEG"), None)
    for nome in ([seg] if seg else []) + list(wb.sheetnames[:3]):
        if nome in wb.sheetnames:
            varrer(wb[nome], 1, 12)
    for k in ("data_inventario", "data_base"):
        if k in proj:
            proj[k] = nz.data(proj[k])
    for k in list(proj):
        if k not in ("data_inventario", "data_base", "prazo"):
            proj[k] = nz.texto(proj[k])
    if "prazo" in proj:
        proj["prazo"] = nz.numero(proj["prazo"])
    return proj


def importar_xlsx(path: str) -> dict:
    """Lê todas as abas de inventário reconhecidas. Não grava nada."""
    merges = _merges(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        nomes_cat = {}
        ignoradas = []
        for nome in wb.sheetnames:
            cat = categoria_da_aba(nome)
            if cat:
                nomes_cat[nome] = cat
            else:
                ignoradas.append(nome)
        relatorio = []
        usadas: set[str] = set()
        vocab: Counter = Counter()
        for nome, cat in nomes_cat.items():
            rel = ler_aba(wb[nome], nome, cat, merges.get(nome, []), vocab)
            if cat in usadas and rel["linhas"]:
                rel["avisos"].append("Categoria já lida em outra aba; itens somados.")
            usadas.add(cat)
            relatorio.append(rel)
        projeto = _projeto(wb, nomes_cat)
    finally:
        wb.close()
    itens = [it for rel in relatorio for it in rel["itens"]]
    segs = [i for i in itens if i["categoria"] == "SEG" and i.get("km_ini") is not None]
    if segs:
        projeto["km_ini"] = min(i["km_ini"] for i in segs)
        projeto["km_fim"] = max(i.get("km_fim") or i["km_ini"] for i in segs)
        if not projeto.get("snv"):
            cods = [i["snv"] for i in segs if i.get("snv")]
            if cods:
                projeto["snv"] = cods[0] if len(cods) == 1 else f"{cods[0]} - {cods[-1]}"
    else:
        kms = [i["km_ini"] for i in itens if i.get("km_ini") is not None]
        kms += [i["km_fim"] for i in itens if i.get("km_fim") is not None]
        if kms:
            projeto["km_ini"], projeto["km_fim"] = min(kms), max(kms)
    for rel in relatorio:
        rel.pop("itens")
    log = padrao_dnit.padronizar(itens)
    return dict(projeto=projeto, itens=itens, relatorio=relatorio, ignoradas=ignoradas,
                padronizacao=log, padronizacao_resumo=padrao_dnit.resumir(log),
                vocabulario=padrao_dnit.contar_vocabulario(vocab))


def mapa_colunas(path: str, abas: set[str] | None = None) -> dict[str, dict[str, list[str]]]:
    """{aba: {campo: [letras das colunas]}} das abas de inventário reconhecidas."""
    merges = _merges(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        out = {}
        for nome in wb.sheetnames:
            if abas is not None and nome not in abas:
                continue
            cat = categoria_da_aba(nome)
            if cat:
                out[nome] = ler_aba(wb[nome], nome, cat, merges.get(nome, [])).get("mapa", {})
        return out
    finally:
        wb.close()
