"""Exportação: preencher a planilha modelo do PATO e gerar planilha de inventário."""
from __future__ import annotations

import datetime as _dt
from collections import defaultdict
from typing import Any, Callable

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import normalizar as nz
from .calculo import NOMES_MEDIDA, UNIDADES, dimensao_area, medidas, resumo, tipo_rocada
from .catalogo import CATEGORIAS, COLUNAS_UI, ROTULOS
from .xlsx_patch import XlsxPatcher


def _v(it: dict, campo: str):
    if campo in it:
        return it[campo]
    return it.get("dados", {}).get(campo)


def _masc(c: str | None) -> str | None:
    return {"Satisfatória": "Satisfatório"}.get(c, c) if c else c


def _tem_km(it) -> bool:
    return it.get("km_ini") is not None and it.get("km_fim") is not None and it["km_fim"] != it["km_ini"]


def _ou(formula: str, valor_campo: Callable[[dict], Any]):
    """Fórmula baseada no km, ou o valor informado quando o item não tem km final."""
    def f(it, r):
        return formula.format(r=r) if _tem_km(it) else valor_campo(it)
    return f


def _oac_volume(it, r):
    dim = _v(it, "dimensao")
    if isinstance(dim, (int, float)):
        t = (it.get("tipo") or "").upper()
        import re
        if re.match(r"B[SDTQ]T", t) or "TUBUL" in t:
            return f"=(PI()*(I{r}^2)*M{r}*J{r})/4"
        return f"=I{r}^2*M{r}*J{r}"
    return medidas(it).get("volume")


_BASE8 = lambda c0: dict(zip(["snv", "tipo_pista", "km_ini", "km_fim", "lat_ini", "lon_ini", "lat_fim", "lon_fim"],  # noqa: E731
                             [get_column_letter(c0 + i) for i in range(8)]))

# Layout das abas CAD_* do modelo PATO (linhas de dados = first..last)
LAYOUT: dict[str, dict] = {
    "SEG": dict(sheet="CAD_SEG", first=23, last=30, cols=_BASE8(2), formulas={}),
    "PISTA": dict(sheet="CAD_PISTA", first=22, last=27,
                  cols={**_BASE8(2), "revestimento": "J", "n_faixas": "K", "largura": "L", "rev_fa": "M",
                        "lado_fa": "N", "largura_fa": "O", "obs": "R"},
                  formulas={"P": "=E{r}-D{r}", "Q": "=P{r}*IF(K{r}>0,K{r},2)", "U": "=P{r}*1000*L{r}", "V": None}),
    "ACOST": dict(sheet="CAD_ACOST", first=20, last=27,
                  cols={**_BASE8(1), "revestimento": "I", "larg_e": "J", "larg_d": "K", "obs": "N"},
                  formulas={"L": "=D{r}-C{r}", "M": "=L{r}*((J{r}>0)+(K{r}>0))"}),
    "MFC": dict(sheet="CAD_MFC", first=21, last=243,
                cols={**_BASE8(1), "lado": "I", "condicao": "J", "tipo": "K", "largura": "M", "obs": "P"},
                formulas={"L": _ou("=(D{r}-C{r})*1000", lambda it: medidas(it)["ext_m"]),
                          "N": "=M{r}*L{r}", "O": "=0.2*L{r}"}),
    "STC": dict(sheet="CAD_STC", first=21, last=29,
                cols={**_BASE8(1), "lado": "I", "condicao": "J", "tipo": "K", "largura": "M", "obs": "P"},
                formulas={"L": _ou("=(D{r}-C{r})*1000", lambda it: medidas(it)["ext_m"]),
                          "N": "=L{r}*M{r}", "O": "=L{r}*0.2"}),
    "DAR": dict(sheet="CAD_DAR", first=19, last=2815, masc=True,
                cols={"snv": "A", "tipo_pista": "B", "km_ini": "C", "lat_ini": "D", "lon_ini": "E", "lado": "F",
                      "condicao": "G", "tipo": "H", "extensao": "I", "obs": "J"}, formulas={}),
    "EDA": dict(sheet="CAD_EDA", first=19, last=2815, masc=True,
                cols={"snv": "A", "tipo_pista": "B", "km_ini": "C", "lat_ini": "D", "lon_ini": "E", "lado": "F",
                      "condicao": "G", "tipo": "H", "extensao": "I", "obs": "J"}, formulas={}),
    "OAC": dict(sheet="CAD_OAC", first=25, last=87,
                cols={"snv": "A", "tipo_pista": "B", "km_ini": "C", "lat_ini": "D", "lon_ini": "E", "condicao": "F",
                      "posicao": "G", "tipo": "H", "dimensao": "I", "linhas": "J", "esconsidade": "K",
                      "largura_pista": "L", "extensao": "M", "obs": "O"},
                formulas={"N": _oac_volume}),
    "OAE": dict(sheet="CAD_OAE", first=25, last=29,
                cols={"numero": "A", "nome": "B", "snv": "C", "tipo_pista": "D", "km_ini": "E", "lat_ini": "F",
                      "lon_ini": "G", "extensao": "H", "condicao": "I", "tipo": "J", "posicao": "K", "lado": "L",
                      "obs": "P"},
                formulas={"M": lambda it, r: _v(it, "ext_dispositivo") if _v(it, "ext_dispositivo") is not None else f"=H{r}*2",
                          "N": lambda it, r: _v(it, "area_caiacao") if _v(it, "area_caiacao") is not None else f"=M{r}*1.8"}),
    "ROCADA": dict(sheet="CAD_ROCADA", first=19, last=27,
                   cols={**_BASE8(1), "lado": "I", "tipo": "J", "largura": "K", "obs": "N"},
                   formulas={"L": _ou("=(D{r}-C{r})*1000", lambda it: medidas(it)["ext_m"]),
                             "M": '=(K{r}*L{r})*IF(I{r}="E/D",2,1)/10000'}),
    "DEFENSA": dict(sheet="CAD_DEFENSA", first=21, last=69,
                    cols={**_BASE8(1), "lado": "I", "condicao": "J", "defensa_ancoragem": "K", "tipo": "L",
                          "ext_ancoragem": "N"},
                    formulas={"M": _ou('=(D{r}-C{r})*1000*IF(I{r}="E/D",2,1)', lambda it: medidas(it)["ext_defensa"])}),
    "CERCA": dict(sheet="CAD_CERCA", first=20, last=63,
                  cols={**_BASE8(1), "lado": "I", "condicao": "J", "tipo": "K", "obs": "M"},
                  formulas={"L": _ou("=D{r}-C{r}", lambda it: medidas(it)["ext_km"])}),
    "ALAMB": dict(sheet="CAD_ALAMB", first=20, last=22,
                  cols={**_BASE8(1), "lado": "I", "condicao": "J", "tipo": "K", "obs": "M"},
                  formulas={"L": lambda it, r: medidas(it)["ext_m"]}),
    "ACEIRO": dict(sheet="CAD_ACEIRO", first=4, last=50,
                   cols={"km_ini": "B", "km_fim": "C", "lat_ini": "D", "lon_ini": "E", "lat_fim": "F", "lon_fim": "G",
                         "lado": "H", "cerca": "J", "pct_arboreo": "N", "largura_ate_cerca": "P"},
                   formulas={"I": "=(C{r}-B{r})*1000", "K": "=I{r}*1.5", "L": "=(P{r}-3-1.5)*I{r}", "M": "=I{r}*3"}),
}

SEM_ABA_NO_MODELO = ["GRAMADA", "VALETA", "PLACA", "TACHA", "MARCA", "INSCRICAO", "FUNC"]


def _valor_coluna(cat: str, campo: str, it: dict, lay: dict):
    v = _v(it, campo)
    if campo == "condicao" and lay.get("masc"):
        v = _masc(v)
    elif campo == "tipo" and cat == "ROCADA":
        v = tipo_rocada(v)
    elif campo == "defensa_ancoragem":
        v = (v or "DEFENSA").upper()
    elif campo == "cerca" and v is None:
        v = "SIM"
    elif campo == "linhas" and v is None:
        v = 1
    elif campo == "dimensao" and isinstance(v, str) and nz.numero(v) is not None:
        v = nz.numero(v)
    elif campo == "tipo_pista" and v:
        v = v.upper()
    return v


def _ordenar(itens: list[dict]) -> list[dict]:
    return sorted(itens, key=lambda i: (i.get("km_ini") if i.get("km_ini") is not None else 1e9, i.get("ordem") or 0))


def _texto_km(v: float) -> str:
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def preencher_modelo_pato(projeto: dict, itens: list[dict], modelo: str, saida: str,
                          limpar_ausentes: bool = True) -> list[str]:
    """Grava projeto e inventário nas abas CAD_* (e CAPA) de uma cópia do modelo PATO.

    Fórmulas e demais abas são preservadas; o Excel recalcula tudo ao abrir.
    Retorna o log das operações.
    """
    log: list[str] = []
    p = XlsxPatcher(modelo)
    por_cat: dict[str, list] = defaultdict(list)
    for it in itens:
        por_cat[it["categoria"]].append(it)

    faixas: dict[str, tuple[int, int]] = {}
    for cat in ["ACOST", "SEG", "PISTA", "MFC", "STC", "DAR", "EDA", "OAC", "OAE", "ROCADA", "DEFENSA",
                "CERCA", "ALAMB", "ACEIRO"]:
        lay = LAYOUT[cat]
        sh = lay["sheet"]
        if sh not in p.sheets:
            log.append(f"⚠ Aba {sh} não existe no modelo — {CATEGORIAS[cat]['nome']} não exportado.")
            continue
        lst = _ordenar(por_cat.get(cat, []))
        if not lst and not limpar_ausentes:
            log.append(f"{sh}: mantido como está no modelo (sem itens no inventário).")
            continue
        first, last = lay["first"], lay["last"]
        cap = last - first + 1
        if len(lst) > cap:
            extra = len(lst) - cap
            p.insert_rows(sh, last, extra)
            last += extra
            log.append(f"{sh}: {extra} linha(s) inserida(s) (capacidade {cap} → {cap + extra}); totais e referências ajustados.")
        faixas[cat] = (first, last)
        colunas = list(lay["cols"].values()) + list(lay["formulas"].keys())
        for r in range(first, last + 1):
            for c in colunas:
                p.set(sh, f"{c}{r}", None)
        for k, it in enumerate(lst):
            r = first + k
            for campo, c in lay["cols"].items():
                p.set(sh, f"{c}{r}", _valor_coluna(cat, campo, it, lay))
            for c, f in lay["formulas"].items():
                if f is None:
                    continue
                val = f(it, r) if callable(f) else f.format(r=r)
                p.set(sh, f"{c}{r}", val)
        log.append(f"{sh}: {len(lst)} item(ns) gravado(s) nas linhas {first}–{first + len(lst) - 1 if lst else first}.")

    # ---------------------------------------------------- ajustes de fórmulas
    if "ACOST" in faixas and "CAD_ACOST" in p.sheets:
        a1, a2 = faixas["ACOST"]
        p.set("CAD_ACOST", f"M{p.shifted_row('CAD_ACOST', 28)}", f"=SUM(M{a1}:M{a2})")
        p.set("CAD_ACOST", f"M{p.shifted_row('CAD_ACOST', 29)}", None)
        if "RESUMO_INVENT" in p.sheets:
            for cel, padrao in (("D20", "Concreto*"), ("G20", "Trat*")):
                p.set("RESUMO_INVENT", cel,
                      f'=SUMIF(CAD_ACOST!$I${a1}:$I${a2},"{padrao}",CAD_ACOST!$M${a1}:$M${a2})')
        if "PATO" in p.sheets:
            p.set("PATO", "F59", f"=(SUM(CAD_ACOST!L{a1}:L{a2})*1000)+RESUMO_INVENT!P16+RESUMO_INVENT!M20")
        if "PISTA" in faixas:
            f1 = faixas["PISTA"][0]
            p.set("CAD_PISTA", f"V{f1}",
                  f"=SUMPRODUCT(CAD_ACOST!L{a1}:L{a2},CAD_ACOST!J{a1}:J{a2}+CAD_ACOST!K{a1}:K{a2})*1000")
        log.append("Fórmulas de acostamento por revestimento (RESUMO_INVENT D20/G20, PATO F59, CAD_PISTA V) "
                   "reescritas com SOMASE para acompanhar o novo número de linhas.")
    if "CAD_DAR" in p.sheets:
        p.set("CAD_DAR", "B5", "=CAD_SEG!C5")
    if "RESUMO_INVENT" in p.sheets:
        p.set("RESUMO_INVENT", "AE6", "=CAD_SEG!I7")
        p.set("RESUMO_INVENT", "AE7", "=CAD_SEG!I8")
        gram = sum(medidas(i).get("area_m2", 0) for i in por_cat.get("GRAMADA", []))
        p.set("RESUMO_INVENT", "B44", round(gram, 3))

    # ---------------------------------------------------- dados do projeto
    km0, km1 = projeto.get("km_ini"), projeto.get("km_fim")
    seg_txt = projeto.get("segmento") or (f"km {_texto_km(km0)} ao km {_texto_km(km1)}" if km0 is not None and km1 is not None else None)
    if "CAPA" in p.sheets:
        for cel, campo in (("D74", "rodovia"), ("D75", "trecho"), ("D76", "subtrecho"), ("D79", "snv"),
                           ("D81", "prazo"), ("B86", "estado")):
            if projeto.get(campo) not in (None, ""):
                p.set("CAPA", cel, projeto[campo])
        if seg_txt:
            p.set("CAPA", "D77", seg_txt)
        if km0 is not None and km1 is not None:
            p.set("CAPA", "D78", f"={km1}-{km0}")
        if projeto.get("data_base"):
            try:
                p.set("CAPA", "B85", _dt.datetime.strptime(projeto["data_base"][:10], "%Y-%m-%d"))
            except ValueError:
                log.append(f"⚠ Data-base inválida: {projeto['data_base']}")
    if "CAD_SEG" in p.sheets:
        for cel, campo in (("C5", "rodovia"), ("C6", "trecho"), ("C7", "subtrecho"), ("I7", "snv_versao"),
                           ("I9", "unidade_local")):
            if projeto.get(campo) not in (None, ""):
                p.set("CAD_SEG", cel, projeto[campo])
        if seg_txt:
            p.set("CAD_SEG", "C8", seg_txt)
        if km0 is not None and km1 is not None:
            p.set("CAD_SEG", "C9", f"{_texto_km(km1 - km0)} KM")
        if projeto.get("data_inventario"):
            try:
                p.set("CAD_SEG", "I8", _dt.datetime.strptime(projeto["data_inventario"][:10], "%Y-%m-%d"))
            except ValueError:
                p.set("CAD_SEG", "I8", projeto["data_inventario"])

    sem_aba = [CATEGORIAS[c]["nome"] for c in SEM_ABA_NO_MODELO if por_cat.get(c) and c != "GRAMADA"]
    if sem_aba:
        log.append("Sem aba correspondente no modelo PATO (ficam na planilha de inventário exportada): "
                   + ", ".join(sem_aba) + ".")
    if por_cat.get("GRAMADA"):
        log.append("Área gramada total gravada em RESUMO_INVENT!B44.")
    for cat in ("ACEIRO",):
        if not por_cat.get(cat) and limpar_ausentes:
            log.append("⚠ Sem itens de ACEIRO: os serviços 01.02.08.x (regularização da faixa, roçada costal, "
                       "capina, carga) ficam zerados. Use 'Gerar aceiro a partir da cerca' se aplicável.")
    log.append("⚠ As abas IP / CALCULO_IP (condição do pavimento por km) não são alteradas — revise se o trecho mudou.")
    p.save(saida)
    log.append("Arquivo salvo; o Excel recalcula todas as fórmulas ao abrir.")
    return log


# ---------------------------------------------------------------- inventário
_AZUL = PatternFill("solid", fgColor="1F3A5F")
_CINZA = PatternFill("solid", fgColor="E5E7EB")
_FINA = Side(style="thin", color="9CA3AF")
_BORDA = Border(left=_FINA, right=_FINA, top=_FINA, bottom=_FINA)


def _cabecalho(ws, projeto: dict, titulo: str) -> int:
    ws["A1"] = "DEPARTAMENTO NACIONAL DE INFRAESTRUTURA DE TRANSPORTES"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = titulo
    ws["A2"].font = Font(bold=True, size=11, color="1F3A5F")
    linhas = [("Rodovia", projeto.get("rodovia")), ("Trecho", projeto.get("trecho")),
              ("Subtrecho", projeto.get("subtrecho")), ("Segmento", projeto.get("segmento")),
              ("SNV", projeto.get("snv")), ("Data do inventário", projeto.get("data_inventario"))]
    for k, (a, b) in enumerate(linhas):
        ws.cell(3 + k, 1, a + ":").font = Font(bold=True)
        ws.cell(3 + k, 2, b)
    return 3 + len(linhas) + 1


def exportar_inventario(projeto: dict, itens: list[dict], saida: str) -> str:
    wb = openpyxl.Workbook()
    res = resumo(projeto, itens)

    ws = wb.active
    ws.title = "RESUMO"
    r = _cabecalho(ws, projeto, "RESUMO DO INVENTÁRIO")
    for c, h in enumerate(["Grupo", "Descrição", "Quantidade"], 1):
        cell = ws.cell(r, c, h)
        cell.font, cell.fill, cell.border = Font(bold=True, color="FFFFFF"), _AZUL, _BORDA
    for x in res["resumo_invent"]:
        r += 1
        ws.cell(r, 1, x["grupo"]).border = _BORDA
        ws.cell(r, 2, x["descricao"]).border = _BORDA
        c = ws.cell(r, 3, round(x["valor"], 3))
        c.number_format, c.border = "#,##0.00", _BORDA
    r += 2
    ws.cell(r, 1, "REFLEXO NO PATO (coluna Inventário)").font = Font(bold=True, color="1F3A5F")
    r += 1
    for c, h in enumerate(["Item", "Código", "Serviço", "Regra", "Quantidade", "Unid."], 1):
        cell = ws.cell(r, c, h)
        cell.font, cell.fill, cell.border = Font(bold=True, color="FFFFFF"), _AZUL, _BORDA
    for x in res["pato"]:
        r += 1
        for c, k in enumerate(["item", "codigo", "servico", "regra", "valor", "unidade"], 1):
            v = round(x[k], 3) if k == "valor" else x[k]
            cell = ws.cell(r, c, v)
            cell.border = _BORDA
            if k == "valor":
                cell.number_format = "#,##0.000"
    for col, w in zip("ABCDEF", (14, 48, 56, 44, 16, 8)):
        ws.column_dimensions[col].width = w

    por_cat: dict[str, list] = defaultdict(list)
    for it in itens:
        por_cat[it["categoria"]].append(it)
    for cat, info in CATEGORIAS.items():
        lst = _ordenar(por_cat.get(cat, []))
        if not lst:
            continue
        ws = wb.create_sheet(cat[:31])
        r0 = _cabecalho(ws, projeto, f"CADASTRO — {info['nome'].upper()}")
        campos = ["snv", "tipo_pista", "km_ini", "km_fim", "lat_ini", "lon_ini", "lat_fim", "lon_fim",
                  "lado", "condicao", "tipo"] + COLUNAS_UI.get(cat, []) + ["obs"]
        if info["geom"] == "ponto":
            campos = [c for c in campos if c not in ("km_fim", "lat_fim", "lon_fim")]
        meds = list(medidas(lst[0]).keys())
        cab = [ROTULOS.get(c, c) for c in campos] + [f"{NOMES_MEDIDA.get(m, m)} ({UNIDADES.get(m, '')})" for m in meds]
        for c, h in enumerate(cab, 1):
            cell = ws.cell(r0, c, h)
            cell.font, cell.fill, cell.border = Font(bold=True, color="FFFFFF"), _AZUL, _BORDA
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[r0].height = 32
        for k, it in enumerate(lst, 1):
            mv = medidas(it)
            vals = [_v(it, c) for c in campos] + [round(mv.get(m, 0), 4) for m in meds]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(r0 + k, c, v)
                cell.border = _BORDA
        rt = r0 + len(lst) + 1
        ws.cell(rt, 1, "TOTAL").font = Font(bold=True)
        for j, m in enumerate(meds):
            c = len(campos) + 1 + j
            L = get_column_letter(c)
            cell = ws.cell(rt, c, f"=SUM({L}{r0 + 1}:{L}{rt - 1})")
            cell.font, cell.fill, cell.number_format = Font(bold=True), _CINZA, "#,##0.00"
        ws.freeze_panes = ws.cell(r0 + 1, 3)
        for c in range(1, len(cab) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 13 if c > 2 else 16
    wb.save(saida)
    return saida
