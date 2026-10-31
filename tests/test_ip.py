"""Levantamento da condição do pavimento (abas IP e CALCULO_IP) e ligação com a Planilha PATO."""
import pytest

from pato_app import ip
from pato_app.servico import MODELO_PADRAO, Servico


def test_icp_condicao_e_sugestao():
    assert ip.icp(0, 1, 0) == 7.5 and ip.condicao(7.5) == "BOM" and ip.condicao(30) == "REGULAR" and ip.condicao(75) == "PÉSSIMO"
    assert ip.sugerir(0, 3, 0) == ("R. PROFUNDO", "FS CONT")
    assert ip.sugerir(1, 0, 0) == ("TP/RL/RM", "FS DESC+MICRO")
    assert ip.sugerir(0, 1, 0) == ("CORREÇÃO DE DEFEITOS", "MICRORREVESTIMENTO")
    assert ip.sugerir(0, 0, 0) == ("", "")
    assert [ip.nivel(v) for v in ("Péssimo", "ruim", 0.25, None, 2)] == [3, 2, 1, 0, 2]


def test_calculo_ip_igual_planilha_br158():
    d = ip.importar_planilha(str(MODELO_PADRAO))
    r = ip.calcular(d)
    assert len(r["linhas"]) == 256 and r["kmf_total"] == pytest.approx(255)
    q = {c["chave"]: (c["quantidade"], c["pato"]) for c in r["calculo"]}
    # coluna H (calculado) e L (PATO = MAIOR(calculado; histórico)) da aba CALCULO_IP
    assert q == {"REPARO LOCALIZADO": (16, 107.99875), "TAPA BURACO": (32, 32), "REMENDO PROFUNDO": (495, 992.5645),
                 "CORREÇÃO DE DEFEITOS": (846, 846), "FRESAGEM DESCONTÍNUA": (2250, 2250),
                 "FRESAGEM CONTÍNUA": (1260, 3604.8485), "MICRORREVESTIMENTO": (190800, 259524.999)}
    assert r["divergencias"] == 5  # km em que o engenheiro escolheu diferente da regra sugerida


def test_pato_recalcula_com_o_levantamento(tmp_path):
    sv = Servico(pasta=tmp_path)
    pid = sv.criar_projeto(dict(nome="t", rodovia="BR-158/MT", km_ini=85, km_fim=212.5))["id"]
    sv.orcamento.importar_referencia(pid, str(MODELO_PADRAO))
    c = sv.orcamento.calcular(pid)
    it = {x["codigo"]: x for x in c["itens"] if x["tipo"] == "servico" and (x.get("base") or {}).get("tipo") == "ip"}
    assert it["4915705"]["quantidade"] == 846 and it["5214001"]["quantidade"] == 5850  # (7 + 53) / 2 × 195
    d = sv.db.dados_projeto(pid, "ip")
    for ln in d["linhas"]:
        if ln["primaria"] == "CORREÇÃO DE DEFEITOS":
            ln["primaria"] = ""
    sv.orcamento.ip_salvar(pid, d["linhas"])
    c2 = sv.orcamento.calcular(pid)
    corr = next(x for x in c2["itens"] if x.get("codigo") == "4915705")
    assert corr["quantidade"] == pytest.approx(486.851)  # sem levantamento, vale o histórico
    g = sv.orcamento.ip_gerar_grade(pid)
    assert len(g["linhas"]) == 2 * 128 and g["linhas"][0]["km_ini"] == 85 and g["linhas"][-1]["km_fim"] == 85


def test_importa_formato_antigo_icp(tmp_path):
    """Modelo antigo: aba ICP (Alta/Média/Baixa), CALCULO_ICP com nível de serventia, intervalo de linhas e soma de linhas."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ICP"
    ws["E2"], ws["F2"], ws["G2"], ws["H2"], ws["T2"], ws["U2"] = ("Rodovia", "km inicial", "km final", "Extensão (km)",
                                                                  "INTERVENÇÃO PRIMÁRIA", "INTERVENÇÃO SECUNDÁRIA")
    ws["I3"], ws["L3"], ws["O3"] = "Panela", "Remendo", "Trincamento"
    for c0 in ("I", "L", "O"):
        for k, g in enumerate(("Alta", "Média", "Baixa")):
            ws[f"{chr(ord(c0) + k)}4"] = g
    dados = [("BR-174", 0, 1, "", "X", "", "RL / TP / RP", "FS CONT"),       # linha 5: remendo médio
             ("BR-174", 1, 2, "", "", "X", "CORREÇÃO+FS DESC", ""),          # linha 6: remendo baixo
             ("BR-174", 2, 3, "X", "", "", "RL / TP / RP", "")]               # linha 7: panela alta (fora do intervalo 5:6)
    for i, (rod, a, b, pa, rm, rb, prim, sec) in enumerate(dados, start=5):
        ws[f"E{i}"], ws[f"F{i}"], ws[f"G{i}"], ws[f"H{i}"] = rod, a, b, b - a
        ws[f"I{i}"], ws[f"M{i}"], ws[f"N{i}"] = pa, rm, rb
        ws[f"T{i}"], ws[f"U{i}"] = prim, sec
    c = wb.create_sheet("CALCULO_ICP")
    c["C2"], c["O2"], c["H3"], c["O3"] = "MEMÓRIA DE CÁLCULO - ICP", "LARGURA", "NÍVEL DE SERVENTIA", 3.5
    c["A5"], c["C5"], c["D5"], c["F5"], c["H5"], c["J5"], c["O5"] = ("RL / TP / RP", "REPARO LOCALIZADO",
        "=ROUNDUP(SUMIFS(ICP!$H$5:$H$6,ICP!$T$5:$T$6,CALCULO_ICP!A5),0)", 14, 2, 3, 14)
    c["A6"], c["C6"], c["D6"], c["F6"], c["H6"], c["J6"], c["O6"] = ("CORREÇÃO+FS DESC", "CORREÇÃO DE DEFEITOS",
        "=ROUNDUP(SUMIFS(ICP!$H$5:$H$7,ICP!$T$5:$T$7,CALCULO_ICP!A6),0)", "=1000*$O$3*0.05*O6", 2, 0, 0.05)
    c["A7"], c["C7"], c["D7"], c["F7"], c["H7"], c["O7"] = ("FS CONT", "FRESAGEM CONTÍNUA",
        "=ROUNDUP(SUMIFS(ICP!$H$5:$H$7,ICP!$U$5:$U$7,A7),0)", "=1000*$O$3*0.05*O7", 1, 1)
    c["A8"], c["C8"], c["D8"], c["F8"], c["H8"], c["O8"] = ("MICRO", "MICRORREVESTIMENTO A FRIO", "=D6+D7",
                                                            "=1000*$O$3*O8", 1.1, 1)
    arq = tmp_path / "antigo.xlsx"
    wb.save(arq)
    # valores em cache (o openpyxl não calcula): grava D como número na cópia lida com data_only
    from pato_app import recalculo
    if recalculo.soffice():
        arq = recalculo.recalcular(str(arq), str(tmp_path / "antigo_calc.xlsx"))
    else:
        pytest.skip("LibreOffice ausente: sem valores calculados para a aba CALCULO_ICP")
    d = ip.importar_planilha(str(arq))
    assert d["parametros"]["modelo"] == "antigo" and [ln["remendo"] for ln in d["linhas"]] == [2, 1, 0]
    r = ip.calcular(d)
    q = {c["chave"]: c for c in r["calculo"]}
    assert q["REPARO LOCALIZADO"]["inventario"] == 1 and q["REPARO LOCALIZADO"]["quantidade"] == 28  # linha 7 fora de 5:6; ×2
    assert q["CORREÇÃO DE DEFEITOS"]["ne_calc"] == pytest.approx(8.75) and q["CORREÇÃO DE DEFEITOS"]["quantidade"] == 17.5
    assert q["FRESAGEM CONTÍNUA"]["quantidade"] == 175
    assert q["MICRORREVESTIMENTO"]["inventario"] == 2 and q["MICRORREVESTIMENTO"]["quantidade"] == pytest.approx(7700)


def test_icm_resolucao_6_2025():
    """Exemplo do Anexo I da Res. DNIT 6/2025: sem panela, remendos 'Ruim' (3 a 5/km), trincas < 10%, roçada e drenagem
    'Bom', sinalização vertical e horizontal 'Péssimo' → IP 20, IC 50, ICM 29 (BOM)."""
    from pato_app import ip as ipm
    r = ipm.calcular(dict(linhas=[dict(km_ini=0, km_fim=1, n_panelas=0, n_remendos=4, pct_trinca=5, rocada="Bom",
                                       drenagem="Bom", sinal_vertical="Péssimo", sinal_horizontal="Péssimo")]))
    ln = r["linhas"][0]
    assert (ln["icp"], ln["ic"], ln["icm"], ln["condicao_icm"]) == (20.0, 50.0, 29.0, "BOM")
    assert [ipm.nivel_contagem(n) for n in (0, 2, 3, 6)] == [0, 1, 2, 3]
    assert [ipm.nivel_trinca(v) for v in (0, 9.9, 10, 50)] == [0, 1, 2, 3]
