"""Exportação da Planilha PATO (Res. DNIT nº 8/2023): Anexo XIX (planilha), Anexo XX (% dos grupos), Anexo XXI (curva ABC),
cronograma físico-financeiro e grupo de desempenho — com fórmulas para conferência."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from . import pato_orcamento as po

FINA = Side(style="thin", color="999999")
BORDA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
CAB = PatternFill("solid", fgColor="1F4E79")
TIT = PatternFill("solid", fgColor="DDEBF7")
BRANCO = Font(bold=True, color="FFFFFF")
NEG = Font(bold=True)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
MOEDA = "#,##0.00"


def _ident(ws, p, cfg, titulo, ate):
    ws.cell(2, 2, "MINISTÉRIO DOS TRANSPORTES").font = NEG
    ws.cell(3, 2, "DEPARTAMENTO NACIONAL DE INFRAESTRUTURA DE TRANSPORTES").font = NEG
    ws.cell(4, 2, f"SUPERINTENDÊNCIA REGIONAL DO DNIT NO ESTADO DE {cfg.get('uf') or p.get('uf') or ''}").font = NEG
    linhas = [("RODOVIA:", p.get("rodovia"), "MÊS BASE:", cfg.get("data_base")),
              ("TRECHO:", p.get("trecho"), "EXTENSÃO:", f"{(p.get('km_fim') or 0) - (p.get('km_ini') or 0):.3f} km"),
              ("SUBTRECHO:", p.get("subtrecho"), "", ""),
              ("SEGMENTO:", f"km {p.get('km_ini')} <--> km {p.get('km_fim')}", "", ""),
              ("SNV (Versão):", f"{p.get('snv') or ''} {('(' + p['snv_versao'] + ')') if p.get('snv_versao') else ''}", "", "")]
    for i, (a, b, c, d) in enumerate(linhas):
        ws.cell(5 + i, 2, a).font = NEG
        ws.cell(5 + i, 3, b)
        ws.cell(5 + i, ate - 1, c).font = NEG
        ws.cell(5 + i, ate, d)
    ws.cell(11, 2, titulo).font = Font(bold=True, size=12)
    ws.merge_cells(start_row=11, start_column=2, end_row=11, end_column=ate)


def _cab(ws, linha, cols, col0=2):
    for j, t in enumerate(cols):
        c = ws.cell(linha, col0 + j, t)
        c.fill, c.font, c.alignment, c.border = CAB, BRANCO, CENTRO, BORDA
    ws.row_dimensions[linha].height = 34


def exportar_xlsx(orc, pid: int, saida: str) -> str:
    p = orc.sv.projeto(pid)
    c = orc.calcular(pid)
    cfg = c["config"]
    res = orc.resumo(pid)
    wb = Workbook()

    # ------------------------------------------------ Anexo XIX — planilha PATO
    ws = wb.active
    ws.title = "PATO"
    _ident(ws, p, cfg, "PLANO ANUAL DE TRABALHO E ORÇAMENTO - PATO (Anexo XIX — Res. DNIT nº 8/2023)", 14)
    _cab(ws, 13, ["ITEM", "CÓDIGO", "DISCRIMINAÇÃO", "BASE", "UND", "NÍVEL DE ESFORÇO", "QUANTIDADE", "UND",
                  "PREÇO UNIT. SEM DESONERAÇÃO", "PARCIAL SEM DESONERAÇÃO", "PREÇO UNIT. COM DESONERAÇÃO",
                  "PARCIAL COM DESONERAÇÃO", "ORIGEM DA QUANTIDADE"])
    lin = 14
    ini = lin
    for r in c["itens"]:
        if r["tipo"] == "titulo":
            ws.cell(lin, 2, r["item"]).font = NEG
            ws.cell(lin, 4, r["descricao"]).font = NEG
            for col in range(2, 15):
                ws.cell(lin, col).fill = TIT
        else:
            direta = (r.get("base") or {}).get("tipo") == "direta"
            fq = "ROUND" if r.get("arred") == "arred" else "TRUNC"
            vals = [r["item"], r["codigo"], r.get("descricao"), r.get("base_valor"), r.get("unidade_base") or "",
                    (r["quantidade"] / r["base_valor"]) if direta and r.get("base_valor") else r.get("ne"),
                    r["quantidade"] if direta else f"={fq}(E{lin}*G{lin},3)", r.get("unidade"),
                    r.get("preco_SD"), f"=TRUNC(H{lin}*J{lin},2)", r.get("preco_CD"), f"=TRUNC(H{lin}*L{lin},2)",
                    r.get("base_fonte")]
            for j, v in enumerate(vals):
                cel = ws.cell(lin, 2 + j, v)
                cel.border = BORDA
                if j in (8, 9, 10, 11):
                    cel.number_format = MOEDA
                elif j in (3, 6):
                    cel.number_format = "#,##0.000"
        lin += 1
    fim = lin - 1
    lin += 1
    ws.cell(lin, 4, "TOTAL GERAL DO ORÇAMENTO").font = NEG
    for col in ("K", "M"):
        cel = ws[f"{col}{lin}"]
        cel.value = f"=SUM({col}{ini}:{col}{fim})"
        cel.font, cel.number_format = NEG, MOEDA
    ws.cell(lin + 1, 4, "CONDIÇÃO MAIS VANTAJOSA").font = NEG
    ws.cell(lin + 1, 11, f'=IF(K{lin}<=M{lin},"SEM DESONERAÇÃO","COM DESONERAÇÃO")').font = NEG
    for col, w in zip(range(2, 15), [14, 11, 60, 13, 7, 10, 14, 7, 13, 15, 13, 15, 30]):
        ws.column_dimensions[L(col)].width = w
    ws.freeze_panes = "E14"

    # ------------------------------------------------ Anexo XX — % dos grupos
    wg = wb.create_sheet("PERCENTUAL_GRUPOS")
    _ident(wg, p, cfg, "DISTRIBUIÇÃO PERCENTUAL DO ORÇAMENTO CONFORME GRUPOS DE TAREFAS (Anexo XX)", 7)
    _cab(wg, 13, ["SEÇÃO", "GRUPOS DE TAREFAS", "TOTAL SEM DESONERAÇÃO", "%", "TOTAL COM DESONERAÇÃO", "%"])
    n = len(res["secoes"])
    for i, s in enumerate(res["secoes"]):
        x = 14 + i
        for j, v in enumerate([s["secao"], s["nome"], s["SD"], f"=D{x}/D${14 + n}", s["CD"], f"=F{x}/F${14 + n}"]):
            cel = wg.cell(x, 2 + j, v)
            cel.border = BORDA
            cel.number_format = MOEDA if j in (2, 4) else ("0.00%" if j in (3, 5) else "General")
    t = 14 + n
    wg.cell(t, 3, "TOTAL GERAL DE SERVIÇOS").font = NEG
    for col in ("D", "F"):
        wg[f"{col}{t}"] = f"=SUM({col}14:{col}{t - 1})"
        wg[f"{col}{t}"].number_format = MOEDA
    for col, w in zip("BCDEFG", [8, 60, 18, 9, 18, 9]):
        wg.column_dimensions[col].width = w

    # ------------------------------------------------ Anexo XXI — curva ABC
    wa = wb.create_sheet("CURVA_ABC")
    reg = res["vantajosa"]
    _ident(wa, p, cfg, f"CURVA ABC (Anexo XXI) — {po.REG_TXT[reg]}", 11)
    _cab(wa, 13, ["GRUPO DE TAREFAS", "CÓDIGO", "DESCRIÇÃO", "UND.", "QUANTIDADE TOTAL", "PREÇO UNITÁRIO", "CUSTO TOTAL",
                  "PROPORÇÃO (%)", "PROPORÇÃO ACUMULADA (%)", "FAIXA"])
    na = len(res["abc"])
    for i, a in enumerate(res["abc"]):
        x = 14 + i
        vals = [po.SECOES.get(a["secao"], a["secao"]), a["codigo"], a["descricao"], a["unidade"], a["quantidade"], a["preco"],
                a["total"], f"=H{x}/H${14 + na}", f"=SUM(I$14:I{x})", a["faixa"]]
        for j, v in enumerate(vals):
            cel = wa.cell(x, 2 + j, v)
            cel.border = BORDA
            cel.number_format = MOEDA if j in (5, 6) else ("0.00%" if j in (7, 8) else "#,##0.000" if j == 4 else "General")
    wa.cell(14 + na, 4, "ORÇAMENTO TOTAL").font = NEG
    wa[f"H{14 + na}"] = f"=SUM(H14:H{13 + na})"
    wa[f"H{14 + na}"].number_format = MOEDA
    wa.cell(15 + na, 2, "Faixa A: até 80% do custo acumulado · Faixa B: de 80% a 95% · Faixa C: demais itens.")
    for col, w in zip("BCDEFGHIJK", [30, 10, 60, 7, 14, 12, 15, 10, 12, 7]):
        wa.column_dimensions[col].width = w

    # ------------------------------------------------ cronograma
    wc = wb.create_sheet("CRONOGRAMA")
    cr = res["cronograma"]
    _ident(wc, p, cfg, f"CRONOGRAMA FÍSICO-FINANCEIRO — {po.REG_TXT[reg]}", 4 + cr["meses"])
    _cab(wc, 13, ["GRUPO DE TAREFAS", "TOTAL"] + [f"MÊS {i + 1}" for i in range(cr["meses"])])
    for i, ln in enumerate(cr["linhas"]):
        x = 14 + i
        wc.cell(x, 2, ln["nome"]).border = BORDA
        wc.cell(x, 3, f"=SUM(D{x}:{L(3 + cr['meses'])}{x})").number_format = MOEDA
        for k, v in enumerate(ln["meses"]):
            cel = wc.cell(x, 4 + k, v)
            cel.number_format, cel.border = MOEDA, BORDA
    t = 14 + len(cr["linhas"])
    wc.cell(t, 2, "TOTAL DO MÊS").font = NEG
    wc.cell(t + 1, 2, "% ACUMULADO").font = NEG
    for k in range(cr["meses"] + 1):
        col = L(3 + k)
        wc[f"{col}{t}"] = f"=SUM({col}14:{col}{t - 1})"
        wc[f"{col}{t}"].number_format = MOEDA
        if k:
            wc[f"{col}{t + 1}"] = f"=SUM($D{t}:{col}{t})/$C{t}"
            wc[f"{col}{t + 1}"].number_format = "0.00%"
    wc.column_dimensions["B"].width = 50
    wc.column_dimensions["C"].width = 16
    for k in range(cr["meses"]):
        wc.column_dimensions[L(4 + k)].width = 14

    # ------------------------------------------------ desempenho
    for cod, g in (c.get("desempenho") or {}).items():
        wd = wb.create_sheet(f"DESEMPENHO_{cod}"[:31])
        _ident(wd, p, cfg, f"{cod} - {g['descricao']}", 12)
        _cab(wd, 13, ["SERVIÇO", "DESCRIÇÃO", "QTDE TOTAL", "NÍVEL DE ESFORÇO", "FREQ. ANUAL", "DESCRIÇÃO DA QTDE",
                      "QTDE POR KM", "UND", "PREÇO SD", "R$/KM/ANO SD", "PREÇO CD", "R$/KM/ANO CD"])
        x = 14
        for ln in g["linhas"]:
            fmlh = (ln["qtd_km_calc"] if ln.get("qtd_total_calc") is None
                    else f"=TRUNC(D{x}*E{x}*F{x}/{c['extensao']},3)")
            vals = [ln["codigo"], ln.get("descricao"), ln.get("qtd_total_calc") if ln.get("qtd_total_calc") is not None else "-",
                    ln.get("ne") or 1, ln.get("freq") or 1, ln.get("desc_qtd"), fmlh, ln.get("unidade"),
                    ln.get("preco"), f"=H{x}*J{x}", ln.get("preco_CD"), f"=H{x}*L{x}"]
            for j, v in enumerate(vals):
                wd.cell(x, 2 + j, v).border = BORDA
            x += 1
        for k, (rot, fs, fc) in enumerate((("R$ / KM / ANO:", f"=SUM(K14:K{x - 1})", f"=SUM(M14:M{x - 1})"),
                                           ("R$ / KM / MÊS:", f"=TRUNC(K{x + 1}/12,2)", f"=TRUNC(M{x + 1}/12,2)"),
                                           ("R$ / MÊS:", f"=TRUNC(K{x + 2}*{c['extensao']},2)", f"=TRUNC(M{x + 2}*{c['extensao']},2)"))):
            wd.cell(x + 1 + k, 10, rot).font = NEG
            wd.cell(x + 1 + k, 11, fs).number_format = MOEDA
            wd.cell(x + 1 + k, 13, fc).number_format = MOEDA
        for col, w in zip("BCDEFGHIJKLM", [10, 50, 13, 10, 9, 34, 11, 6, 11, 13, 11, 13]):
            wd.column_dimensions[col].width = w
    wb.save(saida)
    return saida
