"""Verificação do SNV e correção de pontos (coordenadas, km e trecho) do inventário.

Cada ação proposta traz: item, origem (aba!linha da planilha importada), campo, valor original,
valor novo, motivo, método, confiança e situação:
  - "corrigido"              alta confiança (erro de digitação identificado, recorte no limite...)
  - "corrigido — conferir"   valor estimado pelo eixo/km; recomenda-se conferir
  - "verificar em campo"     não foi possível determinar o valor correto; nada é alterado
  - "recortado"              feição cortada no limite do trecho
  - "removido"               feição inteiramente fora do trecho
"""
from __future__ import annotations

import csv
import re
from collections import Counter, defaultdict
from typing import Any

import openpyxl

from . import normalizar as nz
from .calculo import ext_m
from .catalogo import CATEGORIAS
from .geo import Eixo, dist_km

TOL_OK = 1.0      # km: divergência aceitável entre km declarado e coordenada (odômetro × GPS)
TOL_GRAVE = 2.0   # km: acima disso é erro
TOL_DIGITO = 0.3  # km: candidato de correção de digitação precisa cair a até 300 m do eixo
SEM_SNV = {"SEG", "ACEIRO"}  # categorias cujo cadastro não tem coluna SNV

_COD_RE = re.compile(r"(\d{3})\s*B?\s*([A-Z]{2})\s*(\d{4})", re.I)
CORRIGIDO, CONFERIR, VERIFICAR, RECORTADO, REMOVIDO = (
    "corrigido", "corrigido — conferir", "verificar em campo", "recortado", "removido")


# ------------------------------------------------------------------ SNV
def codigos(txt) -> list[str]:
    """Extrai e normaliza códigos SNV ('158MT0182' → '158BMT0182')."""
    if not txt:
        return []
    return [f"{a}B{b.upper()}{c}" for a, b, c in _COD_RE.findall(str(txt).upper())]


def _fmt_cods(cods: list[str]) -> str:
    return " / ".join(cods)


def tabela_seg(itens) -> list[dict]:
    segs = sorted((i for i in itens if i["categoria"] == "SEG" and i.get("km_ini") is not None),
                  key=lambda i: i["km_ini"])
    out = []
    for s in segs:
        c = codigos(s.get("snv"))
        if c:
            out.append(dict(codigo=c[0], km_ini=s["km_ini"], km_fim=s.get("km_fim") or s["km_ini"], id=s.get("id")))
    return out


def esperado(tab: list[dict], a: float, b: float | None) -> list[str]:
    lo, hi = (a, a) if b is None else (min(a, b), max(a, b))
    out = []
    for s in tab:
        if hi - lo < 1e-6:
            if s["km_ini"] - 1e-6 <= lo < s["km_fim"] + 1e-6 or (s is tab[-1] and abs(lo - s["km_fim"]) < 1e-6):
                out.append(s["codigo"])
                break
        elif min(hi, s["km_fim"]) - max(lo, s["km_ini"]) > 0.001:
            out.append(s["codigo"])
    return out


def inferir_tabela(itens) -> list[dict]:
    """Limites de cada código SNV como usados nos itens (para quando o cadastro de segmentos diverge)."""
    meio: dict[str, list[float]] = defaultdict(list)
    for it in itens:
        if it["categoria"] == "SEG" or it.get("km_ini") is None:
            continue
        c = codigos(it.get("snv"))
        if len(c) == 1:
            kf = it.get("km_fim") if it.get("km_fim") is not None else it["km_ini"]
            meio[c[0]].append((it["km_ini"] + kf) / 2)
    total = sum(len(v) for v in meio.values())
    validos = {c: sorted(v) for c, v in meio.items() if len(v) >= max(5, 0.01 * total)}
    ordem = sorted(validos, key=lambda c: validos[c][len(validos[c]) // 2])
    tab = []
    for i, c in enumerate(ordem):
        tab.append(dict(codigo=c, km_ini=None, km_fim=None, n=len(validos[c]), inferido=True))
    for i in range(len(ordem) - 1):
        a, b = validos[ordem[i]], validos[ordem[i + 1]]
        cands = sorted(set(a[-20:] + b[:20]))
        melhor = min(cands, key=lambda x: sum(1 for v in a if v > x) + sum(1 for v in b if v < x))
        tab[i]["km_fim"] = tab[i + 1]["km_ini"] = round(melhor, 3)
    if tab:
        tab[0]["km_ini"] = round(min(validos[ordem[0]]), 3)
        tab[-1]["km_fim"] = round(max(validos[ordem[-1]]), 3)
    return tab


def ler_snv_oficial(caminho: str, rodovia: str | None = None, uf: str | None = None) -> list[dict]:
    """Lê a planilha/CSV do SNV publicada pelo DNIT (colunas reconhecidas pelo nome)."""
    linhas: list[list] = []
    if caminho.lower().endswith((".csv", ".txt")):
        with open(caminho, encoding="utf-8-sig", errors="replace") as f:
            amostra = f.read(4096)
            f.seek(0)
            dial = csv.Sniffer().sniff(amostra, delimiters=";,\t")
            linhas = [r for r in csv.reader(f, dial)]
    else:
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        try:
            ws = max(wb.worksheets, key=lambda w: w.max_row or 0)
            linhas = [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    hdr, cols = None, {}
    for i, r in enumerate(linhas[:50]):
        labs = [nz.rotulo(x) for x in r]
        c = {}
        for j, l in enumerate(labs):
            if re.search(r"^(VL_)?COD(IGO)?( SNV)?$|^CODIGO", l) and "cod" not in c:
                c["cod"] = j
            elif re.search(r"KM.?INIC|^KM INICIAL|VL_KM_INIC", l):
                c["ki"] = j
            elif re.search(r"KM.?FIN|^KM FINAL|VL_KM_FINA", l):
                c["kf"] = j
            elif re.search(r"^(VL_)?BR$|^RODOVIA$", l):
                c["br"] = j
            elif re.search(r"^(SG_)?UF$", l):
                c["uf"] = j
            elif re.search(r"LOCAL DE INICIO|DS_LOCAL_I|^INICIO", l):
                c["li"] = j
            elif re.search(r"LOCAL DE FIM|DS_LOCAL_F|^FIM$", l):
                c["lf"] = j
            elif re.search(r"SUPERF|DS_SUP", l):
                c["sup"] = j
        if {"cod", "ki", "kf"} <= set(c):
            hdr, cols = i, c
            break
    if hdr is None:
        raise ValueError("Não encontrei as colunas Código / km inicial / km final na planilha do SNV.")
    br = re.search(r"\d{3}", rodovia or "")
    uf_ = (uf or (rodovia or "").split("/")[-1] if rodovia and "/" in rodovia else uf or "").upper().strip()
    out = []
    for r in linhas[hdr + 1:]:
        if len(r) <= max(cols.values()):
            continue
        cod = codigos(r[cols["cod"]])
        ki, kf = nz.numero(r[cols["ki"]]), nz.numero(r[cols["kf"]])
        if not cod or ki is None or kf is None:
            continue
        cod = cod[0]
        if br and not cod.startswith(br.group(0)):
            continue
        if uf_ and len(uf_) == 2 and cod[4:6] != uf_:
            continue
        out.append(dict(codigo=cod, km_ini=ki, km_fim=kf,
                        local_ini=nz.texto(r[cols["li"]]) if "li" in cols else None,
                        local_fim=nz.texto(r[cols["lf"]]) if "lf" in cols else None,
                        superficie=nz.texto(r[cols["sup"]]) if "sup" in cols else None))
    out.sort(key=lambda s: s["km_ini"])
    return out


# ------------------------------------------------------- correção de digitação
def _variantes(v: float):
    s = f"{v:.6f}".rstrip("0").rstrip(".")
    yield -v, "sinal trocado"
    pos = [i for i, ch in enumerate(s) if ch.isdigit()]
    for i in pos:
        for d in "0123456789":
            if d != s[i]:
                t = s[:i] + d + s[i + 1:]
                try:
                    yield float(t), f"dígito {s[i]}→{d}"
                except ValueError:
                    pass
    for a, b in zip(pos, pos[1:]):
        if b == a + 1 and s[a] != s[b]:
            t = s[:a] + s[b] + s[a] + s[b + 1:]
            yield float(t), f"dígitos invertidos {s[a]}{s[b]}→{s[b]}{s[a]}"


def corrigir_digitacao(lat: float, lon: float, alvo: tuple[float, float]):
    """Procura um único erro de digitação que traga o ponto para perto do eixo."""
    cands = [((lon, lat), "latitude e longitude trocadas")]
    cands += [((x, lon), f"latitude: {m}") for x, m in _variantes(lat)]
    cands += [((lat, x), f"longitude: {m}") for x, m in _variantes(lon)]
    cands += [((-lat, -lon), "sinal trocado em latitude e longitude")]
    ok = []
    for (la, lo), m in cands:
        if -90 <= la <= 90 and -180 <= lo <= 180:
            d = dist_km(alvo, (la, lo))
            if d <= TOL_DIGITO:
                ok.append((d, la, lo, m))
    if not ok:
        return None
    ok.sort()
    if len(ok) > 1 and ok[1][0] < ok[0][0] * 1.5 and (ok[1][1], ok[1][2]) != (ok[0][1], ok[0][2]):
        return None  # ambíguo
    return ok[0]


# ------------------------------------------------------------------ análise
def _r(v, n=6):
    return None if v is None else round(v, n)


def analisar(projeto: dict, itens: list[dict], snv_oficial: list[dict] | None = None,
             versao_oficial: str | None = None) -> dict:
    eixo = Eixo(itens)
    acoes: list[dict] = []
    K0, K1 = projeto.get("km_ini"), projeto.get("km_fim")
    por_cat: dict[str, list] = defaultdict(list)
    for it in itens:
        por_cat[it["categoria"]].append(it)
    # cópia de trabalho (as correções de km influenciam o SNV)
    novo = {it["id"]: dict(it) for it in itens}

    def acao(it, campo, de, para, motivo, metodo, conf, status, **extra):
        acoes.append(dict(id=it["id"], categoria=it["categoria"], nome_categoria=CATEGORIAS[it["categoria"]]["nome"],
                          origem=it.get("origem"), km_ini=it.get("km_ini"), km_fim=it.get("km_fim"), campo=campo,
                          de=de, para=para, motivo=motivo, metodo=metodo, confianca=conf, situacao=status, **extra))
        if status in (CORRIGIDO, CONFERIR, RECORTADO) and campo in ("km_ini", "km_fim", "lat_ini", "lon_ini",
                                                                     "lat_fim", "lon_fim", "snv"):
            novo[it["id"]][campo] = para

    coords_por_ponto: dict[tuple, set] = defaultdict(set)
    for it in itens:
        for p in ("ini", "fim"):
            if it.get(f"lat_{p}") is not None and it.get(f"km_{p}") is not None:
                coords_por_ponto[(round(it[f"lat_{p}"], 5), round(it[f"lon_{p}"], 5))].add(round(it[f"km_{p}"], 1))

    # ---------------------------------------------- 1. limites do trecho
    removidos: set = set()
    if K0 is not None and K1 is not None:
        for it in itens:
            if it["categoria"] == "SEG" or it.get("km_ini") is None:
                continue
            a, b = it["km_ini"], it.get("km_fim")
            lo, hi = (a, a) if b is None else (min(a, b), max(a, b))
            if hi < K0 - 0.001 or lo > K1 + 0.001:
                # coordenada indica que o item está dentro? então o km é que está errado
                sug = None
                if it.get("lat_ini") is not None:
                    pr = eixo.projetar(it["lat_ini"], it["lon_ini"])
                    if pr and pr[1] <= TOL_DIGITO and K0 + 0.2 <= pr[0] <= K1 - 0.2:
                        sug = round(pr[0], 3)
                if sug is not None:
                    acao(it, "km_ini", a, sug, f"km {a} fora do trecho ({K0}–{K1}), mas a coordenada está na rodovia no km ≈{sug}",
                         "projeção da coordenada no eixo", "média", VERIFICAR)
                else:
                    acao(it, "item", f"km {a}" + (f"–{b}" if b is not None else ""), None,
                         f"inteiramente fora do trecho (km {K0}–{K1})", "limite do segmento", "alta", REMOVIDO)
                    removidos.add(it["id"])
            elif b is not None and (lo < K0 - 0.001 or hi > K1 + 0.001):
                for campo, k in (("km_ini", a), ("km_fim", b)):
                    lim = K0 if k < K0 else K1 if k > K1 else None
                    if lim is None:
                        continue
                    ponta = campo[3:]
                    acao(it, campo, k, lim, f"ultrapassa o limite do trecho ({K0}–{K1})", "recorte no limite",
                         "alta", RECORTADO, ext_antes=round(ext_m(it), 3))
                    pos = eixo.posicao(lim)
                    if pos:
                        acao(it, f"lat_{ponta}", it.get(f"lat_{ponta}"), _r(pos[0]), "coordenada do novo limite",
                             "eixo no km do limite", "alta", RECORTADO)
                        acao(it, f"lon_{ponta}", it.get(f"lon_{ponta}"), _r(pos[1]), "coordenada do novo limite",
                             "eixo no km do limite", "alta", RECORTADO)

    # ---------------------------------------------- 2. coordenadas × km
    def continuo(it) -> bool:
        a, b = it.get("km_ini"), it.get("km_fim")
        if b is None:
            return False
        d = it.get("dados", {})
        ext = d.get("extensao")
        if ext and abs(ext - abs(b - a) * 1000) <= max(2.0, 0.01 * ext):
            return True
        for o in por_cat[it["categoria"]]:
            if o is it or o.get("km_fim") is None:
                continue
            if abs(o["km_fim"] - a) < 0.03 or abs(o["km_ini"] - b) < 0.03:
                return True
        return False

    for it in itens:
        if it["categoria"] == "SEG" or it["id"] in removidos:
            continue
        geom = CATEGORIAS[it["categoria"]]["geom"]
        pontas = ("ini", "fim") if it.get("km_fim") is not None else ("ini",)
        for p in pontas:
            k = novo[it["id"]].get(f"km_{p}")
            la, lo = it.get(f"lat_{p}"), it.get(f"lon_{p}")
            esp = eixo.posicao(k) if k is not None else None
            if esp is None:
                continue
            if la is None or lo is None:
                acao(it, f"lat_{p}", la, _r(esp[0]), "sem coordenada", "eixo no km declarado", "média", CONFERIR,
                     preencher=True)
                acao(it, f"lon_{p}", lo, _r(esp[1]), "sem coordenada", "eixo no km declarado", "média", CONFERIR,
                     preencher=True)
                continue
            d = dist_km(esp, (la, lo))
            if d <= TOL_OK:
                continue
            pr = eixo.projetar(la, lo)
            lateral = pr[1] if pr else 99
            # a) erro de digitação na coordenada
            if lateral > TOL_DIGITO:
                fix = corrigir_digitacao(la, lo, esp)
                if fix:
                    _, nla, nlo, m = fix
                    if nla != la:
                        acao(it, f"lat_{p}", la, _r(nla), f"coordenada a {d:.1f} km da rodovia", m, "alta", CORRIGIDO)
                    if nlo != lo:
                        acao(it, f"lon_{p}", lo, _r(nlo), f"coordenada a {d:.1f} km da rodovia", m, "alta", CORRIGIDO)
                    continue
            if d <= TOL_GRAVE:
                acao(it, f"lat_{p}/lon_{p}", f"{la}, {lo}", None,
                     f"coordenada e km divergem {d:.1f} km (coordenada corresponde ao km ≈{pr[0]:.3f})" if pr else
                     f"coordenada a {d:.1f} km do eixo", "sem como decidir qual está certo", "—", VERIFICAR,
                     km_pela_coordenada=round(pr[0], 3) if pr else None)
                continue
            # b) km confiável (continuidade / extensão) ou coordenada repetida de outro local → refaz a coordenada
            copiada = pr is not None and lateral <= TOL_DIGITO and any(
                abs(kk - pr[0]) < 0.6 for kk in coords_por_ponto.get((round(la, 5), round(lo, 5)), ()))
            if lateral > TOL_DIGITO:
                motivo, conf = f"coordenada fora da rodovia ({lateral:.1f} km do eixo)", "média"
            elif geom == "linha" and continuo(it):
                outra = "fim" if p == "ini" else "ini"
                perto_outra = pr and novo[it["id"]].get(f"km_{outra}") is not None and abs(pr[0] - novo[it["id"]][f"km_{outra}"]) < 0.6
                motivo = (f"coordenada {'inicial' if p == 'ini' else 'final'} registrada no km ≈{pr[0]:.3f}"
                          + (" (junto da outra ponta)" if perto_outra else "")
                          + "; km confirmado pela continuidade/extensão")
                conf = "alta"
            elif copiada:
                motivo, conf = f"coordenada repetida de outro ponto (km ≈{pr[0]:.3f})", "alta"
            else:
                motivo, conf = (f"coordenada corresponde ao km ≈{pr[0]:.3f}, a {d:.1f} km do km declarado; "
                                "mantido o km e reposicionado pelo eixo"), "média"
            st = CORRIGIDO if conf == "alta" else CONFERIR
            acao(it, f"lat_{p}", la, _r(esp[0]), motivo, "eixo no km declarado", conf, st)
            acao(it, f"lon_{p}", lo, _r(esp[1]), motivo, "eixo no km declarado", conf, st)

    # ---------------------------------------------- 3. SNV
    seg = tabela_seg(itens)
    fonte = "cadastro de segmentos do inventário"
    ref = seg
    obs = []
    if snv_oficial:
        ref = [s for s in snv_oficial if K0 is None or (s["km_fim"] > K0 - 0.001 and s["km_ini"] < K1 + 0.001)]
        fonte = f"SNV oficial{(' ' + versao_oficial) if versao_oficial else ''}"
    concord = 0
    n_av = 0
    for it in itens:
        if it["categoria"] == "SEG" or it["id"] in removidos or not it.get("snv") or it.get("km_ini") is None:
            continue
        n_av += 1
        e = esperado(ref, novo[it["id"]]["km_ini"], novo[it["id"]].get("km_fim"))
        if set(codigos(it["snv"])) == set(e):
            concord += 1
    taxa = concord / n_av if n_av else 1.0
    inferida = None
    if not snv_oficial and taxa < 0.6:
        inferida = inferir_tabela(itens)
        obs.append(f"Apenas {taxa:.0%} dos itens usam o código SNV do cadastro de segmentos para o seu km. "
                   "O inventário segue outra numeração (provavelmente outra versão do SNV). Os códigos dos itens foram "
                   "verificados contra a numeração que eles próprios usam (limites inferidos abaixo); o cadastro de "
                   "segmentos precisa ser atualizado com a tabela oficial do SNV vigente.")
        ref, fonte = inferida, "numeração usada nos itens (limites inferidos)"
    validos = {s["codigo"] for s in ref}
    for it in itens:
        if it["categoria"] in SEM_SNV or it["id"] in removidos or it.get("km_ini") is None:
            continue
        atual = codigos(it.get("snv"))
        e = esperado(ref, novo[it["id"]]["km_ini"], novo[it["id"]].get("km_fim"))
        if not e or set(atual) == set(e):
            if it.get("snv") and atual and _fmt_cods(atual) != it["snv"].strip() and set(atual) == set(e):
                if any(c not in it["snv"].upper().replace(" ", "") for c in atual):
                    acao(it, "snv", it["snv"], _fmt_cods(e), "código SNV incompleto (ex.: falta 'B')",
                         "normalização do código", "alta", CORRIGIDO)
            continue
        if not atual:
            acao(it, "snv", it.get("snv"), _fmt_cods(e), "SNV não informado", f"pelo km ({fonte})", "alta", CORRIGIDO)
            continue
        invalido = [c for c in atual if c not in validos]
        if inferida is not None and not invalido:
            # numeração inferida: só corrige quando o item está claramente fora (> 1 km) do código declarado
            a, b = novo[it["id"]]["km_ini"], novo[it["id"]].get("km_fim") or novo[it["id"]]["km_ini"]
            lim = [s for s in ref if s["codigo"] in atual]
            if lim and min(abs((a + b) / 2 - max(s["km_ini"], min(s["km_fim"], (a + b) / 2))) for s in lim) < 1.0:
                continue
        if invalido and len(atual) == 1:
            motivo = f"código {invalido[0]} não existe no {fonte}"
        elif set(atual) < set(e):
            motivo = "item atravessa mais de um segmento SNV; faltava código"
        else:
            motivo = f"km {novo[it['id']]['km_ini']} pertence a {_fmt_cods(e)}"
        conf = "alta" if snv_oficial or inferida is None else "média"
        acao(it, "snv", it["snv"], _fmt_cods(e), motivo, f"pelo km ({fonte})", conf,
             CORRIGIDO if conf == "alta" else CONFERIR)

    # cadastro de segmentos: continuidade e coordenadas das divisas
    segs = sorted((i for i in itens if i["categoria"] == "SEG"), key=lambda i: i.get("km_ini") or 0)
    for s1, s2 in zip(segs, segs[1:]):
        if s1.get("km_fim") is not None and s2.get("km_ini") is not None and abs(s1["km_fim"] - s2["km_ini"]) > 0.001:
            acao(s2, "km_ini", s2["km_ini"], None, f"descontinuidade no cadastro de segmentos ({s1['km_fim']} → {s2['km_ini']})",
                 "conferir com o SNV oficial", "—", VERIFICAR)
        if None not in (s1.get("lat_fim"), s2.get("lat_ini")):
            d = dist_km((s1["lat_fim"], s1["lon_fim"]), (s2["lat_ini"], s2["lon_ini"]))
            if d > 0.1:
                acao(s2, "lat_ini/lon_ini", f"{s2['lat_ini']}, {s2['lon_ini']}", None,
                     f"coordenada da divisa difere {d * 1000:.0f} m da coordenada final do segmento anterior "
                     f"({s1['lat_fim']}, {s1['lon_fim']})", "conferir com o SNV oficial", "—", VERIFICAR)
    if snv_oficial:
        tab_of = {s["codigo"]: s for s in ref}
        tab_inv = {s["codigo"]: s for s in seg}
        for cod in sorted(set(tab_of) | set(tab_inv), key=lambda c: (tab_of.get(c) or tab_inv.get(c))["km_ini"]):
            o, v = tab_of.get(cod), tab_inv.get(cod)
            if o and v and (abs(o["km_ini"] - v["km_ini"]) > 0.001 or abs(o["km_fim"] - v["km_fim"]) > 0.001):
                obs.append(f"{cod}: inventário km {v['km_ini']}–{v['km_fim']} × oficial km {o['km_ini']}–{o['km_fim']}")
            elif o and not v:
                obs.append(f"{cod}: consta no SNV oficial (km {o['km_ini']}–{o['km_fim']}) e não no cadastro de segmentos")
            elif v and not o:
                obs.append(f"{cod}: está no cadastro de segmentos e não consta no SNV oficial")

    cont = Counter(a["situacao"] for a in acoes)
    itens_afetados = defaultdict(set)
    for a in acoes:
        itens_afetados[a["situacao"]].add(a["id"])
    return dict(
        acoes=acoes,
        resumo={k: dict(acoes=cont[k], itens=len(itens_afetados[k])) for k in
                (CORRIGIDO, CONFERIR, RECORTADO, REMOVIDO, VERIFICAR)},
        snv=dict(versao_inventario=projeto.get("snv_versao"), fonte_referencia=fonte,
                 versao_oficial=versao_oficial, concordancia=round(taxa, 3),
                 cadastro_segmentos=[{k: s[k] for k in ("codigo", "km_ini", "km_fim")} for s in seg],
                 tabela_referencia=[{k: s.get(k) for k in ("codigo", "km_ini", "km_fim")} for s in ref],
                 observacoes=obs),
        tolerancias=dict(aceitavel_km=TOL_OK, erro_km=TOL_GRAVE, digitacao_km=TOL_DIGITO),
    )


def aplicar(db, pid: int, analise: dict, atualizar_segmentos: list[dict] | None = None) -> dict:
    """Grava as ações (exceto 'verificar em campo'). Guarda o valor original em dados['_original']."""
    por_item: dict[int, list] = defaultdict(list)
    for a in analise["acoes"]:
        por_item[a["id"]].append(a)
    alterados, removidos = 0, []
    for iid, acs in por_item.items():
        if any(a["situacao"] == REMOVIDO for a in acs):
            removidos.append(iid)
            continue
        mud: dict[str, Any] = {}
        notas = []
        for a in acs:
            if a["situacao"] in (CORRIGIDO, CONFERIR, RECORTADO) and a["campo"] in (
                    "km_ini", "km_fim", "lat_ini", "lon_ini", "lat_fim", "lon_fim", "snv"):
                mud[a["campo"]] = a["para"]
                notas.append(f"{a['campo']}: {a['de']} → {a['para']} ({a['metodo']})")
        if not mud:
            continue
        it = db.item(iid)
        if not it:
            continue
        d = dict(it["dados"])
        orig = dict(d.get("_original") or {})
        for k in mud:
            orig.setdefault(k, it.get(k))
        d["_original"] = orig
        d["_correcao"] = "; ".join(notas)
        if any(k.startswith("km_") for k in mud):
            for k in ("extensao", "extensao_km", "km_faixa", "km_acost", "area", "area_ha", "area_caiacao",
                      "area_capina", "quantidade"):
                if k in d:
                    d.setdefault("_original", {})
                    d["_original"].setdefault(k, d[k])
                    d.pop(k)
        mud["dados"] = d
        db.atualizar_item(iid, mud)
        alterados += 1
    db.excluir_itens(removidos)
    return dict(alterados=alterados, removidos=len(removidos))


# ------------------------------------------------------------------ saídas
CORES = {CORRIGIDO: "FFF59D", CONFERIR: "FFCC80", RECORTADO: "B3E5FC", REMOVIDO: "EF9A9A", VERIFICAR: "F48FB1",
         "padronizado": "C8E6C9", "padronizado — conferir": "FFE0B2", "fora do padrão — verificar": "F8BBD0"}
CAMPOS_EDITAVEIS = ("km_ini", "km_fim", "lat_ini", "lon_ini", "lat_fim", "lon_fim", "snv", "tipo", "linhas", "dimensao")
_SO_DESTACAR = {REMOVIDO, VERIFICAR, "fora do padrão — verificar"}


def _valor(v):
    """Valores do log vêm como texto; devolve número quando for número."""
    if isinstance(v, str):
        n = nz.numero(v) if re.fullmatch(r"-?\d+(\.\d+)?", v.strip()) else None
        return n if n is not None else v
    return v


def gravar_no_original(arquivo: str, saida: str, entradas: list[dict]) -> dict:
    """Grava correções/padronizações numa cópia da planilha de inventário original e pinta as células:
    amarelo = corrigido · laranja = corrigido (conferir) · azul = recortado no limite · vermelho = fora do trecho ·
    rosa = verificar em campo · verde = nomenclatura padronizada · rosa-claro = fora do padrão (verificar)."""
    from .importador import mapa_colunas
    from .xlsx_patch import XlsxPatcher

    acoes = [a for a in entradas if a.get("origem") and "!" in a["origem"]]
    abas = {a["origem"].rsplit("!", 1)[0] for a in acoes}
    mapa = mapa_colunas(arquivo, abas)
    p = XlsxPatcher(arquivo)
    gravadas = sem_coluna = 0
    km_novos: dict[str, dict] = defaultdict(dict)
    for a in acoes:
        aba, linha = a["origem"].rsplit("!", 1)
        if aba not in p.sheets:
            continue
        cols = mapa.get(aba, {})
        cor = CORES.get(a["situacao"], "FFF59D")
        if a["situacao"] in _SO_DESTACAR or a.get("para") in (None, ""):
            if a["situacao"] == REMOVIDO:
                campos = [c for c in ("km_ini", "km_fim", "lat_ini", "lon_ini", "lat_fim", "lon_fim", "snv") if c in cols]
            else:
                campos = [x for x in (a.get("campo") or "").split("/") if x in cols]
            for campo in campos:
                col = cols[campo][-1] if campo.startswith("km") else cols[campo][0]
                p.destacar(aba, f"{col}{linha}", cor)
            continue
        if a["campo"] not in CAMPOS_EDITAVEIS or a["campo"] not in cols:
            sem_coluna += 1
            continue
        col = cols[a["campo"]][-1] if a["campo"].startswith("km") else cols[a["campo"]][0]
        p.set(aba, f"{col}{linha}", _valor(a["para"]))
        p.destacar(aba, f"{col}{linha}", cor)
        gravadas += 1
        if a["campo"].startswith("km"):
            km_novos[a["origem"]][a["campo"]] = _valor(a["para"])
            km_novos[a["origem"]]["_ref"] = a
    # extensão / quantidade das linhas cujo km mudou (somente se não forem fórmulas)
    for origem, km in km_novos.items():
        aba, linha = origem.rsplit("!", 1)
        cols = mapa.get(aba, {})
        a = km["_ref"]
        ki = km.get("km_ini", _valor(a.get("km_ini")))
        kf = km.get("km_fim", _valor(a.get("km_fim")))
        if ki is None or kf is None:
            continue
        ext = round(abs(kf - ki) * 1000, 3)
        cor = CORES.get(a["situacao"], "FFF59D")
        if "extensao" in cols:
            ref = f"{cols['extensao'][0]}{linha}"
            atual = p.get(aba, ref)
            if not (isinstance(atual, str) and atual.startswith("=")):
                p.set(aba, ref, ext)
                p.destacar(aba, ref, cor)
        if "quantidade" in cols and "cadencia" in cols:
            ref = f"{cols['quantidade'][0]}{linha}"
            atual = p.get(aba, ref)
            cad = nz.numero(p.get(aba, f"{cols['cadencia'][0]}{linha}"))
            if cad and not (isinstance(atual, str) and atual.startswith("=")):
                p.set(aba, ref, round(ext * cad, 3))
                p.destacar(aba, ref, cor)
    p.save(saida)
    return dict(arquivo=saida, celulas=gravadas, sem_coluna=sem_coluna)


def _maps(lat, lon):
    if lat is None or lon is None:
        return None
    return f"https://www.google.com/maps?q={lat},{lon}"


def relatorio_xlsx(projeto: dict, analise: dict, saida: str, itens_por_id: dict | None = None) -> str:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    azul = PatternFill("solid", fgColor="1F3A5F")
    branco = Font(bold=True, color="FFFFFF")

    def cab(ws, linha, titulos, larguras):
        for j, (t, w) in enumerate(zip(titulos, larguras), 1):
            c = ws.cell(linha, j, t)
            c.fill, c.font = azul, branco
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = ws.cell(linha + 1, 1)

    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "VERIFICAÇÃO DO SNV E CORREÇÃO DE PONTOS DO INVENTÁRIO"
    ws["A1"].font = Font(bold=True, size=13)
    linhas = [("Rodovia", projeto.get("rodovia")), ("Segmento", projeto.get("segmento")),
              ("Trecho (km)", f"{projeto.get('km_ini')} a {projeto.get('km_fim')}"),
              ("Versão SNV declarada no inventário", analise["snv"]["versao_inventario"]),
              ("Referência usada para o SNV", analise["snv"]["fonte_referencia"]),
              ("Itens com SNV coerente com o cadastro de segmentos", f"{analise['snv']['concordancia']:.1%}"),
              ("Tolerância km × coordenada (aceitável / erro)",
               f"{analise['tolerancias']['aceitavel_km']} km / {analise['tolerancias']['erro_km']} km")]
    for i, (a, b) in enumerate(linhas, 3):
        ws.cell(i, 1, a).font = Font(bold=True)
        ws.cell(i, 2, b)
    r = 3 + len(linhas) + 1
    cab(ws, r, ["Situação", "Itens", "Alterações", "Significado"], [44, 12, 12, 90])
    sig = {CORRIGIDO: "Erro identificado e corrigido com segurança (digitação, coordenada trocada de ponta, SNV).",
           CONFERIR: "Valor estimado pelo eixo da rodovia / km; recomenda-se conferir.",
           RECORTADO: "Feição cortada no limite do trecho (km e coordenada do limite).",
           REMOVIDO: "Feição inteiramente fora do trecho; retirada do inventário do segmento.",
           VERIFICAR: "Não foi possível determinar o valor correto; verificar em campo. Nada foi alterado."}
    for k, v in analise["resumo"].items():
        r += 1
        ws.cell(r, 1, k)
        ws.cell(r, 2, v["itens"])
        ws.cell(r, 3, v["acoes"])
        ws.cell(r, 4, sig[k])
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=CORES[k])
    r += 2
    ws.cell(r, 1, "Observações sobre o SNV").font = Font(bold=True)
    for o in analise["snv"]["observacoes"] or ["Códigos SNV dos itens coerentes com o cadastro de segmentos."]:
        r += 1
        ws.cell(r, 1, o).alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 45

    titulos = ["Situação", "Confiança", "Categoria", "Origem (aba!linha)", "km inicial", "km final", "Campo",
               "Valor original", "Valor corrigido", "Motivo", "Método", "Mapa (original)", "Mapa (corrigido)"]
    larg = [20, 10, 24, 34, 10, 10, 14, 22, 22, 70, 36, 18, 18]
    grupos = [("Corrigidos", (CORRIGIDO, CONFERIR)), ("Fora do trecho", (RECORTADO, REMOVIDO)),
              ("Verificar em campo", (VERIFICAR,))]
    for nome, sits in grupos:
        ws = wb.create_sheet(nome)
        cab(ws, 1, titulos, larg)
        lin = 1
        for a in analise["acoes"]:
            if a["situacao"] not in sits:
                continue
            lin += 1
            it = (itens_por_id or {}).get(a["id"], {})
            p = "ini" if a["campo"].endswith("ini") else "fim" if a["campo"].endswith("fim") else "ini"
            orig = _maps(it.get(f"lat_{p}"), it.get(f"lon_{p}"))
            corr = None
            if a["campo"].startswith(("lat", "lon")) and a["para"] is not None:
                outro = ("lon_" if a["campo"].startswith("lat") else "lat_") + p
                par = next((b["para"] for b in analise["acoes"] if b["id"] == a["id"] and b["campo"] == outro
                            and b["situacao"] == a["situacao"]), it.get(outro))
                corr = _maps(a["para"], par) if a["campo"].startswith("lat") else _maps(par, a["para"])
            vals = [a["situacao"], a["confianca"], a["nome_categoria"], a["origem"], a["km_ini"], a["km_fim"],
                    a["campo"], a["de"], a["para"], a["motivo"], a["metodo"], orig, corr]
            for j, v in enumerate(vals, 1):
                c = ws.cell(lin, j, v)
                if j in (12, 13) and v:
                    c.hyperlink = v
                    c.value = "abrir mapa"
                    c.font = Font(color="1F4E79", underline="single")
            ws.cell(lin, 1).fill = PatternFill("solid", fgColor=CORES[a["situacao"]])
        if nome == "Verificar em campo":
            lin += 2
            ws.cell(lin, 1, "Para cada ponto: confirmar no local o km e a coordenada; se o km pela coordenada "
                            "(coluna Motivo) estiver certo, corrigir o km; caso contrário, refazer a coordenada.")

    ws = wb.create_sheet("SNV")
    ws["A1"] = "Cadastro de segmentos do inventário"
    ws["A1"].font = Font(bold=True)
    cab(ws, 2, ["Código", "km inicial", "km final"], [16, 12, 12])
    for i, s in enumerate(analise["snv"]["cadastro_segmentos"], 3):
        ws.cell(i, 1, s["codigo"]); ws.cell(i, 2, s["km_ini"]); ws.cell(i, 3, s["km_fim"])
    ws["E1"] = f"Referência: {analise['snv']['fonte_referencia']}"
    ws["E1"].font = Font(bold=True)
    for j, t in enumerate(["Código", "km inicial", "km final"], 5):
        c = ws.cell(2, j, t)
        c.fill, c.font = azul, branco
        ws.column_dimensions[get_column_letter(j)].width = 14
    for i, s in enumerate(analise["snv"]["tabela_referencia"], 3):
        ws.cell(i, 5, s["codigo"]); ws.cell(i, 6, s["km_ini"]); ws.cell(i, 7, s["km_fim"])
    wb.save(saida)
    return saida
